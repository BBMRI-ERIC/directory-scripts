"""Read and validate SO2 descriptive-survey workbooks without external services."""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from hashlib import sha256
import json
from html import unescape
from math import atan2, ceil, cos, degrees, radians, sin
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from types import MappingProxyType
from typing import Any, Literal, Mapping, Sequence
from unicodedata import normalize
from zipfile import BadZipFile
from xml.etree import ElementTree as ET

import openpyxl
import pandas as pd


class InputError(Exception):
    """Raised when a workbook or its descriptive schema is invalid."""


QuestionType = Literal["single_choice", "multi_choice", "ordinal", "free_text"]
_QUESTION_TYPES = frozenset({"single_choice", "multi_choice", "ordinal", "free_text"})
_UPSET_SINGLE_DEFINITIONS = ("q_009", "q_012", "q_018", "q_035", "q_042", "q_050")
_UPSET_SHARED_DEFINITIONS = {
    "q_042_q_044": ("q_042", "q_044"),
    "q_035_q_050": ("q_035", "q_050"),
}


@dataclass(frozen=True)
class QuestionDefinition:
    """Validated classification and presentation metadata for one survey question.

    ``parent_context_values`` maps a free-text parent column to the canonical
    structured selections that semantically enable that follow-up. An absent
    mapping retains the full parent response as context rather than inventing a
    routing rule.
    """

    question_id: str
    column: str
    question_type: QuestionType
    label: str
    categories: tuple[str, ...]
    delimiter: str | None
    parent_columns: tuple[str, ...]
    applicability: Mapping[str, Any] | None
    parent_context_values: Mapping[str, tuple[str, ...]] = field(
        default_factory=lambda: MappingProxyType({})
    )
    exclusive_categories: tuple[str, ...] = ()
    category_aliases: Mapping[str, str] = field(default_factory=lambda: MappingProxyType({}))
    form_uid: str = ""


@dataclass(frozen=True)
class AssociationHeatmapDefinition:
    """Validated registry metadata for one approved association heatmap."""

    definition_id: str
    row_question_id: str
    column_question_id: str
    row_form_uid: str
    column_form_uid: str
    row_categories: tuple[str, ...]
    column_categories: tuple[str, ...]
    mode: Literal["default", "exploratory"]
    title: str
    interpretation: str
    conditional_summary: bool


@dataclass(frozen=True)
class SurveyWorkbook:
    """Validated workbook provenance and the retained response rows."""

    source_path: str
    source_sha256: str
    worksheet: str
    alias: str
    export_date: str
    header_row: int
    responses: pd.DataFrame
    total_data_rows: int
    excluded_blank_rows: int


def _reject_duplicate_json_keys(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    """Build a JSON object while rejecting ambiguous duplicate keys.

    Args:
        pairs: Decoder-supplied ordered key/value pairs for one JSON object.

    Returns:
        Dictionary preserving the decoded object values after duplicate checking.

    Raises:
        InputError: If a key occurs more than once in the same object."""
    result = {}
    for key, value in pairs:
        if key in result:
            raise InputError(f"Schema contains duplicate JSON key: {key!r}.")
        result[key] = value
    return result


def load_descriptive_schema(path: str | Path) -> dict[str, Any]:
    """Load a descriptive-workbook schema while rejecting duplicate JSON keys.

    Args:
        path: JSON schema path.

    Returns:
        Parsed object-shaped schema.

    Raises:
        InputError: If the file cannot be read, is malformed JSON, contains
            duplicate keys, or does not have an object root.
    """
    schema_path = Path(path)
    try:
        with schema_path.open(encoding="utf-8") as schema_file:
            schema = json.load(schema_file, object_pairs_hook=_reject_duplicate_json_keys)
    except OSError as exc:
        raise InputError(f"Could not read descriptive schema {schema_path}: {exc}") from exc
    except json.JSONDecodeError as exc:
        raise InputError(f"Invalid JSON in descriptive schema {schema_path}: {exc}") from exc
    if not isinstance(schema, dict):
        raise InputError("Descriptive schema root must be a JSON object.")
    return schema


def load_association_heatmap_registry(path: str | Path) -> tuple[AssociationHeatmapDefinition, ...]:
    """Load a syntactically valid versioned association-heatmap registry.

    Args:
        path: JSON registry path.

    Returns:
        Immutable association definitions in declared rendering order.

    Raises:
        InputError: If the registry cannot be read, is malformed, or violates
            the registry's self-contained JSON contract. Question semantics are
            intentionally checked separately by ``validate_association_definitions``.
    """
    registry_path = Path(path)
    try:
        with registry_path.open(encoding="utf-8") as registry_file:
            registry = json.load(registry_file, object_pairs_hook=_reject_duplicate_json_keys)
    except OSError as exc:
        raise InputError(f"Could not read association registry {registry_path}: {exc}") from exc
    except json.JSONDecodeError as exc:
        raise InputError(f"Invalid JSON in association registry {registry_path}: {exc}") from exc

    root = _required_mapping(registry, "association registry")
    if root.get("schema_version") != "1":
        raise InputError(f"Unknown association registry version: {root.get('schema_version')!r}.")
    definitions_raw = root.get("definitions")
    if not isinstance(definitions_raw, list):
        raise InputError("Association registry field definitions must be an array.")

    definitions = []
    definition_ids = set()
    question_pairs = set()
    required_fields = (
        "definition_id", "row_question_id", "column_question_id", "row_form_uid",
        "column_form_uid", "row_categories", "column_categories", "mode", "title",
        "interpretation", "conditional_summary",
    )
    for position, raw_definition in enumerate(definitions_raw):
        definition = _required_mapping(raw_definition, f"association definitions[{position}]")
        for name in required_fields:
            if name not in definition:
                raise InputError(
                    f"Association definition {position} is missing required field {name}."
                )
        definition_id = _required_text(
            definition["definition_id"], f"association definitions[{position}].definition_id"
        )
        row_question_id = _required_text(
            definition["row_question_id"], f"association definitions[{position}].row_question_id"
        )
        column_question_id = _required_text(
            definition["column_question_id"],
            f"association definitions[{position}].column_question_id",
        )
        if row_question_id == column_question_id:
            raise InputError("Association definition must contain two different question IDs.")
        if definition_id in definition_ids:
            raise InputError(f"Association registry has duplicate definition ID {definition_id!r}.")
        pair = tuple(sorted((row_question_id, column_question_id)))
        if pair in question_pairs:
            raise InputError("Association registry has duplicate unordered question pair.")
        mode = _required_text(definition["mode"], f"association definitions[{position}].mode")
        if mode not in {"default", "exploratory"}:
            raise InputError(f"Association definition has unknown mode {mode!r}.")
        conditional_summary = definition["conditional_summary"]
        if not isinstance(conditional_summary, bool):
            raise InputError(
                f"Association definitions[{position}].conditional_summary must be boolean."
            )
        definitions.append(AssociationHeatmapDefinition(
            definition_id=definition_id,
            row_question_id=row_question_id,
            column_question_id=column_question_id,
            row_form_uid=_required_text(
                definition["row_form_uid"], f"association definitions[{position}].row_form_uid"
            ),
            column_form_uid=_required_text(
                definition["column_form_uid"],
                f"association definitions[{position}].column_form_uid",
            ),
            row_categories=_text_sequence(
                definition["row_categories"], f"association definitions[{position}].row_categories"
            ),
            column_categories=_text_sequence(
                definition["column_categories"],
                f"association definitions[{position}].column_categories",
            ),
            mode=mode,
            title=_required_text(definition["title"], f"association definitions[{position}].title"),
            interpretation=_required_text(
                definition["interpretation"], f"association definitions[{position}].interpretation"
            ),
            conditional_summary=conditional_summary,
        ))
        definition_ids.add(definition_id)
        question_pairs.add(pair)
    return tuple(definitions)


def validate_association_definitions(
    definitions: Sequence[AssociationHeatmapDefinition],
    questions_by_id: Mapping[str, QuestionDefinition],
) -> None:
    """Validate association definitions against already validated survey questions.

    Args:
        definitions: Syntactically valid association registry definitions.
        questions_by_id: Validated descriptive questions indexed by their stable IDs.

    Returns:
        ``None`` after every definition matches the descriptive schema.

    Raises:
        InputError: If a question is unknown, unsuitable for an association,
            prohibited, or differs in form UID or declared category order.
    """
    for definition in definitions:
        for axis, question_id, form_uid, categories in (
            ("row", definition.row_question_id, definition.row_form_uid, definition.row_categories),
            ("column", definition.column_question_id, definition.column_form_uid,
             definition.column_categories),
        ):
            question = questions_by_id.get(question_id)
            if question is None:
                raise InputError(
                    f"Association definition {definition.definition_id!r} references unknown {axis} "
                    f"question {question_id!r}."
                )
            if question.question_type not in {"single_choice", "ordinal"}:
                raise InputError(
                    f"Association {axis} question {question_id!r} must be single_choice or ordinal."
                )
            if question_id == "q_111" or question_id.startswith("q_111_"):
                raise InputError("Association definitions must not include q_111 contact questions.")
            if question.form_uid != form_uid:
                raise InputError(
                    f"Association {axis} question {question_id!r} has mismatched form UID."
                )
            if question.categories != categories:
                raise InputError(
                    f"Association {axis} question {question_id!r} has mismatched category sequence."
                )


def _form_label(value: str) -> str:
    """Normalize an XML/form label for stable schema matching.

    Args:
        value: Literal question or answer text from XML or the descriptive schema.

    Returns:
        Case-folded alphanumeric comparison key with markup and whitespace removed.
    """
    previous = value
    while True:
        decoded = unescape(previous)
        if decoded == previous:
            break
        previous = decoded
    plain = re.sub(r"<[^>]+>", " ", decoded)
    return re.sub(r"[^\w]+", "", plain, flags=re.UNICODE).casefold()


def load_form_structure(path: str | Path) -> dict[str, Any]:
    """Load authoritative response types and dependencies from an SO2 form XML export.

    Args:
        path: XML form-definition export containing ``Survey/Elements``.

    Returns:
        Source provenance and question metadata indexed by normalized form label.

    Raises:
        InputError: If the XML cannot be read or lacks a consistent Elements definition.
    """
    source = Path(path)
    try:
        source_bytes = source.read_bytes()
        root = ET.fromstring(source_bytes)
    except (OSError, ET.ParseError) as exc:
        raise InputError(f"Could not read form-structure XML {source}: {exc}") from exc
    elements = root.find("./Survey/Elements")
    if elements is None:
        raise InputError("Form-structure XML lacks Survey/Elements.")
    questions_by_id: dict[str, dict[str, Any]] = {}
    answers: list[tuple[str, str, tuple[str, ...]]] = []
    current_question_id: str | None = None
    matrix_title: str | None = None

    def register_question(question_id: str | None, question_type: str | None, label: str) -> None:
        """Register one XML question after validating its identity and display text."""
        if not question_id or not question_type or not label:
            raise InputError("Form-structure XML has a question without id, type, or text.")
        if question_id in questions_by_id:
            raise InputError(f"Form-structure XML repeats question id {question_id!r}.")
        questions_by_id[question_id] = {
            "id": question_id,
            "label": label,
            "response_type": question_type,
            # This export contains no question-level requiredness attribute.
            "requiredness": "not declared by form export",
            "dependencies": [],
        }

    for element in elements:
        if element.tag == "MatrixTitle":
            matrix_title = (element.text or "").strip()
            current_question_id = None
        elif element.tag == "MatrixQuestion":
            row_label = (element.text or "").strip()
            if not matrix_title:
                raise InputError("Form-structure XML has a matrix row without a matrix title.")
            register_question(
                element.get("id"), element.get("type"), f"{matrix_title}: {row_label}",
            )
            current_question_id = None
        elif element.tag == "Question":
            label = (element.text or "").strip()
            register_question(element.get("id"), element.get("type"), label)
            current_question_id = element.get("id")
            matrix_title = None
        elif element.tag == "Answer" and current_question_id is not None:
            dependent = tuple(filter(
                None,
                re.split(r"[;,\s]+", element.get("dependentElements") or ""),
            ))
            answers.append((current_question_id, (element.text or "").strip(), dependent))
    for parent_id, answer_text, children in answers:
        parent = questions_by_id[parent_id]
        for child_id in children:
            child = questions_by_id.get(child_id)
            if child is None:
                raise InputError(f"Form-structure XML references missing dependent question {child_id!r}.")
            child["dependencies"].append({"parent_question": parent["label"], "answer": answer_text})
    by_label: dict[str, dict[str, Any]] = {}
    for item in questions_by_id.values():
        key = _form_label(item["label"])
        if key in by_label:
            raise InputError(f"Form-structure XML has duplicate normalized question text: {item['label']!r}.")
        by_label[key] = item
    return {"source_path": str(source), "source_sha256": sha256(source_bytes).hexdigest(), "questions": by_label}


def load_form_manifest_structure(path: str | Path) -> dict[str, Any]:
    """Load report-facing field metadata from a generated EUS JSON manifest.

    Args:
        path: Path to the checked-in JSON form manifest generated by
            ``export-eus-form-json``.

    Returns:
        Provenance and question metadata indexed by normalized report labels;
        matrix rows are exposed as individual report questions.

    Raises:
        InputError: If the manifest cannot be loaded or does not expose a
            unique label for each report-facing question.
    """
    try:
        from so2_eusurvey import FormInputError, load_form_manifest

        form = load_form_manifest(path)
    except FormInputError as exc:
        raise InputError(str(exc)) from exc
    response_types = {
        "free_text": "Free Text",
        "single_choice": "Single Choice",
        "multi_choice": "Multiple Choice",
        "matrix": "Single Choice Matrix Question",
    }
    questions: dict[str, dict[str, Any]] = {}

    def add_question(label: str, response_type: str, mandatory: bool, dependencies: list[dict[str, str]], matrix_parent: str | None = None, matrix_row: str | None = None) -> None:
        """Store one report-facing question after label uniqueness validation."""
        key = _form_label(label)
        if key in questions:
            raise InputError(f"Form manifest has duplicate normalized question text: {label!r}.")
        questions[key] = {
            "id": key,
            "label": label,
            "response_type": response_type,
            "requiredness": "mandatory" if mandatory else "optional",
            "dependencies": dependencies,
            "matrix_parent": matrix_parent,
            "matrix_row": matrix_row,
        }

    for field in form.fields_by_uid.values():
        dependencies = []
        if field.shown_when is not None:
            dependencies = []
            for atom in field.shown_when.atoms:
                parent = form.fields_by_uid.get(atom.parent_uid)
                choice = None if parent is None else next(
                    (item for item in parent.choices if item.uid == atom.choice_uid), None
                )
                if parent is None or choice is None:
                    raise InputError("Form manifest dependency does not resolve to readable labels.")
                dependencies.append({"parent_question": parent.title, "answer": choice.label})
        if field.field_type == "matrix":
            for row in field.matrix_rows:
                add_question(
                    f"{field.title}: {row.label}", response_types[field.field_type],
                    field.mandatory, dependencies, field.title, row.label,
                )
        else:
            add_question(field.title, response_types[field.field_type], field.mandatory, dependencies)
    return {
        "source_path": str(path),
        "source_sha256": sha256(Path(path).read_bytes()).hexdigest(),
        "archive_sha256": form.archive_sha256,
        "active_member_sha256": form.active_member_sha256,
        "questions": questions,
    }


def _required_mapping(value: Any, name: str) -> Mapping[str, Any]:
    """Validate one required object-shaped schema value.

    Args:
        value: Candidate value supplied by the parsed schema.
        name: Fully qualified schema field name used in error messages.

    Returns:
        The same value narrowed to a mapping.

    Raises:
        InputError: If the value is not mapping-shaped."""
    if not isinstance(value, Mapping):
        raise InputError(f"Schema field {name} must be an object.")
    return value


def _required_text(value: Any, name: str) -> str:
    """Validate one required nonblank schema string.

    Args:
        value: Candidate value supplied by the parsed schema.
        name: Fully qualified schema field name used in error messages.

    Returns:
        The original nonblank string.

    Raises:
        InputError: If the value is not a nonblank string."""
    if not isinstance(value, str) or not value.strip():
        raise InputError(f"Schema field {name} must be a nonblank string.")
    return value


def _text_sequence(value: Any, name: str) -> tuple[str, ...]:
    """Validate an ordered sequence of unique schema strings.

    Args:
        value: Candidate JSON array from the schema.
        name: Fully qualified schema field name used in error messages.

    Returns:
        Immutable ordered strings after nonblank and duplicate validation.

    Raises:
        InputError: If the value is not an array of unique nonblank strings."""
    if not isinstance(value, list):
        raise InputError(f"Schema field {name} must be an array of strings.")
    values = tuple(_required_text(item, name) for item in value)
    if len(set(values)) != len(values):
        raise InputError(f"Schema field {name} must not contain duplicate columns.")
    return values


def _is_blank(value: Any) -> bool:
    """Classify one spreadsheet scalar as blank or answered.

    Args:
        value: Raw cell value, including pandas missing-value sentinels.

    Returns:
        ``True`` only when the value is absent, whitespace-only, or pandas-null."""
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip()
    return bool(pd.isna(value))


def is_nonblank_response_row(row: pd.Series, classified_columns: Sequence[str]) -> bool:
    """Return whether context or question cells contain a nonblank response value.

    Args:
        row: One workbook data row.
        classified_columns: Context and question columns that define a response.

    Returns:
        True when at least one classified cell is nonblank.
    """
    return any(not _is_blank(row[column]) for column in classified_columns)


def _validate_question(question: Any, position: int) -> QuestionDefinition:
    """Validate one schema question and construct its immutable definition.

    Args:
        question: Candidate object from the schema ``questions`` array.
        position: Zero-based question position used for contextual validation errors.

    Returns:
        Validated immutable definition with normalized optional fields.

    Raises:
        InputError: If required fields, category rules, aliases, or options are invalid."""
    item = _required_mapping(question, f"questions[{position}]")
    required_fields = ("question_id", "column", "question_type", "label", "categories")
    for field in required_fields:
        if field not in item:
            raise InputError(f"Schema question {position} is missing required field {field}.")

    question_type = _required_text(item["question_type"], f"questions[{position}].question_type")
    if question_type not in _QUESTION_TYPES:
        raise InputError(
            f"Schema question {position} has invalid question_type {question_type!r}."
        )
    categories = _text_sequence(item["categories"], f"questions[{position}].categories")
    if question_type != "free_text" and not categories:
        raise InputError(f"Schema question {position} requires at least one category.")
    if question_type == "free_text" and categories:
        raise InputError(f"Schema free_text question {position} must not define categories.")

    delimiter = item.get("delimiter")
    if question_type == "multi_choice":
        delimiter = _required_text(delimiter, f"questions[{position}].delimiter")
    elif delimiter is not None:
        raise InputError(
            f"Schema question {position} may define delimiter only for multi_choice."
        )

    parent_columns = _text_sequence(
        item.get("parent_columns", []), f"questions[{position}].parent_columns"
    )
    parent_context_values_raw = _required_mapping(
        item.get("parent_context_values", {}),
        f"questions[{position}].parent_context_values",
    )
    parent_context_values = {
        _required_text(column, f"questions[{position}].parent_context_values key"):
        _text_sequence(values, f"questions[{position}].parent_context_values[{column!r}]")
        for column, values in parent_context_values_raw.items()
    }
    exclusive_categories = _text_sequence(
        item.get("exclusive_categories", []),
        f"questions[{position}].exclusive_categories",
    )
    if exclusive_categories and question_type != "multi_choice":
        raise InputError(
            f"Schema question {position} may define exclusive categories only for multi_choice."
        )
    unknown_exclusive = sorted(set(exclusive_categories) - set(categories))
    if unknown_exclusive:
        raise InputError(
            f"Schema question {position} has unknown exclusive category {unknown_exclusive[0]!r}.")
    aliases = _required_mapping(item.get("category_aliases", {}), f"questions[{position}].category_aliases")
    category_aliases = {
        _required_text(source, f"questions[{position}].category_aliases key"):
        _required_text(target, f"questions[{position}].category_aliases[{source!r}]")
        for source, target in aliases.items()
    }
    unknown_alias_targets = sorted(set(category_aliases.values()) - set(categories))
    if unknown_alias_targets:
        raise InputError(
            f"Schema question {position} aliases unknown category {unknown_alias_targets[0]!r}."
        )
    applicability = item.get("applicability")
    if applicability is not None:
        applicability = MappingProxyType(dict(_required_mapping(
            applicability, f"questions[{position}].applicability"
        )))
    return QuestionDefinition(
        question_id=_required_text(item["question_id"], f"questions[{position}].question_id"),
        column=_required_text(item["column"], f"questions[{position}].column"),
        question_type=question_type,
        label=_required_text(item["label"], f"questions[{position}].label"),
        categories=categories,
        delimiter=delimiter,
        parent_columns=parent_columns,
        parent_context_values=MappingProxyType(parent_context_values),
        applicability=applicability,
        exclusive_categories=exclusive_categories,
        category_aliases=MappingProxyType(category_aliases),
        form_uid=(
            _required_text(item["form_uid"], f"questions[{position}].form_uid")
            if "form_uid" in item else ""
        ),
    )


def validate_descriptive_schema(
    schema: Mapping[str, Any], columns: Sequence[str]
) -> list[QuestionDefinition]:
    """Validate the schema against source headers.

    Args:
        schema: Parsed descriptive schema.
        columns: Exact workbook headers.

    Returns:
        Immutable question definitions in report order.

    Raises:
        InputError: If the schema, classifications, routing, or references are invalid.
    """
    root = _required_mapping(schema, "schema")
    for field in ("schema_version", "input", "columns", "questions", "output"):
        if field not in root:
            raise InputError(f"Schema is missing required field {field}.")
    if root["schema_version"] != "1":
        raise InputError(f"Unknown descriptive schema version: {root['schema_version']!r}.")

    headers = tuple(columns)
    if any(not isinstance(header, str) or not header.strip() for header in headers):
        raise InputError("Workbook header row 4 contains a blank header.")
    if len(set(headers)) != len(headers):
        raise InputError("Workbook header row 4 contains nonunique headers.")

    input_schema = _required_mapping(root["input"], "input")
    for field in ("worksheet", "alias", "header_row"):
        if field not in input_schema:
            raise InputError(f"Schema input is missing required field {field}.")
    _required_text(input_schema["worksheet"], "input.worksheet")
    _required_text(input_schema["alias"], "input.alias")
    if input_schema["header_row"] != 4:
        raise InputError("Schema input.header_row must be 4.")

    column_schema = _required_mapping(root["columns"], "columns")
    for field in (
        "respondent_context", "institution_column", "country_column",
        "administrative_exclusions", "administrative_exclusion_reasons",
    ):
        if field not in column_schema:
            raise InputError(f"Schema columns is missing required field {field}.")
    context_columns = _text_sequence(
        column_schema["respondent_context"], "columns.respondent_context"
    )
    administrative_columns = _text_sequence(
        column_schema["administrative_exclusions"], "columns.administrative_exclusions"
    )
    institution_column = _required_text(
        column_schema["institution_column"], "columns.institution_column"
    )
    country_column = _required_text(
        column_schema["country_column"], "columns.country_column"
    )
    if institution_column not in context_columns:
        raise InputError("Schema columns.institution_column must be respondent context.")
    if country_column not in context_columns:
        raise InputError("Schema columns.country_column must be respondent context.")
    reasons = _required_mapping(
        column_schema["administrative_exclusion_reasons"],
        "columns.administrative_exclusion_reasons",
    )
    if set(reasons) != set(administrative_columns):
        missing = sorted(set(administrative_columns) - set(reasons))
        unexpected = sorted(set(reasons) - set(administrative_columns))
        column = missing[0] if missing else unexpected[0]
        raise InputError(
            f"Schema administrative exclusion reasons do not match administrative exclusion {column!r}."
        )
    for column in administrative_columns:
        _required_text(reasons[column], f"columns.administrative_exclusion_reasons.{column}")
    if not isinstance(root["questions"], list):
        raise InputError("Schema field questions must be an array.")
    questions = [_validate_question(question, index) for index, question in enumerate(root["questions"])]
    question_columns = {question.column for question in questions}
    question_by_column = {question.column: question for question in questions}
    for question in questions:
        if question.applicability is None:
            continue
        applicability_column = question.applicability.get("column")
        applicability_values = question.applicability.get("values")
        if not isinstance(applicability_column, str) or not applicability_column.strip():
            raise InputError(
                f"Schema applicability for {question.column!r} must declare a nonblank column."
            )
        if applicability_column not in question_by_column:
            raise InputError(
                "Schema applicability column must name a declared question: "
                f"{applicability_column!r}."
            )
        values = _text_sequence(
            applicability_values, f"applicability values for {question.column!r}"
        )
        if not values:
            raise InputError(
                f"Schema applicability for {question.column!r} requires at least one value."
            )
        invalid_values = sorted(
            set(values) - set(question_by_column[applicability_column].categories)
        )
        if invalid_values:
            raise InputError(
                f"Schema applicability value {invalid_values[0]!r} is not declared "
                f"for {applicability_column!r}."
            )
    for question in questions:
        if question.parent_context_values and question.question_type != "free_text":
            raise InputError(
                "Schema parent_context_values is allowed only for free-text questions: "
                f"{question.column!r}."
            )
        if question.question_type == "free_text":
            invalid_parents = sorted(set(question.parent_columns) - question_columns)
            if invalid_parents:
                raise InputError(
                    f"Schema free_text parent column must name a declared question column: {invalid_parents[0]!r}."
                )
            invalid_context_columns = sorted(
                set(question.parent_context_values) - set(question.parent_columns)
            )
            if invalid_context_columns:
                raise InputError(
                    "Schema free-text parent_context_values column must be a declared parent column: "
                    f"{invalid_context_columns[0]!r}."
                )
            for parent_column, values in question.parent_context_values.items():
                if not values:
                    raise InputError(
                        "Schema free-text parent_context_values must not be empty for "
                        f"{parent_column!r}."
                    )
                invalid_values = sorted(set(values) - set(question_by_column[parent_column].categories))
                if invalid_values:
                    raise InputError(
                        "Schema free-text parent_context_values value is not declared "
                        f"for {parent_column!r}: {invalid_values[0]!r}."
                    )

    classifications = (
        list(context_columns)
        + list(administrative_columns)
        + [question.column for question in questions]
    )
    duplicates = sorted({column for column in classifications if classifications.count(column) > 1})
    if duplicates:
        raise InputError(f"Source column classified more than once: {duplicates[0]!r}.")
    declared_columns = classifications + [
        parent_column for question in questions for parent_column in question.parent_columns
    ]
    missing = sorted(set(declared_columns) - set(headers))
    if missing:
        raise InputError(f"Schema-declared column is absent: {missing[0]!r}.")
    unclassified = sorted(set(headers) - set(classifications))
    if unclassified:
        raise InputError(f"Schema has unclassified source column: {unclassified[0]!r}.")

    output = _required_mapping(root["output"], "output")
    if "source_row_column" not in output:
        raise InputError("Schema output is missing required field source_row_column.")
    source_row_column = _required_text(output["source_row_column"], "output.source_row_column")
    if source_row_column in headers:
        raise InputError("Schema output.source_row_column must not collide with a source column.")
    return questions


def read_descriptive_workbook(
    path: str | Path, schema: Mapping[str, Any]
) -> SurveyWorkbook:
    """Read an SO2 workbook and retain only meaningful response rows.

    Args:
        path: Source XLSX path.
        schema: Parsed descriptive schema for the workbook.

    Returns:
        Validated workbook provenance and retained response rows.

    Raises:
        InputError: If the workbook, envelope, headers, or schema relationship is invalid.
    """
    source_path = Path(path)
    try:
        source_bytes = source_path.read_bytes()
    except OSError as exc:
        raise InputError(f"Could not read descriptive workbook {source_path}: {exc}") from exc

    try:
        workbook = openpyxl.load_workbook(source_path, read_only=True, data_only=True)
    except (
        OSError,
        ValueError,
        BadZipFile,
        openpyxl.utils.exceptions.InvalidFileException,
    ) as exc:
        raise InputError(f"Could not open descriptive workbook {source_path}: {exc}") from exc
    try:
        root = _required_mapping(schema, "schema")
        if "input" not in root:
            raise InputError("Schema is missing required field input.")
        input_schema = _required_mapping(root["input"], "input")
        worksheet_name = _required_text(input_schema.get("worksheet"), "input.worksheet")
        expected_alias = _required_text(input_schema.get("alias"), "input.alias")
        header_row = input_schema.get("header_row")
        if header_row != 4:
            raise InputError("Schema input.header_row must be 4.")
        if worksheet_name not in workbook.sheetnames:
            raise InputError(f"Workbook has wrong worksheet: expected {worksheet_name!r}.")
        sheet = workbook[worksheet_name]
        if sheet.cell(1, 1).value != "Alias" or sheet.cell(1, 2).value != expected_alias:
            raise InputError("Workbook row 1 alias does not match the descriptive schema.")
        export_date = sheet.cell(2, 2).value
        if sheet.cell(2, 1).value != "Export Date" or not isinstance(export_date, datetime):
            raise InputError("Workbook row 2 export date must be a labeled datetime.")
        export_date = export_date.isoformat()
        if any(cell.value is not None and str(cell.value).strip() for cell in sheet[3]):
            raise InputError("Workbook row 3 must be fully blank.")

        headers = [sheet.cell(4, column).value for column in range(1, sheet.max_column + 1)]
        if any(not isinstance(header, str) or not header.strip() for header in headers):
            raise InputError("Workbook header row 4 contains a blank header.")
        try:
            questions = validate_descriptive_schema(schema, headers)
        except InputError as exc:
            if "Schema-declared column is absent" in str(exc) or "unclassified source column" in str(exc):
                raise InputError("Workbook header row 4 does not match the descriptive schema.") from exc
            raise
        output = _required_mapping(_required_mapping(schema, "schema")["output"], "output")
        source_row_column = _required_text(output.get("source_row_column"), "output.source_row_column")

        records = []
        for source_row, values in enumerate(
            sheet.iter_rows(min_row=5, max_col=len(headers), values_only=True), start=5
        ):
            record = dict(zip(headers, values, strict=True))
            record[source_row_column] = source_row
            records.append(record)
    finally:
        workbook.close()

    data = pd.DataFrame.from_records(records, columns=[*headers, source_row_column])
    context_columns = tuple(schema["columns"]["respondent_context"])
    classified_response_columns = (*context_columns, *(question.column for question in questions))
    response_mask = data.apply(
        lambda row: is_nonblank_response_row(row, classified_response_columns), axis=1
    )
    responses = data.loc[response_mask].reset_index(drop=True)
    return SurveyWorkbook(
        source_path=str(source_path),
        source_sha256=sha256(source_bytes).hexdigest(),
        worksheet=worksheet_name,
        alias=expected_alias,
        export_date=export_date,
        header_row=4,
        responses=responses,
        total_data_rows=len(data.index),
        excluded_blank_rows=int((~response_mask).sum()),
    )


def _display_value(value: Any) -> str:
    """Convert one response value to its report-safe literal representation.

    Args:
        value: Raw spreadsheet value from a response or context column.

    Returns:
        ``Missing`` for blank values, otherwise the literal string representation."""
    return "Missing" if _is_blank(value) else str(value)


def _normalised_respondent_value(value: Any) -> str:
    """Normalize one respondent identity component for repeat detection only.

    Args:
        value: Raw country or institution value from a submitted response.

    Returns:
        Case-folded, whitespace-normalized display value used as a grouping key."""
    return " ".join(_display_value(value).split()).casefold()


def _percentage(count: int, base: int) -> float | None:
    """Calculate a count percentage when a denominator exists.

    Args:
        count: Nonnegative numerator count.
        base: Nonnegative denominator count.

    Returns:
        Percentage rounded to four decimals, or ``None`` when ``base`` is zero."""
    return None if base == 0 else round(count * 100 / base, 4)


def _applicability_summary(
    question: QuestionDefinition,
    responses: pd.DataFrame,
    question_by_column: Mapping[str, QuestionDefinition],
) -> dict[str, Any]:
    """Evaluate an explicitly declared question-routing rule.

    Args:
        question: Validated question whose optional applicability rule is evaluated.
        responses: Included survey response rows.
        question_by_column: Validated questions indexed by their source columns.

    Returns:
        Unknown status when no rule exists, otherwise row counts for each routing outcome.

    Raises:
        InputError: If the declared rule is malformed or references an absent column."""
    if question.applicability is None:
        return {"status": "unknown"}
    column = question.applicability.get("column")
    values = question.applicability.get("values")
    if not isinstance(column, str) or not isinstance(values, list) or not all(
        isinstance(value, str) for value in values
    ):
        raise InputError(
            f"Schema applicability for {question.column!r} must declare column and string values."
        )
    if column not in responses.columns:
        raise InputError(f"Schema applicability column is absent: {column!r}.")
    parent_question = question_by_column[column]
    applicable_values = frozenset(values)
    eligibility_unknown = responses[column].map(_is_blank)
    applicable = responses[column].map(
        lambda value: bool(
            applicable_values.intersection(_structured_answers(parent_question, value)[0])
        )
    ) & ~eligibility_unknown
    inapplicable = ~applicable & ~eligibility_unknown
    answered = ~responses[question.column].map(_is_blank)
    return {
        "status": "evaluated",
        "column": column,
        "values": values,
        "applicable_rows": int(applicable.sum()),
        "inapplicable_rows": int(inapplicable.sum()),
        "eligible_unanswered_rows": int((applicable & ~answered).sum()),
        "structurally_skipped_rows": int((inapplicable & ~answered).sum()),
        "eligibility_unknown_rows": int(eligibility_unknown.sum()),
        "out_of_route_answered_rows": int((inapplicable & answered).sum()),
    }


def _respondent_groups(
    responses: pd.DataFrame, institution_column: str, country_column: str, source_row_column: str
) -> tuple[set[int], list[dict[str, Any]]]:
    """Find repeated normalized country/institution submissions without deduplication.

    Args:
        responses: Included survey response rows.
        institution_column: Source column identifying the reported institution.
        country_column: Source column identifying the reported country.
        source_row_column: Synthetic original-workbook row number column.

    Returns:
        Source-row set flagged as repeated and deterministic repeat-group diagnostics."""
    groups: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for _, row in responses.iterrows():
        key = (
            _normalised_respondent_value(row[country_column]),
            _normalised_respondent_value(row[institution_column]),
        )
        groups.setdefault(key, []).append(
            {
                "country": _display_value(row[country_column]),
                "institution": _display_value(row[institution_column]),
                "source_row": int(row[source_row_column]),
            }
        )
    duplicates = [group for group in groups.values() if len(group) > 1]
    duplicate_rows = {item["source_row"] for group in duplicates for item in group}
    return duplicate_rows, [
        {"responses": sorted(group, key=lambda item: item["source_row"])}
        for group in sorted(duplicates, key=lambda group: min(item["source_row"] for item in group))
    ]


def _structured_answers(
    question: QuestionDefinition, value: Any
) -> tuple[list[str], list[str]]:
    """Parse and canonicalize one structured response value.

    Args:
        question: Validated structured question defining delimiter and category aliases.
        value: Raw response cell value for that question.

    Returns:
        Deduplicated canonical selections and duplicate selections, in source order."""
    if _is_blank(value):
        return [], []
    raw_answers = (
        [part.strip() for part in str(value).split(question.delimiter)]
        if question.question_type == "multi_choice"
        else [str(value)]
    )
    answers: list[str] = []
    duplicates: list[str] = []
    for answer in raw_answers:
        canonical_answer = question.category_aliases.get(answer, answer)
        if canonical_answer in answers:
            duplicates.append(canonical_answer)
        else:
            answers.append(canonical_answer)
    return answers, duplicates


def _association_axis_value(
    question: QuestionDefinition,
    response: Mapping[str, Any],
    questions_by_id: Mapping[str, QuestionDefinition],
) -> tuple[str, str | None]:
    """Classify one association axis without retaining the submitted raw response.

    Args:
        question: Validated structured question used by this axis.
        response: One source response row.
        questions_by_id: Validated questions indexed by stable question ID.

    Returns:
        One of the four association row states and a declared category only for
        an answered state.

    Raises:
        InputError: If a required response or applicability dependency is absent.
    """
    if question.column not in response:
        raise InputError(f"Association response lacks question column {question.column!r}.")
    raw_value = response[question.column]
    if question.applicability is not None:
        parent_column = question.applicability.get("column")
        values = question.applicability.get("values")
        if not isinstance(parent_column, str) or not isinstance(values, list) or not all(
            isinstance(value, str) for value in values
        ):
            raise InputError(
                f"Schema applicability for {question.column!r} must declare column and string values."
            )
        parent_question = next(
            (item for item in questions_by_id.values() if item.column == parent_column), None
        )
        if parent_question is None or parent_column not in response:
            raise InputError(f"Association applicability column is absent: {parent_column!r}.")
        parent_value = response[parent_column]
        parent_answers, _ = _structured_answers(parent_question, parent_value)
        if _is_blank(parent_value):
            return "missing", None
        if not frozenset(values).intersection(parent_answers):
            return ("inapplicable", None) if _is_blank(raw_value) else ("out_of_route", None)
    answers, _ = _structured_answers(question, raw_value)
    if len(answers) == 1 and answers[0] in question.categories:
        return "answered", answers[0]
    return "missing", None


def _association_definition_payload(definition: AssociationHeatmapDefinition) -> dict[str, Any]:
    """Serialize one validated association definition for a renderer payload.

    Args:
        definition: Validated association definition.

    Returns:
        JSON-serializable metadata preserving exact registry category order.
    """
    return {
        "definition_id": definition.definition_id,
        "row_question_id": definition.row_question_id,
        "column_question_id": definition.column_question_id,
        "row_form_uid": definition.row_form_uid,
        "column_form_uid": definition.column_form_uid,
        "row_categories": list(definition.row_categories),
        "column_categories": list(definition.column_categories),
        "mode": definition.mode,
        "title": definition.title,
        "interpretation": definition.interpretation,
        "conditional_summary": definition.conditional_summary,
    }


def _association_definitions_sha256(
    definitions: Sequence[AssociationHeatmapDefinition],
) -> str:
    """Return the deterministic provenance digest for serialized definitions.

    Args:
        definitions: Validated association definitions in rendering order.

    Returns:
        Lowercase SHA-256 over canonical JSON definition metadata.
    """
    serialized = [_association_definition_payload(definition) for definition in definitions]
    canonical_json = json.dumps(
        serialized, ensure_ascii=False, separators=(",", ":"), sort_keys=True
    )
    return sha256(canonical_json.encode("utf-8")).hexdigest()


def build_association_heatmap_payload(
    definitions: Sequence[AssociationHeatmapDefinition],
    responses: pd.DataFrame,
    questions_by_id: Mapping[str, QuestionDefinition],
    source_row_column: str,
) -> dict[str, dict[str, Any]]:
    """Build privacy-safe paired vectors and contingency cells for approved associations.

    Args:
        definitions: Schema-validated association definitions in report order.
        responses: Included survey response rows.
        questions_by_id: Validated questions indexed by stable question ID.
        source_row_column: Column holding original workbook row numbers.

    Returns:
        Association data keyed by definition ID, containing category metadata,
        source-row-only vectors, complete-case cells, and state diagnostics.

    Raises:
        InputError: If definitions, response columns, or source-row values are invalid.
    """
    if source_row_column not in responses.columns:
        raise InputError(f"Association responses lack source-row column {source_row_column!r}.")
    validate_association_definitions(definitions, questions_by_id)
    pairs: dict[str, dict[str, Any]] = {}
    for definition in definitions:
        row_question = questions_by_id[definition.row_question_id]
        column_question = questions_by_id[definition.column_question_id]
        cells = {
            f"{row_category}|{column_category}": 0
            for row_category in definition.row_categories
            for column_category in definition.column_categories
        }
        vectors: list[dict[str, Any]] = []
        diagnostics = {
            "eligible_missing_rows": [],
            "inapplicable_rows": [],
            "out_of_route_rows": [],
        }
        source_rows: set[int] = set()
        for _, response in responses.iterrows():
            source_row = _payload_count(response[source_row_column], source_row_column)
            if source_row in source_rows:
                raise InputError(f"Association responses have duplicate source row {source_row}.")
            source_rows.add(source_row)
            row_state, row_value = _association_axis_value(
                row_question, response, questions_by_id
            )
            column_state, column_value = _association_axis_value(
                column_question, response, questions_by_id
            )
            vector = {
                "source_row": source_row,
                "row_state": row_state,
                "row_value": row_value,
                "column_state": column_state,
                "column_value": column_value,
            }
            vectors.append(vector)
            states = {row_state, column_state}
            if "missing" in states:
                diagnostics["eligible_missing_rows"].append(source_row)
            if "inapplicable" in states:
                diagnostics["inapplicable_rows"].append(source_row)
            if "out_of_route" in states:
                diagnostics["out_of_route_rows"].append(source_row)
            if row_state == column_state == "answered":
                cells[f"{row_value}|{column_value}"] += 1
        pairs[definition.definition_id] = {
            "row_categories": list(definition.row_categories),
            "column_categories": list(definition.column_categories),
            "vectors": vectors,
            "cells": cells,
            "paired_denominator": sum(cells.values()),
            "diagnostics": diagnostics,
        }
    return pairs


def _structured_question_payload(
    question: QuestionDefinition,
    responses: pd.DataFrame,
    institution_column: str,
    country_column: str,
    source_row_column: str,
    duplicate_rows: set[int],
) -> dict[str, Any]:
    """Build counts, evidence, and diagnostics for one structured question.

    Args:
        question: Validated non-free-text question definition.
        responses: Included survey response rows.
        institution_column: Source column identifying the reported institution.
        country_column: Source column identifying the reported country.
        source_row_column: Synthetic original-workbook row number column.
        duplicate_rows: Source rows belonging to repeated respondent groups.

    Returns:
        JSON-serializable population counts, categories, literal evidence, and diagnostics."""
    counts = {category: 0 for category in question.categories}
    declared_categories = frozenset(question.categories)
    contributions: list[dict[str, Any]] = []
    duplicate_selections: list[dict[str, Any]] = []
    unexpected_selections: list[dict[str, Any]] = []
    contradictory_selections: list[dict[str, Any]] = []
    answered_rows = 0
    for _, row in responses.iterrows():
        source_row = int(row[source_row_column])
        answers, duplicates = _structured_answers(question, row[question.column])
        if answers:
            answered_rows += 1
        if duplicates:
            duplicate_selections.append({"source_row": source_row, "values": duplicates})
        unexpected = [answer for answer in answers if answer not in declared_categories]
        if unexpected:
            unexpected_selections.append({"source_row": source_row, "values": unexpected})
            for answer in unexpected:
                counts.setdefault(answer, 0)
        exclusive_values = [
            answer for answer in answers if answer in question.exclusive_categories
        ]
        if exclusive_values and len(answers) > 1:
            contradictory_selections.append(
                {
                    "source_row": source_row,
                    "exclusive_values": exclusive_values,
                    "values": answers,
                }
            )
        selected = answers
        if not selected:
            selected = ["Missing"]
        for answer in selected:
            if answer != "Missing":
                counts[answer] += 1
            contributions.append(
                {
                    "value": answer,
                    "country": _display_value(row[country_column]),
                    "institution": _display_value(row[institution_column]),
                    "source_row": source_row,
                    "repeated_response": source_row in duplicate_rows,
                }
            )
    duplicate_selections.sort(key=lambda item: item["source_row"])
    unexpected_selections.sort(key=lambda item: item["source_row"])
    population = {"N": len(responses), "A": answered_rows, "M": len(responses) - answered_rows}
    population["blank_rows"] = population["M"]
    category_rows = [
        {
            "value": category,
            "count": counts[category],
            "percent_base": population["A"],
            "percent": _percentage(counts[category], population["A"]),
        }
        for category in counts
    ]
    category_rows.append(
        {
            "value": "Missing",
            "count": population["M"],
            "percent_base": population["N"],
            "percent": _percentage(population["M"], population["N"]),
        }
    )
    category_order = {value: index for index, value in enumerate((*counts, "Missing"))}
    contributions.sort(
        key=lambda item: (
            category_order[item["value"]], item["country"], item["institution"], item["source_row"]
        )
    )
    return {
        "population": population,
        "categories": category_rows,
        "pie_categories": [
            {**category, "excluded_from_chart": category["value"] == "Missing"}
            for category in category_rows
        ],
        "contributions": contributions,
        "free_text_rows": [],
        "diagnostics": {
            "duplicate_selections": duplicate_selections,
            "unexpected_selections": unexpected_selections,
            "contradictory_selections": contradictory_selections,
        },
    }


def _canonical_upset_vectors(
    question: QuestionDefinition,
    responses: pd.DataFrame,
    institution_column: str,
    country_column: str,
    source_row_column: str,
    question_by_column: Mapping[str, QuestionDefinition],
) -> list[dict[str, Any]]:
    """Build source-of-truth respondent vectors for one multi-choice question.

    Args:
        question: Validated multi-choice question to serialize.
        responses: Included survey response rows in source-row order.
        institution_column: Source column identifying the reported institution.
        country_column: Source column identifying the reported country.
        source_row_column: Synthetic original-workbook row-number column.
        question_by_column: Validated questions indexed by source column.

    Returns:
        One vector per included response with applicability state and declared selections.

    Raises:
        InputError: If a declared applicability rule is malformed or refers to an unknown question.
    """
    if question.question_type != "multi_choice":
        raise InputError("Canonical UpSet vectors may only be built for multi_choice questions.")
    applicable_values: frozenset[str] | None = None
    parent_question: QuestionDefinition | None = None
    applicability_column: str | None = None
    if question.applicability is not None:
        applicability_column = question.applicability.get("column")
        values = question.applicability.get("values")
        if not isinstance(applicability_column, str) or not isinstance(values, list) or not all(isinstance(value, str) for value in values):
            raise InputError(f"Schema applicability for {question.column!r} must declare column and string values.")
        parent_question = question_by_column.get(applicability_column)
        if parent_question is None:
            raise InputError(f"Schema applicability column is absent: {applicability_column!r}.")
        applicable_values = frozenset(values)
    vectors: list[dict[str, Any]] = []
    for _, row in responses.iterrows():
        state = "answered"
        if applicability_column is not None and parent_question is not None and applicable_values is not None:
            parent_value = row[applicability_column]
            if _is_blank(parent_value) or not applicable_values.intersection(_structured_answers(parent_question, parent_value)[0]):
                state = "inapplicable"
        answers, _ = _structured_answers(question, row[question.column])
        selected = [category for category in question.categories if category in answers]
        if state == "answered" and not selected:
            state = "missing"
        vectors.append({
            "source_row": int(row[source_row_column]),
            "country": _display_value(row[country_column]),
            "institution": _display_value(row[institution_column]),
            "state": state,
            "selected_categories": selected if state == "answered" else [],
        })
    return vectors


def _free_text_question_payload(
    question: QuestionDefinition,
    responses: pd.DataFrame,
    institution_column: str,
    country_column: str,
    source_row_column: str,
    question_by_column: Mapping[str, QuestionDefinition],
) -> dict[str, Any]:
    """Build literal narrative evidence and parent-answer context for one free-text question.

    Args:
        question: Validated free-text question definition and its configured parent columns.
        responses: Included, nonblank survey response rows.
        institution_column: Response column identifying the responding institution.
        country_column: Response column containing the submitted country value.
        source_row_column: Synthetic workbook-row provenance column.
        question_by_column: Complete validated registry, used to diagnose parent values.

    Returns:
        JSON-serializable free-text rows retaining literal narrative, filtered
        display context, raw parent evidence, and parent-context diagnostics.
    """
    rows = []
    missing_parent_answers: list[dict[str, Any]] = []
    unmatched_parent_context: list[dict[str, Any]] = []
    unexpected_parent_answers: list[dict[str, Any]] = []
    for _, response in responses.iterrows():
        if _is_blank(response[question.column]) or _is_empty_free_text(response[question.column]):
            continue
        source_row = int(response[source_row_column])
        parent_answers = []
        raw_parent_answers = []
        for column in question.parent_columns:
            raw_parent_answers.append({"column": column, "value": _display_value(response[column])})
            parent_question = question_by_column[column]
            answers, _ = _structured_answers(parent_question, response[column])
            context_values = question.parent_context_values.get(column)
            matching_answers = (
                [answer for answer in answers if answer in context_values]
                if context_values is not None else answers
            )
            parent_answers.append({
                "column": column,
                "value": ";".join(matching_answers) if matching_answers
                else ("Missing" if _is_blank(response[column]) else "Not selected"),
            })
            if context_values is not None and not matching_answers and not _is_blank(response[column]):
                unmatched_parent_context.append({"source_row": source_row, "columns": [column]})
        missing_columns = [answer["column"] for answer in parent_answers if answer["value"] == "Missing"]
        if missing_columns:
            missing_parent_answers.append({"source_row": source_row, "columns": missing_columns})
        unexpected_values = [
            answer
            for answer in parent_answers
            if answer["value"] != "Missing"
            and any(
                value not in question_by_column[answer["column"]].categories
                for value in _structured_answers(question_by_column[answer["column"]], response[answer["column"]])[0]
            )
        ]
        if unexpected_values:
            unexpected_parent_answers.append({"source_row": source_row, "values": unexpected_values})
        rows.append(
            {
                "country": _display_value(response[country_column]),
                "institution": _display_value(response[institution_column]),
                "source_row": source_row,
                "text": str(response[question.column]),
                "parent_answers": parent_answers,
                "raw_parent_answers": raw_parent_answers,
            }
        )
    rows.sort(key=lambda item: (item["country"], item["institution"], item["source_row"]))
    population = {"N": len(responses), "A": len(rows), "M": len(responses) - len(rows)}
    population["blank_rows"] = population["M"]
    return {
        "population": population,
        "categories": [],
        "pie_categories": [],
        "contributions": [],
        "free_text_rows": rows,
        "diagnostics": {
            "duplicate_selections": [],
            "unexpected_selections": [],
            "contradictory_selections": [],
            "missing_parent_answers": missing_parent_answers,
            "unmatched_parent_context": unmatched_parent_context,
            "unexpected_parent_answers": unexpected_parent_answers,
        },
    }


def build_descriptive_payload(
    workbook: SurveyWorkbook,
    schema: Mapping[str, Any],
    form_structure: Mapping[str, Any] | None = None,
    association_definitions: Sequence[AssociationHeatmapDefinition] = (),
    association_registry_sha256: str | None = None,
) -> dict[str, Any]:
    """Build a self-contained descriptive payload without Directory data.

    Args:
        workbook: Validated source workbook and nonblank survey rows.
        schema: Validated complete descriptive registry.
        form_structure: Optional authoritative XML-derived question metadata.
        association_definitions: Optional schema-validated association definitions.
        association_registry_sha256: SHA-256 of the loaded association registry.

    Returns:
        JSON-serializable provenance, diagnostics, question counts, literal
        contribution rows, and free-text evidence.

    Raises:
        InputError: If the schema, association registry provenance, or explicit
            applicability rule is invalid.
    """
    root = _required_mapping(schema, "schema")
    output = _required_mapping(root["output"], "output")
    source_row_column = _required_text(output["source_row_column"], "output.source_row_column")
    headers = [column for column in workbook.responses.columns if column != source_row_column]
    questions = validate_descriptive_schema(schema, headers)
    if source_row_column not in workbook.responses.columns:
        raise InputError(f"Workbook responses lack source-row column {source_row_column!r}.")
    columns = _required_mapping(root["columns"], "columns")
    institution_column = _required_text(columns["institution_column"], "columns.institution_column")
    country_column = _required_text(columns["country_column"], "columns.country_column")
    duplicate_rows, duplicate_groups = _respondent_groups(
        workbook.responses, institution_column, country_column, source_row_column
    )
    question_by_column = {question.column: question for question in questions}
    questions_by_id = {question.question_id: question for question in questions}
    if association_definitions and (
        not isinstance(association_registry_sha256, str)
        or re.fullmatch(r"[0-9a-f]{64}", association_registry_sha256) is None
    ):
        raise InputError("Association registry SHA-256 must be a lowercase hexadecimal digest.")
    form_questions = form_structure.get("questions", {}) if form_structure is not None else {}
    expected_form_types = {
        "single_choice": frozenset({"Single Choice", "Single Choice Matrix Question"}),
        "multi_choice": frozenset({"Multiple Choice"}),
        "ordinal": frozenset({"Single Choice", "Single Choice Matrix Question"}),
        "free_text": frozenset({"Free Text"}),
    }
    question_payloads = []
    upset_vectors: dict[str, list[dict[str, Any]]] = {}
    for question in questions:
        form_question = form_questions.get(_form_label(question.column))
        if form_structure is not None and form_question is None:
            raise InputError(f"Descriptive schema question is absent from form XML: {question.column!r}.")
        if form_question is not None and form_question["response_type"] not in expected_form_types[question.question_type]:
            allowed_types = ", ".join(sorted(expected_form_types[question.question_type]))
            raise InputError(
                f"Descriptive schema type {question.question_type!r} conflicts with XML response type "
                f"{form_question['response_type']!r} for {question.column!r}; expected {allowed_types}."
            )
        detail = (
            _free_text_question_payload(
                question,
                workbook.responses,
                institution_column,
                country_column,
                source_row_column,
                question_by_column,
            )
            if question.question_type == "free_text"
            else _structured_question_payload(
                question, workbook.responses, institution_column, country_column,
                source_row_column, duplicate_rows
            )
        )
        question_payloads.append(
            {
                "question_id": question.question_id,
                "column": question.column,
                "label": question.label,
                "question_type": question.question_type,
                "form": form_question,
                "applicability": _applicability_summary(
                    question, workbook.responses, question_by_column
                ),
                **detail,
            }
        )
        if question.question_type == "multi_choice":
            upset_vectors[question.question_id] = _canonical_upset_vectors(
                question, workbook.responses, institution_column, country_column,
                source_row_column, question_by_column,
            )
    payload = {
        "payload_type": "so2_descriptive_statistics",
        "payload_version": "1",
        "provenance": {
            "source_path": workbook.source_path,
            "source_sha256": workbook.source_sha256,
            "worksheet": workbook.worksheet,
            "alias": workbook.alias,
            "export_date": workbook.export_date,
            "header_row": workbook.header_row,
            "total_data_rows": workbook.total_data_rows,
            "excluded_blank_rows": workbook.excluded_blank_rows,
            "included_response_rows": len(workbook.responses),
            "schema_version": root["schema_version"],
        },
        "diagnostics": {"duplicate_respondent_groups": duplicate_groups},
        "form_structure": None if form_structure is None else {
            "source_path": form_structure["source_path"],
            "source_sha256": form_structure["source_sha256"],
        },
        "upset_vectors": upset_vectors,
        "questions": question_payloads,
    }
    if association_definitions:
        payload["association_heatmap_definitions"] = {
            "registry_sha256": association_registry_sha256,
            "definitions_sha256": _association_definitions_sha256(
                association_definitions
            ),
            "definitions": [
                _association_definition_payload(definition)
                for definition in association_definitions
            ],
        }
        payload["association_heatmaps"] = build_association_heatmap_payload(
            association_definitions, workbook.responses, questions_by_id, source_row_column
        )
    return payload


@dataclass(frozen=True)
class RenderedDescriptiveReport:
    """Rendered report TeX and shared standalone-chart fragments."""

    tex: str
    chart_fragments: Mapping[str, str]
    chart_documents: Mapping[str, str]
    chart_paths: Mapping[str, str]


_TEX_ESCAPE = {
    "\\": r"\textbackslash{}",
    "&": r"\&",
    "%": r"\%",
    "$": r"\$",
    "#": r"\#",
    "_": r"\_",
    "{": r"\{",
    "}": r"\}",
    "~": r"\textasciitilde{}",
    "^": r"\textasciicircum{}",
}


def _tex(value: Any) -> str:
    """Escape one literal value for ordinary TeX text.

    Args:
        value: Literal value that must not be interpreted as TeX markup.

    Returns:
        TeX-escaped string safe for ordinary text contexts."""
    return "".join(_TEX_ESCAPE.get(character, character) for character in str(value))


_URL_PATTERN = re.compile(r"https?://[^\s<>{}\[\]]+")
_EMPTY_FREE_TEXT = frozenset({"", "-", "n/a", "na", "nil", "none", "no", "no comment", "nothing", "not applicable"})


def _is_empty_free_text(value: Any) -> bool:
    """Return whether a nominal free-text response conveys no narrative information.

    Args:
        value: Literal cell value from a free-text survey question.

    Returns:
        ``True`` when the normalized value is a known empty-response placeholder.
    """
    return str(value).strip().casefold().rstrip(".") in _EMPTY_FREE_TEXT


def _tex_with_links(value: Any) -> str:
    """Escape narrative text while replacing literal HTTP(S) URLs with short hyperlinks.

    Args:
        value: Literal free-text response which may include HTTP(S) URLs.

    Returns:
        TeX-safe text, with each detected URL represented by a ``link`` hyperlink.
    """
    text = str(value)
    parts: list[str] = []
    cursor = 0
    for match in _URL_PATTERN.finditer(text):
        parts.append(_tex(text[cursor:match.start()]))
        url = match.group(0)
        parts.append(r"\href{\detokenize{" + url + r"}}{link}")
        cursor = match.end()
    parts.append(_tex(text[cursor:]))
    return "".join(parts)


def _question_identifier(value: Any) -> str:
    """Render a compact monospace identifier with legal line breaks at underscores.

    Args:
        value: Question identifier from the descriptive-report schema.

    Returns:
        TeX markup for a smaller monospace identifier which may wrap at underscores.
    """
    return r"\texttt{" + _tex(value).replace(r"\_", r"\_\hspace{0pt}") + r"}"


def _slug(value: str) -> str:
    """Create an ASCII filename stem from presentation metadata.

    Args:
        value: Human-readable label or identifier to convert.

    Returns:
        Lowercase hyphenated ASCII filename stem, possibly empty."""
    ascii_value = normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"-+", "-", re.sub(r"[^a-z0-9]+", "-", ascii_value.lower())).strip("-")


def _chart_key(question: Mapping[str, Any], answered_only: bool = False) -> str:
    """Build a stable bounded chart filename key.

    Args:
        question: Renderable question payload containing identifier and label.
        answered_only: Whether this is the Missing-excluded pie variant.

    Returns:
        Stable filename stem with question ordinal and optional variant suffix.

    Raises:
        InputError: If the question identifier or label cannot produce a safe key."""
    match = re.search(r"(?:^|_)q?0*(\d{1,3})(?:_|$)", str(question["question_id"]))
    if match is None:
        raise InputError(f"Question id lacks a numeric chart prefix: {question['question_id']!r}.")
    number = int(match.group(1))
    compact_names = {8: "experience", 9: "institution-type"}
    stem = compact_names.get(number, _slug(str(question["label"])))
    if not stem:
        raise InputError(f"Question label cannot form a chart filename: {question['label']!r}.")
    suffix = "-answered-only" if answered_only else ""
    max_stem_length = 120 - len(f"{number:02d}-{suffix}.tex")
    stem = stem[:max_stem_length].rstrip("-")
    return f"{number:02d}-{stem}{suffix}"


def _category_label(category: Mapping[str, Any]) -> str:
    """Format one category count with its denominator-aware percentage.

    Args:
        category: Renderable category with value, count, and percentage metadata.

    Returns:
        TeX-safe descriptive label for legends or textual category summaries."""
    value = _tex(category["value"])
    count = int(category["count"])
    percentage = category.get("percent")
    if percentage is None:
        return f"{value}: {count} (no percentage base)"
    basis = "all included rows" if category["value"] == "Missing" else "answering rows"
    return f"{value}: {count} ({float(percentage):.1f}\\% of {basis})"


def _bar_end_label(category: Mapping[str, Any]) -> str:
    """Return a short denominator-aware label for one bar endpoint.

    Args:
        category: Renderable category with count, percentage, and value semantics.

    Returns:
        TeX-safe count/percentage text using the report abbreviation glossary.
    """
    count = int(category["count"])
    percentage = category.get("percent")
    if percentage is None:
        return f"{count} (no percentage base)"
    basis = "oIR" if category["value"] == "Missing" else "oAR"
    return f"{count} ({float(percentage):.1f}\\% {basis})"


_ORDINAL_EXCEPTIONS = frozenset({
    "missing", "not applicable", "n/a", "don't know", "i don't know", "i don’t know",
})

# The default 11pt article text block is 541.4pt (about 19.03cm). The axis
# height formula adds 1cm to 0.52cm per coordinate, so 34.6 uses 19.0cm.
_BAR_PAGE_PRINTABLE_HEIGHT_CM = 19.0
_BAR_AXIS_BASE_HEIGHT_CM = 1.0
_BAR_AXIS_COORDINATE_HEIGHT_CM = 0.52
_BAR_PAGE_MAX_COORDINATE_HEIGHT = (
    _BAR_PAGE_PRINTABLE_HEIGHT_CM - _BAR_AXIS_BASE_HEIGHT_CM
) / _BAR_AXIS_COORDINATE_HEIGHT_CM


def _bar_row_height(category: Mapping[str, Any]) -> float:
    """Return vertical axis space needed by a category label.

    Args:
        category: Renderable descriptive category with a literal ``value``.

    Returns:
        Axis-coordinate height that prevents neighbouring wrapped labels touching.
    """
    return max(1, ceil(len(str(category["value"])) / 26)) + 0.35


def _bar_chunks(
    categories: Sequence[Mapping[str, Any]],
    max_height: float = _BAR_PAGE_MAX_COORDINATE_HEIGHT,
) -> list[list[Mapping[str, Any]]]:
    """Partition ordered categories into page-sized chart groups.

    Args:
        categories: Categories in the semantic display order to preserve.
        max_height: Maximum sum of adaptive row heights in one chart fragment.
            The default consumes the complete printable height of the default
            report page, excluding page margins, headers, and footers.

    Returns:
        Non-empty ordered groups; each fits the requested height unless one
        category alone requires more space.
    """
    chunks: list[list[Mapping[str, Any]]] = []
    current: list[Mapping[str, Any]] = []
    current_height = 0.0
    for category in categories:
        row_height = _bar_row_height(category)
        if current and current_height + row_height > max_height:
            chunks.append(current)
            current = []
            current_height = 0.0
        current.append(category)
        current_height += row_height
    if current:
        chunks.append(current)
    return chunks


def _bar_axis(
    categories: Sequence[Mapping[str, Any]], *, shared_max_count: int | None = None
) -> str:
    """Render one horizontal count-bar axis with adaptive row spacing.

    Args:
        categories: Non-empty ordered categories to render in this axis.
        shared_max_count: Optional maximum count across sibling chunks, ensuring
            their horizontal axes use an identical scale.

    Returns:
        A complete TikZ/PGFPlots axis fragment, or explanatory TeX when empty.
    """
    if not categories:
        return r"\emph{No categories in this group.}"
    max_count = shared_max_count or max(int(category["count"]) for category in categories)
    rows = []
    labels = []
    ticks = []
    cursor = 0.0
    for category in categories:
        row_height = _bar_row_height(category)
        position = cursor + row_height / 2
        cursor += row_height
        count = int(category["count"])
        normalized = str(category["value"]).strip().casefold()
        color = (
            "bbmriRed" if normalized == "missing"
            else "bbmriGray" if normalized in _ORDINAL_EXCEPTIONS
            else "bbmriBlue"
        )
        rows.append(
            f"\\addplot+[fill={color},draw={color},bar shift=0pt] coordinates "
            f"{{({count},{position:.2f})}};"
        )
        labels.append(
            f"\\node[anchor=west,font=\\scriptsize,align=left] "
            f"at (axis cs:{count + 0.08},{position:.2f}) "
            f"{{{_bar_end_label(category)}}};"
        )
        ticks.append(f"{position:.2f}")
    tick_labels = ",".join(
        rf"{{\raggedleft {_tex(category['value'])}}}"
        for category in categories
    )
    return "\n".join([
        r"\begin{tikzpicture}",
        r"\begin{axis}[xbar, xmin=0, xmax=" + f"{max(1, max_count) * 1.18:.2f}, "
        r"ymin=0, ymax=" + f"{cursor + 0.70:.2f}, width=0.42\\textwidth, xshift=0.46\\textwidth, "
        r"scale only axis, height=" + f"{max(3.0, _BAR_AXIS_COORDINATE_HEIGHT_CM * cursor + _BAR_AXIS_BASE_HEIGHT_CM):.1f}cm,",
        f"ytick={{{','.join(ticks)}}}, yticklabels={{{tick_labels}}},",
        r"xlabel={Count}, y dir=reverse, axis x line*=bottom, axis y line*=left,",
        r"yticklabel style={text width=0.40\textwidth,align=right,font=\scriptsize},",
        r"enlarge x limits={upper,value=0.14}, clip=false]",
        *rows,
        *labels,
        r"\end{axis}",
        r"\end{tikzpicture}",
    ])


def _bar_charts(categories: Sequence[Mapping[str, Any]]) -> str:
    """Render page-sized bar charts with an identical count scale.

    Args:
        categories: Ordered categories for one conceptual chart.

    Returns:
        TeX fragments separated by page breaks when multiple chunks are needed.
    """
    if not categories:
        return _bar_axis(categories)
    shared_max_count = max(int(category["count"]) for category in categories)
    fragments = []
    for index, chunk in enumerate(_bar_chunks(categories)):
        if index:
            fragments.extend([r"\clearpage", r"\noindent\textit{Continued}\par"])
        fragments.append(_bar_axis(chunk, shared_max_count=shared_max_count))
    return "\n".join(fragments)


def _bar_fragment(question: Mapping[str, Any]) -> str:
    """Render a question's categorical bars, separating ordinal exceptions.

    Args:
        question: Validated question payload containing category counts.

    Returns:
        TeX for one or more compatible bar-chart pages.
    """
    categories = question["categories"]
    if question["question_type"] != "ordinal":
        return _bar_charts(categories)
    ordered = [
        category for category in categories
        if str(category["value"]).strip().casefold() not in _ORDINAL_EXCEPTIONS
    ]
    exceptional = [category for category in categories if category not in ordered]
    fragments = [r"\noindent\textit{Ordered scale}\par", _bar_charts(ordered)]
    if any(int(category["count"]) for category in exceptional):
        fragments.extend([r"\noindent\textit{Exceptional categories}\par", _bar_charts(exceptional)])
    return "\n".join(fragments)

_PIE_RADIUS = 1.5
_PIE_LABEL_BOUND = 1.45
_PIE_LABEL_GAP = 0.08


def _pie_label_height(lines: int) -> float:
    """Return the vertical space reserved for a pie label.

    Args:
        lines: Wrapped-line count calculated from the literal category label.

    Returns:
        Height in TikZ coordinates, including a small separation allowance.
    """
    return 0.42 * lines + 0.16


def _pie_label_layout(labels: Sequence[Mapping[str, Any]], label_bound: float = _PIE_LABEL_BOUND) -> dict[int, tuple[str, float]]:
    """Place pie labels near their slice centers without exceeding pie height.

    Args:
        labels: Segment descriptors containing a stable ``index``, middle angle,
            and wrapped-line count.
        label_bound: Maximum absolute vertical label center in TikZ coordinates.

    Returns:
        Mapping from segment index to its selected horizontal side and bounded
        vertical label center. Labels are moved to the opposite side only when
        their preferred side cannot fit within the pie's vertical extent.

    Raises:
        InputError: If label blocks cannot fit within the fixed pie height.
    """
    groups = {"left": [], "right": []}
    for label in labels:
        preferred = "right" if cos(radians(float(label["middle"]))) >= 0 else "left"
        groups[preferred].append({**label, "side": preferred})

    def occupied_height(items: Sequence[Mapping[str, Any]]) -> float:
        return sum(_pie_label_height(int(item["lines"])) for item in items) + max(0, len(items) - 1) * _PIE_LABEL_GAP

    available = 2 * label_bound
    for side in ("left", "right"):
        while occupied_height(groups[side]) > available:
            other = "left" if side == "right" else "right"
            candidates = sorted(
                groups[side],
                key=lambda item: (abs(cos(radians(float(item["middle"])))), -int(item["lines"])),
            )
            candidate = next(
                (
                    item for item in candidates
                    if occupied_height([*groups[other], item]) <= available
                ),
                None,
            )
            if candidate is None:
                raise InputError("Pie-chart labels cannot fit within the fixed pie height.")
            groups[side].remove(candidate)
            candidate["side"] = other
            groups[other].append(candidate)

    layout: dict[int, tuple[str, float]] = {}
    for side, items in groups.items():
        ordered = sorted(items, key=lambda item: _PIE_RADIUS * sin(radians(float(item["middle"]))))
        centers: list[float] = []
        heights = [_pie_label_height(int(item["lines"])) for item in ordered]
        for item, height in zip(ordered, heights, strict=True):
            desired = _PIE_RADIUS * sin(radians(float(item["middle"])))
            minimum = -label_bound + height / 2
            maximum = label_bound - height / 2
            center = min(max(desired, minimum), maximum)
            if centers:
                center = max(center, centers[-1] + heights[len(centers) - 1] / 2 + _PIE_LABEL_GAP + height / 2)
            centers.append(center)
        # Clamp the upper label first, then preserve the required gap backward.
        for index in range(len(centers) - 1, -1, -1):
            maximum = _PIE_LABEL_BOUND - heights[index] / 2
            if index < len(centers) - 1:
                maximum = min(maximum, centers[index + 1] - heights[index + 1] / 2 - _PIE_LABEL_GAP - heights[index] / 2)
            centers[index] = min(centers[index], maximum)
        if centers and centers[0] - heights[0] / 2 < -_PIE_LABEL_BOUND:
            raise InputError("Pie-chart labels cannot fit within the fixed pie height.")
        for item, center in zip(ordered, centers, strict=True):
            layout[int(item["index"])] = (side, center)
    return layout


def _pie_fragment(question: Mapping[str, Any], answered_only: bool, max_ratio: float = 4 / 3) -> str:
    """Render a pie whose leader lines connect each segment to its label.

    Args:
        question: Validated categorical question payload with pie categories.
        answered_only: Whether the Missing category is excluded from this chart.
        max_ratio: Maximum permitted horizontal-to-vertical chart ratio.

    Returns:
        Complete TikZ markup with labels near slice centers and within the pie's
        fixed vertical extent.
    """
    categories = [
        category for category in question["pie_categories"]
        if int(category["count"]) > 0
        and not (answered_only and category.get("excluded_from_chart"))
    ]
    total = sum(int(category["count"]) for category in categories)
    title = "answered rows only" if answered_only else "including Missing"
    if total == 0:
        return "\n".join([
            r"\begin{tikzpicture}",
            rf"\node {{No responses ({title})}};",
            r"\end{tikzpicture}",
        ])
    start = 0.0
    colors = (
        "bbmriBlue", "bbmriTeal", "bbmriGold", "bbmriGray",
        "bbmriGreen", "bbmriOrange",
    )
    segments = []
    non_missing_index = 0
    for index, category in enumerate(categories):
        count = int(category["count"])
        end = start + 360 * count / total
        middle = (start + end) / 2
        if str(category["value"]).strip().casefold() == "missing":
            color = "bbmriRed"
        else:
            color = colors[non_missing_index % len(colors)]
            non_missing_index += 1
        segments.append({
            "index": index, "category": category, "start": start, "end": end,
            "middle": middle, "color": color,
            "label": f"{_tex(category['value'])}: {count} ({100 * count / total:.1f}\\%)",
            "lines": max(1, ceil(len(str(category["value"])) / 42)),
        })
        start = end
    keyed_legend = False
    max_label_bound = max(_PIE_LABEL_BOUND, min(5.0, 2.8 * (4 / 3) / max_ratio))
    for label_bound in (_PIE_LABEL_BOUND, min(2.1, max_label_bound), max_label_bound):
        try:
            layout = _pie_label_layout(segments, label_bound)
            break
        except InputError:
            continue
    else:
        keyed_legend = True
        for segment in segments:
            segment["label"] = rf"\#{segment['index'] + 1}"
            segment["lines"] = 1
        layout = _pie_label_layout(segments)
    slices = [r"\begin{tikzpicture}", rf"\node[above] at (0,1.9) {{{title}}};"]
    for segment in segments:
        side, y = layout[segment["index"]]
        slices.append(
            rf"\path[fill={segment['color']},draw=white] (0,0) -- ({segment['start']:.3f}:1.5)"
            rf" arc ({segment['start']:.3f}:{segment['end']:.3f}:1.5) -- cycle;"
        )
        swatch_x = 1.76 if side == "right" else -1.76
        # Intersect the centre-to-swatch ray with the pie arc. This keeps the
        # leader line visually radial even after label collision adjustments.
        contact_angle = degrees(atan2(y, swatch_x))
        if side == "right":
            slices.extend([
                rf"\draw[{segment['color']},dashed,thin] ({contact_angle:.3f}:1.5) -- ({swatch_x:.2f},{y:.2f});",
                rf"\fill[{segment['color']}] (1.76,{y - 0.06:.2f}) rectangle (1.90,{y + 0.06:.2f});",
                rf"\node[anchor=west,align=left,text width=0.40\linewidth,font=\scriptsize] at (1.98,{y:.2f}) {{{segment['label']}}};",
            ])
        else:
            slices.extend([
                rf"\draw[{segment['color']},dashed,thin] ({contact_angle:.3f}:1.5) -- ({swatch_x:.2f},{y:.2f});",
                rf"\fill[{segment['color']}] (-1.90,{y - 0.06:.2f}) rectangle (-1.76,{y + 0.06:.2f});",
                rf"\node[anchor=east,align=right,text width=0.40\linewidth,font=\scriptsize] at (-1.98,{y:.2f}) {{{segment['label']}}};",
            ])
    slices.append(r"\end{tikzpicture}")
    if keyed_legend:
        slices.append(r"\noindent\textit{Legend}\par")
        slices.extend(rf"\noindent\#{segment['index'] + 1} -- {_tex(segment['category']['value'])}\par" for segment in segments)
    return "\n".join(slices)

def _table(rows: Sequence[Sequence[str]], columns: str) -> str:
    """Render a repeated-header longtable from already escaped cells.

    Args:
        rows: Header followed by body rows, with TeX-safe cell values.
        columns: TeX longtable column specification.

    Returns:
        Complete longtable TeX, or an empty string when no rows are supplied."""
    if not rows:
        return ""
    header, *body = rows
    return "\n".join([
        rf"\begin{{longtable}}{{{columns}}}",
        r"\toprule",
        " & ".join(header) + r" \\",
        r"\midrule",
        r"\endfirsthead",
        r"\toprule",
        " & ".join(header) + r" \\",
        r"\midrule",
        r"\endhead",
        *[" & ".join(row) + r" \\" for row in body],
        r"\bottomrule",
        r"\end{longtable}",
    ])


_COUNTRY_CODES = {
    "austria": "AT", "belgium": "BE", "bulgaria": "BG", "czech republic": "CZ",
    "cz": "CZ", "denmark": "DK", "estonia": "EE", "finland": "FI", "france": "FR",
    "germany": "DE", "hungary": "HU", "italia": "IT", "italy": "IT", "latvia": "LV",
    "lithuania": "LT", "malta": "MT", "netherlands": "NL", "the netherlands": "NL",
    "norway": "NO", "poland": "PL", "qatar": "QA", "slovakia": "SK", "spain": "ES",
    "sweden": "SE", "switzerland": "CH",
}


def _country_code(value: Any) -> str:
    """Return the ISO 3166-1 alpha-2 code for a reported country or ``??`` if unknown.

    Args:
        value: Literal country value from one submitted survey row.

    Returns:
        A two-letter code after whitespace, case, and accent normalisation.
    """
    normalized = normalize("NFKD", str(value)).encode("ascii", "ignore").decode("ascii")
    return _COUNTRY_CODES.get(normalized.strip().casefold(), "??")


def _question_tables(
    question: Mapping[str, Any], *, include_parent_context: bool = False,
) -> str:
    """Render an optional evidence table for one validated question payload.

    Args:
        question: Validated question payload with either free-text rows or structured contributions.
        include_parent_context: Whether free-text tables include complete raw parent
            answers; the default omits parent context entirely.

    Returns:
        TeX for the table, or an empty string if no free-text evidence exists.
    """
    if question["question_type"] == "free_text":
        evidence = question["free_text_rows"]
        if not evidence:
            return ""
        parent_key = "raw_parent_answers"
        parent_columns = (
            tuple(parent["column"] for parent in evidence[0].get(parent_key, []))
            if include_parent_context else ()
        )
        if include_parent_context and any(
            tuple(parent["column"] for parent in item.get(parent_key, [])) != parent_columns
            for item in evidence
        ):
            raise InputError("Free-text table mixes different parent questions and cannot use one parent-context header.")
        show_parent = bool(parent_columns)
        headers = [r"CC", r"Institution"]
        if show_parent:
            parent_header = r"{\raggedright Parent context\par\smaller[3]" + r"\par ".join(
                _tex(column) for column in parent_columns
            ) + r"\par}"
            headers.append(parent_header)
        headers.append(r"Response")
        rows = [headers]
        for item in evidence:
            parent_values = r"{\raggedright " + r"\par ".join(
                _tex(parent["value"]) for parent in item.get(parent_key, [])
            ) + r"\par}"
            row = [_tex(_country_code(item["country"])), _tex(item["institution"])]
            if show_parent:
                row.append(parent_values)
            row.append(_tex_with_links(item["text"]))
            rows.append(row)
        columns = (
            r"p{0.04\linewidth}p{0.27\linewidth}p{0.61\linewidth}"
            if not show_parent else r"p{0.04\linewidth}p{0.21\linewidth}p{0.20\linewidth}p{0.47\linewidth}"
        )
        return "\n".join([r"\smaller[1]", _table(rows, columns), r"\normalsize"])
    grouped: dict[tuple[str, str], list[str]] = {}
    for item in question["contributions"]:
        entry = _tex(item["institution"])
        if item["repeated_response"]:
            entry += " [suspected repeated response]"
        grouped.setdefault((str(item["value"]), str(item["country"])), []).append(entry)
    rows = [[r"Value", r"Country", r"Institutions"]]
    for (value, country), institutions in grouped.items():
        rows.append([
            _tex(value), _tex(country), r"{\raggedright " + "; ".join(institutions) + r"\par}",
        ])
    return "\n".join([
        r"\smaller[3]",
        _table(rows, r"p{0.22\linewidth}p{0.18\linewidth}p{0.52\linewidth}"),
        r"\normalsize",
    ])


def _preamble() -> str:
    """Build the shared XeLaTeX preamble for reports and standalone charts.

    Returns:
        TeX declarations for fonts, tables, links, TikZ/PGFPlots, and report colors."""
    return r"""\documentclass[11pt]{scrartcl}
\KOMAoptions{parskip=half}
\setlength{\parindent}{0pt}
\usepackage{fontspec}
\usepackage{libertine}
\usepackage{relsize}
\usepackage{longtable}
\usepackage{booktabs}
\usepackage{xcolor}
\usepackage{hyperref}
\usepackage{xurl}
\usepackage{tikz}
\usepackage{pgfplots}
\pgfplotsset{compat=1.18}
\definecolor{bbmriBlue}{HTML}{005A9C}
\definecolor{bbmriTeal}{HTML}{008C95}
\definecolor{bbmriGold}{HTML}{D8A000}
\definecolor{bbmriGray}{HTML}{7A7A7A}
\definecolor{bbmriRed}{HTML}{B63A3A}
\definecolor{bbmriGreen}{HTML}{3F7F4C}
\definecolor{bbmriOrange}{HTML}{C76D1E}
"""


def _response_structure_text(question: Mapping[str, Any]) -> str:
    """Format authoritative XML response metadata for a report question.

    Args:
        question: Renderable payload question with optional XML form metadata.

    Returns:
        Response type and requiredness text without inferred applicability.
    """
    form = question.get("form")
    if not isinstance(form, Mapping):
        return "Response type: unavailable (form structure not supplied); mandatory/optional: unavailable."
    response_type = str(form.get("response_type", "unavailable")).casefold()
    response_type = {
        "single choice matrix question": "single-choice matrix",
        "multiple choice matrix question": "multiple-choice matrix",
    }.get(response_type, response_type)
    requiredness = str(form.get("requiredness", "unavailable")).capitalize()
    dependencies = form.get("dependencies", [])
    result = f"Response type: {response_type}; {requiredness}."
    if dependencies:
        conditions = "; ".join(
            f"{item['parent_question']} = {item['answer']}" for item in dependencies
        )
        result += f" Shown when: {conditions}."
    return result


def _chart_unit(question: Mapping[str, Any]) -> str:
    """Return the respondent-level observation unit represented by one chart mark.

    Args:
        question: Renderable question payload containing its declared response type.

    Returns:
        A human-readable unit. Multi-choice marks count rows selecting each
        value; all other chart marks count submitted response rows.
    """
    if question["question_type"] == "multi_choice":
        return "submitted response rows selecting each value"
    return "submitted response rows"



def _standalone_tex(
    question: Mapping[str, Any], fragment: str, answered_only: bool,
) -> str:
    """Build a self-contained TeX document for one chart variant.

    Args:
        question: Renderable question payload providing title and population counts.
        fragment: Complete TikZ/PGFPlots fragment for the selected chart variant.
        answered_only: Whether the chart excludes Missing responses.

    Returns:
        Standalone XeLaTeX document including interpretation metadata and the chart."""
    population = question["population"]
    variant = " (answered rows only)" if answered_only else ""
    note = (
        f"Missing: {population['M']} of {population['N']} included rows. "
        "Blank means blank; mandatory/optional status is reported from the form definition."
    )
    return "\n".join([
        _preamble(),
        r"\begin{document}",
        rf"\section*{{{_tex(question['label'])}{_tex(variant)}}}",
        rf"\noindent\smaller[3] Question identifier: {_question_identifier(question['question_id'])}\normalsize\par\vspace{{0.35\baselineskip}}",
        rf"Denominator: {population['A']} answering rows; Missing uses "
        rf"{population['N']} included rows\par",
        rf"Unit: {_tex(_chart_unit(question))}\par",
        rf"\emph{{{_tex(note)}}}",
        fragment,
        r"\end{document}",
        "",
    ])


def _association_value(definition: AssociationHeatmapDefinition | Mapping[str, Any], field: str) -> Any:
    """Return one association definition field from validated metadata.

    Args:
        definition: Validated immutable definition or its serialized payload form.
        field: Definition field required by the renderer.

    Returns:
        The requested definition value.

    Raises:
        InputError: If serialized metadata does not provide the requested field.
    """
    if isinstance(definition, AssociationHeatmapDefinition):
        return getattr(definition, field)
    if field not in definition:
        raise InputError(f"Association definition lacks {field!r}.")
    return definition[field]


def _association_heatmap_fragment(
    definition: AssociationHeatmapDefinition | Mapping[str, Any], pair: Mapping[str, Any],
) -> str:
    """Render one zero-inclusive annotated association contingency heatmap.

    Args:
        definition: Validated association definition controlling title and axes.
        pair: Validated aggregate pair payload with declared cells and denominator.

    Returns:
        PGFPlots/TikZ markup for the association heatmap and interpretation note.
    """
    row_categories = tuple(_association_value(definition, "row_categories"))
    column_categories = tuple(_association_value(definition, "column_categories"))
    cells = pair["cells"]
    cell_rows = []
    labels = []
    for row_index, row_category in enumerate(row_categories, start=1):
        for column_index, column_category in enumerate(column_categories, start=1):
            count = int(cells[f"{row_category}|{column_category}"])
            cell_rows.append(f"{column_index} {row_index} {count}")
            labels.append(
                rf"\node at (axis cs:{column_index},{row_index}) "
                rf"{{\scriptsize ({_tex(row_category)}, {_tex(column_category)}): {count}}};"
            )
    x_labels = ",".join(_tex(category) for category in column_categories)
    y_labels = ",".join(_tex(category) for category in row_categories)
    title = _tex(_association_value(definition, "title"))
    interpretation = _tex(_association_value(definition, "interpretation"))
    row_axis = _tex(_association_value(definition, "row_question_id"))
    column_axis = _tex(_association_value(definition, "column_question_id"))
    denominator = int(pair["paired_denominator"])
    return "\n".join([
        rf"\subsection*{{{title}}}",
        rf"\noindent Paired answering rows: {denominator}\par",
        r"\begin{center}",
        r"\begin{tikzpicture}",
        r"\begin{axis}[",
        r"width=0.82\linewidth, height=0.54\textheight,",
        rf"xlabel={{Response to {column_axis}}}, ylabel={{Response to {row_axis}}},",
        rf"xtick={{1,...,{len(column_categories)}}}, ytick={{1,...,{len(row_categories)}}},",
        rf"xticklabels={{{x_labels}}}, yticklabels={{{y_labels}}},",
        r"y dir=reverse, colorbar, colormap={associationSequential}{color(0cm)=(white); color(1cm)=(bbmriTeal)},",
        r"point meta min=0,",
        r"]",
        r"\addplot[matrix plot*, mesh/cols=" + str(len(column_categories)) + r", point meta=explicit] table[row sep=\\,meta index=2] {",
        *[f"{row} \\\\" for row in cell_rows],
        r"};",
        *labels,
        r"\end{axis}",
        r"\end{tikzpicture}",
        r"\end{center}",
        rf"\noindent\emph{{{interpretation} This panel describes submitted response rows, not causal association.}}\par",
    ])


def _association_conditional_summary_fragment(
    definition: AssociationHeatmapDefinition | Mapping[str, Any], pair: Mapping[str, Any],
) -> str:
    """Render row-conditional count and percentage summaries for one association.

    Args:
        definition: Validated association definition controlling category order.
        pair: Validated aggregate pair payload with zero-inclusive contingency cells.

    Returns:
        TeX summary rows using each displayed row's exact denominator.
    """
    if not _association_value(definition, "conditional_summary"):
        return ""
    row_categories = tuple(_association_value(definition, "row_categories"))
    column_categories = tuple(_association_value(definition, "column_categories"))
    cells = pair["cells"]
    rows = [r"\noindent\textit{Conditional policy/workflow response by returned-data experience:}\par"]
    for row_category in row_categories:
        denominator = sum(int(cells[f"{row_category}|{column_category}"]) for column_category in column_categories)
        values = []
        for column_category in column_categories:
            count = int(cells[f"{row_category}|{column_category}"])
            percentage = 0.0 if denominator == 0 else 100 * count / denominator
            values.append(
                rf"({_tex(row_category)}, {_tex(column_category)}): "
                rf"{count} / {denominator} ({percentage:.1f}\%)"
            )
        rows.append(rf"\noindent {'; '.join(values)}\par")
    return "\n".join(rows)


def _association_standalone_tex(
    definition: AssociationHeatmapDefinition, pair: Mapping[str, Any],
) -> str:
    """Build a self-contained association panel chart document.

    Args:
        definition: Validated association definition represented by this chart.
        pair: Validated aggregate pair payload.

    Returns:
        Standalone XeLaTeX document containing the panel and any conditional summary.
    """
    return "\n".join([
        _preamble(),
        r"\begin{document}",
        _association_heatmap_fragment(definition, pair),
        _association_conditional_summary_fragment(definition, pair),
        r"\end{document}",
        "",
    ])


def _payload_mapping(value: Any, path: str) -> Mapping[str, Any]:
    """Validate one object-shaped descriptive payload value.

    Args:
        value: Candidate value from a serialized descriptive payload.
        path: Dot/bracket payload path used in actionable error messages.

    Returns:
        The same value narrowed to a mapping.

    Raises:
        InputError: If the value is not mapping-shaped."""
    if not isinstance(value, Mapping):
        raise InputError(f"Descriptive statistics payload {path} must be an object.")
    return value


def _payload_array(value: Any, path: str) -> list[Any]:
    """Validate one array-shaped descriptive payload value.

    Args:
        value: Candidate value from a serialized descriptive payload.
        path: Dot/bracket payload path used in actionable error messages.

    Returns:
        The same value narrowed to a list.

    Raises:
        InputError: If the value is not an array."""
    if not isinstance(value, list):
        raise InputError(f"Descriptive statistics payload {path} must be an array.")
    return value


def _payload_text(value: Any, path: str) -> str:
    """Validate one nonblank descriptive payload string.

    Args:
        value: Candidate value from a serialized descriptive payload.
        path: Dot/bracket payload path used in actionable error messages.

    Returns:
        The original nonblank string.

    Raises:
        InputError: If the value is not a nonblank string."""
    if not isinstance(value, str) or not value.strip():
        raise InputError(f"Descriptive statistics payload {path} must be a nonblank string.")
    return value


def _payload_count(value: Any, path: str) -> int:
    """Validate one nonnegative integer descriptive payload count.

    Args:
        value: Candidate value from a serialized descriptive payload.
        path: Dot/bracket payload path used in actionable error messages.

    Returns:
        The original nonnegative integer.

    Raises:
        InputError: If the value is boolean, nonintegral, or negative."""
    if isinstance(value, bool) or not isinstance(value, int) or value < 0:
        raise InputError(
            f"Descriptive statistics payload {path} must be a nonnegative integer."
        )
    return value


def _association_definition_from_payload(value: Any, path: str) -> AssociationHeatmapDefinition:
    """Deserialize one association definition embedded in a report payload.

    Args:
        value: Candidate JSON definition metadata.
        path: Payload path used in validation errors.

    Returns:
        Immutable association definition with exact serialized category order.

    Raises:
        InputError: If metadata is malformed.
    """
    definition = _payload_mapping(value, path)
    expected_fields = {
        "definition_id", "row_question_id", "column_question_id", "row_form_uid",
        "column_form_uid", "row_categories", "column_categories", "mode", "title",
        "interpretation", "conditional_summary",
    }
    if set(definition) != expected_fields:
        raise InputError(
            f"Descriptive statistics payload {path} has missing or unknown fields."
        )
    mode = _payload_text(definition.get("mode"), f"{path}.mode")
    if mode not in {"default", "exploratory"}:
        raise InputError(f"Descriptive statistics payload {path}.mode is invalid.")
    conditional_summary = definition.get("conditional_summary")
    if not isinstance(conditional_summary, bool):
        raise InputError(
            f"Descriptive statistics payload {path}.conditional_summary must be boolean."
        )
    def categories(field: str) -> tuple[str, ...]:
        values = _payload_array(definition.get(field), f"{path}.{field}")
        if not values or not all(isinstance(item, str) and item.strip() for item in values):
            raise InputError(
                f"Descriptive statistics payload {path}.{field} must contain nonblank strings."
            )
        if len(values) != len(set(values)):
            raise InputError(f"Descriptive statistics payload {path}.{field} has duplicate categories.")
        return tuple(values)

    return AssociationHeatmapDefinition(
        definition_id=_payload_text(definition.get("definition_id"), f"{path}.definition_id"),
        row_question_id=_payload_text(
            definition.get("row_question_id"), f"{path}.row_question_id"
        ),
        column_question_id=_payload_text(
            definition.get("column_question_id"), f"{path}.column_question_id"
        ),
        row_form_uid=_payload_text(definition.get("row_form_uid"), f"{path}.row_form_uid"),
        column_form_uid=_payload_text(
            definition.get("column_form_uid"), f"{path}.column_form_uid"
        ),
        row_categories=categories("row_categories"),
        column_categories=categories("column_categories"),
        mode=mode,
        title=_payload_text(definition.get("title"), f"{path}.title"),
        interpretation=_payload_text(
            definition.get("interpretation"), f"{path}.interpretation"
        ),
        conditional_summary=conditional_summary,
    )


def validate_association_heatmap_payload(payload: Mapping[str, Any]) -> None:
    """Validate serialized association vectors against build-time provenance.

    Args:
        payload: Complete descriptive report payload containing association keys.

    Returns:
        ``None`` after every vector, cell, diagnostic, and embedded definition is valid.

    Raises:
        InputError: If association data is incomplete, includes identifying data,
            or differs from the approved category and definition contract.
    """
    metadata = _payload_mapping(
        payload.get("association_heatmap_definitions"), "association_heatmap_definitions"
    )
    if set(metadata) != {"registry_sha256", "definitions_sha256", "definitions"}:
        raise InputError(
            "Descriptive statistics payload association_heatmap_definitions has "
            "missing or unknown fields."
        )
    registry_sha256 = _payload_text(
        metadata.get("registry_sha256"), "association_heatmap_definitions.registry_sha256"
    )
    if re.fullmatch(r"[0-9a-f]{64}", registry_sha256) is None:
        raise InputError(
            "Descriptive statistics payload association registry SHA-256 must be lowercase hexadecimal."
        )
    definitions_sha256 = _payload_text(
        metadata.get("definitions_sha256"),
        "association_heatmap_definitions.definitions_sha256",
    )
    if re.fullmatch(r"[0-9a-f]{64}", definitions_sha256) is None:
        raise InputError(
            "Descriptive statistics payload association definition SHA-256 must be "
            "lowercase hexadecimal."
        )
    serialized_definitions = _payload_array(
        metadata.get("definitions"), "association_heatmap_definitions.definitions"
    )
    definitions = tuple(
        _association_definition_from_payload(value, f"association_heatmap_definitions.definitions[{index}]")
        for index, value in enumerate(serialized_definitions)
    )
    if definitions_sha256 != _association_definitions_sha256(definitions):
        raise InputError(
            "Descriptive statistics payload association definition digest does not match "
            "the embedded definitions."
        )
    heatmaps = _payload_mapping(payload.get("association_heatmaps"), "association_heatmaps")
    expected_ids = {definition.definition_id for definition in definitions}
    if set(heatmaps) != expected_ids:
        raise InputError("Descriptive statistics payload association_heatmaps must cover exactly the registry definitions.")
    valid_states = {"answered", "missing", "inapplicable", "out_of_route"}
    for definition in definitions:
        path = f"association_heatmaps.{definition.definition_id}"
        pair = _payload_mapping(heatmaps[definition.definition_id], path)
        if pair.get("row_categories") != list(definition.row_categories):
            raise InputError(f"Descriptive statistics payload {path} has mismatched row category order.")
        if pair.get("column_categories") != list(definition.column_categories):
            raise InputError(f"Descriptive statistics payload {path} has mismatched column category order.")
        vectors = _payload_array(pair.get("vectors"), f"{path}.vectors")
        expected_cells = {
            f"{row_category}|{column_category}"
            for row_category in definition.row_categories
            for column_category in definition.column_categories
        }
        cells = _payload_mapping(pair.get("cells"), f"{path}.cells")
        if set(cells) != expected_cells:
            raise InputError(f"Descriptive statistics payload {path} has undeclared association category.")
        cell_total = sum(_payload_count(value, f"{path}.cells.{key}") for key, value in cells.items())
        if _payload_count(pair.get("paired_denominator"), f"{path}.paired_denominator") != cell_total:
            raise InputError(f"Descriptive statistics payload {path} has mismatched association cell total.")
        source_rows: set[int] = set()
        computed_cells = {key: 0 for key in expected_cells}
        computed_diagnostics = {
            "eligible_missing_rows": [],
            "inapplicable_rows": [],
            "out_of_route_rows": [],
        }
        for index, value in enumerate(vectors):
            vector_path = f"{path}.vectors[{index}]"
            vector = _payload_mapping(value, vector_path)
            if set(vector) != {
                "source_row", "row_state", "row_value", "column_state", "column_value"
            }:
                raise InputError(f"Descriptive statistics payload {vector_path} has identifying or unknown fields.")
            source_row = _payload_count(vector.get("source_row"), f"{vector_path}.source_row")
            if source_row in source_rows:
                raise InputError(f"Descriptive statistics payload {path} has duplicate source row {source_row}.")
            source_rows.add(source_row)
            states = []
            for axis, categories in (("row", definition.row_categories), ("column", definition.column_categories)):
                state = vector.get(f"{axis}_state")
                value = vector.get(f"{axis}_value")
                if state not in valid_states:
                    raise InputError(f"Descriptive statistics payload {vector_path}.{axis}_state is invalid.")
                if state == "answered":
                    if value not in categories:
                        raise InputError(f"Descriptive statistics payload {vector_path} has undeclared association category.")
                elif value is not None:
                    raise InputError(f"Descriptive statistics payload {vector_path}.{axis}_value must be null unless answered.")
                states.append(state)
            if "missing" in states:
                computed_diagnostics["eligible_missing_rows"].append(source_row)
            if "inapplicable" in states:
                computed_diagnostics["inapplicable_rows"].append(source_row)
            if "out_of_route" in states:
                computed_diagnostics["out_of_route_rows"].append(source_row)
            if states == ["answered", "answered"]:
                computed_cells[f"{vector['row_value']}|{vector['column_value']}"] += 1
        if dict(cells) != computed_cells:
            raise InputError(f"Descriptive statistics payload {path} has cells inconsistent with vectors.")
        diagnostics = _payload_mapping(pair.get("diagnostics"), f"{path}.diagnostics")
        if set(diagnostics) != set(computed_diagnostics):
            raise InputError(f"Descriptive statistics payload {path}.diagnostics is incomplete.")
        for name, expected_rows in computed_diagnostics.items():
            rows = _payload_array(diagnostics.get(name), f"{path}.diagnostics.{name}")
            if rows != expected_rows or any(isinstance(row, bool) or not isinstance(row, int) for row in rows):
                raise InputError(f"Descriptive statistics payload {path}.diagnostics.{name} is inconsistent.")


def _validate_render_payload(payload: Mapping[str, Any]) -> Sequence[Mapping[str, Any]]:
    """Validate the complete payload contract consumed by the renderer.

    Args:
        payload: Candidate JSON-deserialized descriptive statistics payload.

    Returns:
        Validated question payloads in declared report order.

    Raises:
        InputError: If provenance, question counts, evidence, or chart metadata is invalid."""
    if not isinstance(payload, Mapping):
        raise InputError("Descriptive statistics payload root must be a JSON object.")
    if payload.get("payload_type") != "so2_descriptive_statistics":
        raise InputError("Expected a so2_descriptive_statistics payload.")
    if payload.get("payload_version") != "1":
        raise InputError("Unsupported descriptive statistics payload version.")
    provenance = _payload_mapping(payload.get("provenance"), "provenance")
    for field in ("source_path", "source_sha256", "worksheet", "alias", "export_date"):
        _payload_text(provenance.get(field), f"provenance.{field}")
    for field in (
        "header_row", "total_data_rows", "excluded_blank_rows", "included_response_rows",
    ):
        _payload_count(provenance.get(field), f"provenance.{field}")
    if not isinstance(provenance.get("schema_version"), (str, int)):
        raise InputError(
            "Descriptive statistics payload provenance.schema_version must be text or an integer."
        )
    _payload_mapping(payload.get("diagnostics"), "diagnostics")
    questions = _payload_array(payload.get("questions"), "questions")
    validated_questions: list[Mapping[str, Any]] = []
    for index, value in enumerate(questions):
        path = f"questions[{index}]"
        question = _payload_mapping(value, path)
        for field in ("question_id", "column", "label"):
            _payload_text(question.get(field), f"{path}.{field}")
        question_type = question.get("question_type")
        if question_type not in _QUESTION_TYPES:
            raise InputError(
                f"Descriptive statistics payload {path}.question_type is invalid: "
                f"{question_type!r}."
            )
        population = _payload_mapping(question.get("population"), f"{path}.population")
        population_counts = {
            field: _payload_count(population.get(field), f"{path}.population.{field}")
            for field in ("N", "A", "M")
        }
        if population_counts["N"] != population_counts["A"] + population_counts["M"]:
            raise InputError(
                f"Descriptive statistics payload {path}.population must satisfy N = A + M."
            )
        if "blank_rows" in population:
            blank_rows = _payload_count(
                population["blank_rows"], f"{path}.population.blank_rows"
            )
            if blank_rows != population_counts["M"]:
                raise InputError(
                    f"Descriptive statistics payload {path}.population.blank_rows must equal M."
                )
        categories = _payload_array(question.get("categories"), f"{path}.categories")
        for category_index, category_value in enumerate(categories):
            category_path = f"{path}.categories[{category_index}]"
            category = _payload_mapping(category_value, category_path)
            _payload_text(category.get("value"), f"{category_path}.value")
            _payload_count(category.get("count"), f"{category_path}.count")
            _payload_count(category.get("percent_base"), f"{category_path}.percent_base")
            percentage = category.get("percent")
            if percentage is not None and (
                isinstance(percentage, bool) or not isinstance(percentage, (int, float))
            ):
                raise InputError(
                    f"Descriptive statistics payload {category_path}.percent "
                    "must be numeric or null."
                )
        for field in ("contributions", "free_text_rows"):
            rows = _payload_array(question.get(field), f"{path}.{field}")
            if not all(isinstance(row, Mapping) for row in rows):
                raise InputError(
                    f"Descriptive statistics payload {path}.{field} must contain objects."
                )
            for row_index, row in enumerate(rows):
                row_path = f"{path}.{field}[{row_index}]"
                for text_field in ("country", "institution"):
                    _payload_text(row.get(text_field), f"{row_path}.{text_field}")
                _payload_count(row.get("source_row"), f"{row_path}.source_row")
                if field == "contributions":
                    _payload_text(row.get("value"), f"{row_path}.value")
                    if not isinstance(row.get("repeated_response"), bool):
                        raise InputError(
                            f"Descriptive statistics payload "
                            f"{row_path}.repeated_response must be boolean."
                        )
                else:
                    _payload_text(row.get("text"), f"{row_path}.text")
                    parent_answers = _payload_array(
                        row.get("parent_answers"), f"{row_path}.parent_answers"
                    )
                    parent_fields = ["parent_answers"]
                    if "raw_parent_answers" in row:
                        parent_fields.append("raw_parent_answers")
                    for parent_field in parent_fields:
                        parent_answers = _payload_array(
                            row.get(parent_field), f"{row_path}.{parent_field}"
                        )
                        for parent_index, parent_value in enumerate(parent_answers):
                            parent_path = f"{row_path}.{parent_field}[{parent_index}]"
                            parent = _payload_mapping(parent_value, parent_path)
                            _payload_text(parent.get("column"), f"{parent_path}.column")
                            _payload_text(parent.get("value"), f"{parent_path}.value")
        _payload_mapping(question.get("diagnostics"), f"{path}.diagnostics")
        applicability = _payload_mapping(
            question.get("applicability"), f"{path}.applicability"
        )
        status = applicability.get("status")
        if status not in {"unknown", "evaluated"}:
            raise InputError(
                f"Descriptive statistics payload {path}.applicability.status is invalid."
            )
        if status == "evaluated":
            _payload_text(applicability.get("column"), f"{path}.applicability.column")
            applicability_values = _payload_array(
                applicability.get("values"), f"{path}.applicability.values"
            )
            if not applicability_values or not all(
                isinstance(item, str) and item.strip() for item in applicability_values
            ):
                raise InputError(
                    f"Descriptive statistics payload {path}.applicability.values "
                    "must contain nonblank strings."
                )
            for field in (
                "applicable_rows", "inapplicable_rows", "eligible_unanswered_rows",
                "structurally_skipped_rows", "eligibility_unknown_rows",
                "out_of_route_answered_rows",
            ):
                _payload_count(
                    applicability.get(field), f"{path}.applicability.{field}"
                )
        if question_type == "free_text" and categories:
            raise InputError(
                f"Descriptive statistics payload {path}.categories must be empty for free_text."
            )
        if question_type != "free_text":
            pie_categories = _payload_array(
                question.get("pie_categories"), f"{path}.pie_categories"
            )
            for category_index, category_value in enumerate(pie_categories):
                category_path = f"{path}.pie_categories[{category_index}]"
                category = _payload_mapping(category_value, category_path)
                _payload_text(category.get("value"), f"{category_path}.value")
                _payload_count(category.get("count"), f"{category_path}.count")
                _payload_count(
                    category.get("percent_base"), f"{category_path}.percent_base"
                )
                percentage = category.get("percent")
                if percentage is not None and (
                    isinstance(percentage, bool)
                    or not isinstance(percentage, (int, float))
                ):
                    raise InputError(
                        f"Descriptive statistics payload {category_path}.percent "
                        "must be numeric or null."
                    )
                if not isinstance(category.get("excluded_from_chart"), bool):
                    raise InputError(
                        f"Descriptive statistics payload "
                        f"{category_path}.excluded_from_chart must be boolean."
                    )
        validated_questions.append(question)
    association_keys = {"association_heatmap_definitions", "association_heatmaps"}
    present_association_keys = association_keys.intersection(payload)
    if present_association_keys and present_association_keys != association_keys:
        raise InputError("Descriptive statistics payload association data is incomplete.")
    if present_association_keys:
        validate_association_heatmap_payload(payload)
    if "upset_vectors" not in payload:
        return validated_questions
    vectors = _payload_mapping(payload.get("upset_vectors"), "upset_vectors")
    multi_choice_questions = {
        str(question["question_id"]): question for question in validated_questions
        if question["question_type"] == "multi_choice"
    }
    if set(vectors) != set(multi_choice_questions):
        raise InputError("Descriptive statistics payload upset_vectors must cover exactly the multi-choice questions.")
    for question_id, rows in vectors.items():
        rows = _payload_array(rows, f"upset_vectors.{question_id}")
        categories = {str(item["value"]) for item in multi_choice_questions[question_id]["categories"] if item["value"] != "Missing"}
        source_rows: set[int] = set()
        for index, vector_value in enumerate(rows):
            vector = _payload_mapping(vector_value, f"upset_vectors.{question_id}[{index}]")
            source_row = _payload_count(vector.get("source_row"), f"upset_vectors.{question_id}[{index}].source_row")
            if source_row in source_rows:
                raise InputError(f"Descriptive statistics payload upset_vectors.{question_id} has duplicate source_row {source_row}.")
            source_rows.add(source_row)
            for text_field in ("country", "institution"):
                _payload_text(vector.get(text_field), f"upset_vectors.{question_id}[{index}].{text_field}")
            state = vector.get("state")
            if state not in {"answered", "missing", "inapplicable"}:
                raise InputError(f"Descriptive statistics payload upset_vectors.{question_id}[{index}].state is invalid.")
            selected = _payload_array(vector.get("selected_categories"), f"upset_vectors.{question_id}[{index}].selected_categories")
            if len(selected) != len(set(selected)) or not all(isinstance(item, str) and item in categories for item in selected):
                raise InputError(f"Descriptive statistics payload upset_vectors.{question_id}[{index}] has undeclared or duplicate selected categories.")
            if (state == "answered") != bool(selected):
                raise InputError(f"Descriptive statistics payload upset_vectors.{question_id}[{index}] has inconsistent state and selections.")
    return validated_questions


def render_descriptive_tex(
    payload: Mapping[str, Any],
    chart_dir: str | Path | None,
    report_path: str | Path | None = None,
    include_contribution_tables: bool = False,
    include_parent_context: bool = False,
    max_piechart_ratio: float = 4 / 3,
    upset_assets: Any | None = None,
) -> RenderedDescriptiveReport:
    """Render a self-contained report and reusable standalone chart documents.

    Args:
        payload: Validated descriptive-statistics payload.
        chart_dir: Optional target directory used to derive printable chart paths.
        report_path: Optional report path used as the base for relative chart paths.
        include_contribution_tables: Whether optional grouped structured-response
            evidence tables are included in the report body.
        include_parent_context: Whether free-text tables include complete raw parent
            answers. The default omits the parent-context column.
        max_piechart_ratio: Maximum permitted horizontal-to-vertical pie ratio.
        upset_assets: Optional prevalidated external UpSet asset state.

    Returns:
        Report TeX, shared chart fragments, standalone documents, and chart paths.

    Raises:
        InputError: If the payload is malformed or chart metadata cannot be
            converted into stable filenames.
    """
    if max_piechart_ratio <= 0:
        raise InputError("max_piechart_ratio must be positive")
    questions = _validate_render_payload(payload)
    association_definitions: tuple[AssociationHeatmapDefinition, ...] = ()
    association_pairs: Mapping[str, Any] = {}
    if "association_heatmap_definitions" in payload:
        metadata = _payload_mapping(
            payload["association_heatmap_definitions"], "association_heatmap_definitions"
        )
        association_definitions = tuple(
            _association_definition_from_payload(
                value, f"association_heatmap_definitions.definitions[{index}]"
            )
            for index, value in enumerate(metadata["definitions"])
        )
        association_pairs = _payload_mapping(
            payload["association_heatmaps"], "association_heatmaps"
        )
    question_positions = {
        str(question["question_id"]): position for position, question in enumerate(questions)
    }
    association_after_question: dict[str, list[AssociationHeatmapDefinition]] = {}
    for definition in association_definitions:
        if definition.mode != "default":
            continue
        positions = [
            question_positions.get(definition.row_question_id),
            question_positions.get(definition.column_question_id),
        ]
        if None in positions:
            continue
        later_question_id = max(
            (definition.row_question_id, definition.column_question_id),
            key=lambda question_id: question_positions[question_id],
        )
        association_after_question.setdefault(later_question_id, []).append(definition)
    fragments: dict[str, str] = {}
    chart_documents: dict[str, str] = {}
    chart_paths: dict[str, str] = {}
    report = [_preamble(), r"\begin{document}", r"\section*{SO2 descriptive statistics}"]
    provenance = payload.get("provenance", {})
    report.append(rf"Source path: {_tex(provenance['source_path'])}\par")
    report.append(rf"Alias: {_tex(provenance.get('alias', 'Unknown'))}\par")
    report.append(rf"Export date: {_tex(provenance.get('export_date', 'Unknown'))}\par")
    report.append(rf"Source SHA-256: {_tex(provenance.get('source_sha256', 'Unknown'))}\par")
    report.append(rf"Worksheet: {_tex(provenance.get('worksheet', 'Unknown'))}\par")
    report.append(rf"Header row: {_tex(provenance['header_row'])}\par")
    report.append(rf"Total data rows: {_tex(provenance['total_data_rows'])}\par")
    report.append(
        rf"Included response rows: {_tex(provenance.get('included_response_rows', 'Unknown'))}\par"
    )
    report.append(
        rf"Excluded blank rows: {_tex(provenance.get('excluded_blank_rows', 'Unknown'))}\par"
    )
    report.append(
        rf"Descriptive schema version: {_tex(provenance.get('schema_version', 'Unknown'))}\par"
    )
    report.extend([
        r"\tableofcontents",
        r"\clearpage",
        r"\section*{Abbreviations}",
        r"\noindent CC: ISO 3166-1 alpha-2 country code; N: included response rows; A: rows with a nonblank answer; M: missing responses (N - A).\par",
        r"\noindent oAR: percentage of answering rows; oIR: percentage of included response rows.\par",
        r"\clearpage",
    ])
    active_matrix_parent: str | None = None

    def append_upset_pair(
        definition_id: str,
        heading: str,
        anchor: bool = False,
        source_selectors: Sequence[str] = (),
    ) -> None:
        """Append one validated figure pair or its explicit omission note.

        Args:
            definition_id: Approved UpSet definition identifier.
            heading: TeX subsection title introducing the figure pair.
            anchor: Whether to create a cross-reference target for shared figures.
            source_selectors: Source question prefixes whose full labels are
                printed below a shared figure title.

        Returns:
            ``None`` after appending report TeX fragments.
        """
        if upset_assets is None:
            return
        label = rf"\label{{upset-{definition_id}}}" if anchor else ""
        report.append(rf"\subsection{{{_tex(heading)}}}{label}")
        for selector in source_selectors:
            source_question = next((
                item for item in questions
                if str(item["question_id"]).startswith(f"{selector}_")
            ), None)
            if source_question is None:
                continue
            report.append(rf"\noindent {_tex(selector)}: {_tex(source_question['label'])}\par")
        if upset_assets.state == "empty":
            report.append("UpSet charts omitted: no rendered assets.")
            return
        figure = upset_assets.figures[definition_id]
        report.extend([
            rf"\includegraphics[width=\linewidth]{{\detokenize{{{figure.upset_pdf}}}}}",
            rf"\includegraphics[width=\linewidth]{{\detokenize{{{figure.deviation_pdf}}}}}",
        ])

    for question in questions:
        for field in ("question_id", "label", "question_type", "categories", "population",
                      "contributions", "free_text_rows", "applicability"):
            if field not in question:
                raise InputError(f"Descriptive statistics question lacks {field}: {question!r}.")
        matrix_parent = question.get("form", {}).get("matrix_parent") if isinstance(question.get("form"), Mapping) else None
        matrix_row = question.get("form", {}).get("matrix_row") if isinstance(question.get("form"), Mapping) else None
        if matrix_parent:
            if matrix_parent != active_matrix_parent:
                report.extend([r"\clearpage", rf"\section{{{_tex(matrix_parent)}}}"])
                report.append(rf"{_tex(_response_structure_text(question))}\par")
                active_matrix_parent = matrix_parent
            report.append(rf"\subsection{{{_tex(matrix_row or question['label'])}}}")
        else:
            report.extend([r"\clearpage", rf"\section{{{_tex(question['label'])}}}"])
            active_matrix_parent = None
        population = question["population"]
        report.append(rf"\noindent\smaller[3] Question identifier: {_question_identifier(question['question_id'])}\normalsize\par\vspace{{0.35\baselineskip}}")
        report.append(rf"N/A/M (included/answered/missing): {population['N']} / {population['A']} / {population['M']}\par")
        if not matrix_parent:
            report.append(rf"{_tex(_response_structure_text(question))}\par")
        if question["question_type"] == "free_text":
            report.append(
                _question_tables(question, include_parent_context=include_parent_context)
                or r"\emph{No text responses.}"
            )
            continue
        use_pie = question["question_type"] == "single_choice"
        variants = (False, True) if use_pie and population["M"] else ((True,) if use_pie else (False,))
        question_paths: list[str] = []
        for answered_only in variants:
            key = _chart_key(question, answered_only)
            fragment = _pie_fragment(question, answered_only, max_piechart_ratio) if use_pie else _bar_fragment(question)
            fragments[key] = fragment
            chart_documents[key] = _standalone_tex(question, fragment, answered_only)
            report.append(fragment)
            if chart_dir is not None:
                chart_pdf = Path(chart_dir) / f"{key}.pdf"
                base = Path(report_path).parent if report_path is not None else Path.cwd()
                chart_paths[key] = Path(os.path.relpath(chart_pdf, start=base)).as_posix()
                question_paths.append(chart_paths[key])
                report.extend([
                    r"\noindent",
                    r"\begingroup\raggedright\smaller[3]",
                    rf"\url{{{_tex(chart_paths[key])}}}\par",
                    r"\endgroup",
                ])
        if chart_dir is not None and isinstance(question, dict):
            question["chart_paths"] = question_paths
        tables = _question_tables(
            question, include_parent_context=include_parent_context,
        ) if (question["question_type"] == "free_text" or include_contribution_tables) else ""
        if tables:
            report.append(tables)
        question_id = str(question["question_id"])
        matching_single = next((item for item in _UPSET_SINGLE_DEFINITIONS if question_id.startswith(f"{item}_")), None)
        if matching_single is not None:
            append_upset_pair(matching_single, "UpSet intersections")
        for definition_id, source_selectors in _UPSET_SHARED_DEFINITIONS.items():
            if any(question_id.startswith(f"{selector}_") for selector in source_selectors):
                report.append(
                    rf"\noindent Shared UpSet intersections: "
                    rf"\hyperref[upset-{definition_id}]{{see { _tex(definition_id) }}}.\par"
                )
        for definition in association_after_question.get(question_id, []):
            pair = _payload_mapping(
                association_pairs[definition.definition_id],
                f"association_heatmaps.{definition.definition_id}",
            )
            key = f"association-{definition.definition_id}"
            fragment = "\n".join(filter(None, [
                _association_heatmap_fragment(definition, pair),
                _association_conditional_summary_fragment(definition, pair),
            ]))
            fragments[key] = fragment
            chart_documents[key] = _association_standalone_tex(definition, pair)
            report.append(fragment)
            if chart_dir is not None:
                chart_pdf = Path(chart_dir) / f"{key}.pdf"
                base = Path(report_path).parent if report_path is not None else Path.cwd()
                chart_paths[key] = Path(os.path.relpath(chart_pdf, start=base)).as_posix()
                report.extend([
                    r"\noindent",
                    r"\begingroup\raggedright\smaller[3]",
                    rf"\url{{{_tex(chart_paths[key])}}}\par",
                    r"\endgroup",
                ])
    if upset_assets is not None:
        report.extend([r"\clearpage", r"\section{Shared UpSet charts}"])
        for definition_id, source_selectors in _UPSET_SHARED_DEFINITIONS.items():
            append_upset_pair(
                definition_id, definition_id, anchor=True,
                source_selectors=source_selectors,
            )
    report.extend([r"\end{document}", ""])
    if chart_dir is not None and isinstance(payload, dict):
        payload["chart_paths"] = dict(chart_paths)
    return RenderedDescriptiveReport(
        tex="\n".join(report),
        chart_fragments=MappingProxyType(fragments),
        chart_documents=MappingProxyType(chart_documents),
        chart_paths=MappingProxyType(chart_paths),
    )


def _run_xelatex(source: Path, output_dir: Path, passes: int = 1) -> Path:
    """Compile one staged TeX source the required number of times.

    Args:
        source: Staged TeX source file to compile.
        output_dir: Directory receiving compiler intermediates and the PDF.
        passes: Positive number of XeLaTeX runs, normally two for the report ToC.

    Returns:
        Completed PDF path in ``output_dir``.

    Raises:
        AssertionError: If ``passes`` is less than one.
        InputError: If XeLaTeX is unavailable or compilation fails."""
    if passes < 1:
        raise AssertionError("XeLaTeX compilation requires at least one pass.")
    compiler = shutil.which("xelatex")
    if compiler is None:
        raise InputError("XeLaTeX is required to render descriptive report PDFs.")
    pdf = output_dir / f"{source.stem}.pdf"
    for _ in range(passes):
        result = subprocess.run(
            [compiler, "-interaction=nonstopmode", "-halt-on-error", "-output-directory", str(output_dir), str(source)],
            cwd=source.parent, capture_output=True, text=True, check=False,
        )
        if result.returncode != 0 or not pdf.is_file():
            details = (result.stderr or result.stdout).strip()
            raise InputError(f"XeLaTeX failed for {source.name}: {details or 'no PDF was produced'}")
    return pdf


def _require_new_or_empty_chart_dir(chart_dir: Path, overwrite: bool = False) -> None:
    """Validate chart-directory safety before publication.

    Args:
        chart_dir: Target directory for standalone chart PDFs.
        overwrite: Whether a nonempty existing directory may be transactionally replaced.

    Returns:
        ``None`` after the path satisfies the publication precondition.

    Raises:
        InputError: If the path is a symbolic link, non-directory, or unsafe existing directory."""
    if chart_dir.is_symlink():
        raise InputError(f"Chart directory must not be a symbolic link: {chart_dir}")
    if chart_dir.exists() and (not chart_dir.is_dir() or (not overwrite and any(chart_dir.iterdir()))):
        raise InputError(f"Chart directory must be new or empty: {chart_dir}")


def _write_compilation_source(path: Path, content: str, description: str) -> None:
    """Write one temporary TeX source with an actionable error boundary.

    Args:
        path: Temporary source path to create or replace.
        content: Complete TeX text to write using UTF-8.
        description: Human-readable source role included in write errors.

    Returns:
        ``None`` after the source is written.

    Raises:
        InputError: If the temporary source cannot be written."""
    try:
        path.write_text(content, encoding="utf-8")
    except OSError as exc:
        raise InputError(f"Could not write {description} {path.name}: {exc}") from exc


def _stage_publication_file(
    target: Path, *, text: str | None = None, source: Path | None = None,
) -> Path:
    """Create one complete hidden staging file on the target filesystem.

    Args:
        target: Final output path whose parent hosts the staging file.
        text: Optional UTF-8 content to stage directly.
        source: Optional existing file to copy into the stage.

    Returns:
        Hidden complete staging-file path ready for atomic promotion.

    Raises:
        AssertionError: If neither or both source forms are provided.
        OSError: If staging content cannot be written or copied."""
    if (text is None) == (source is None):
        raise AssertionError("Exactly one publication source must be provided.")
    descriptor, stage_name = tempfile.mkstemp(
        prefix=f".{target.name}.so2-stage-", dir=target.parent,
    )
    os.close(descriptor)
    stage = Path(stage_name)
    try:
        if text is not None:
            stage.write_text(text, encoding="utf-8")
        else:
            shutil.copy2(source, stage)
    except OSError:
        stage.unlink(missing_ok=True)
        raise
    return stage


def _discard_publication_stages(stages: Sequence[Path]) -> None:
    """Delete publication staging files that were not promoted.

    Args:
        stages: Staging-file paths to remove if they still exist.

    Returns:
        ``None`` after best-effort stage cleanup."""
    for stage in stages:
        stage.unlink(missing_ok=True)


def render_descriptive_pdf(
    rendered: RenderedDescriptiveReport,
    tex_path: str | Path,
    pdf_path: str | Path | None,
    chart_dir: str | Path | None,
    overwrite: bool = False,
) -> None:
    """Compile and transactionally publish a descriptive report and charts.

    Args:
        rendered: Self-contained TeX and standalone chart documents.
        tex_path: Target for the published report source.
        pdf_path: Optional target for the compiled report PDF.
        chart_dir: Optional output directory for standalone chart PDFs.
        overwrite: Whether existing outputs may be replaced after successful staging.

    Returns:
        None.

    Raises:
        InputError: If targets are unsafe, staging or compilation fails, or
            publication cannot complete without partial output.
    """
    target_tex = Path(tex_path)
    target_pdf = Path(pdf_path) if pdf_path is not None else None
    target_charts = Path(chart_dir) if chart_dir is not None else None
    if target_charts is not None:
        _require_new_or_empty_chart_dir(target_charts, overwrite)
    targets = [target_tex, *(path for path in (target_pdf, target_charts) if path is not None)]
    if any(not target.parent.exists() for target in targets):
        raise InputError("Output parent directory does not exist.")
    with tempfile.TemporaryDirectory(prefix="so2-report-") as report_temporary, \
            tempfile.TemporaryDirectory(prefix="so2-charts-") as chart_temporary:
        report_stage = Path(report_temporary)
        chart_stage = Path(chart_temporary)

        chart_pdfs: dict[str, Path] = {}
        if target_charts is not None:
            for key, document in rendered.chart_documents.items():
                source = chart_stage / f"{key}.tex"
                _write_compilation_source(source, document, "chart source")
                chart_pdfs[key] = _run_xelatex(source, chart_stage)

        report_pdf = None
        if target_pdf is not None:
            report_source = report_stage / "report.tex"
            _write_compilation_source(report_source, rendered.tex, "report source")
            report_pdf = _run_xelatex(report_source, report_stage, passes=2)
        publication_stages: list[Path] = []
        try:
            staged_tex = _stage_publication_file(target_tex, text=rendered.tex)
            publication_stages.append(staged_tex)
            staged_pdf = None
            if target_pdf is not None and report_pdf is not None:
                staged_pdf = _stage_publication_file(target_pdf, source=report_pdf)
                publication_stages.append(staged_pdf)
        except OSError as exc:
            _discard_publication_stages(publication_stages)
            raise InputError(f"Could not stage descriptive report outputs: {exc}") from exc

        publish_charts = None
        if target_charts is not None:
            _require_new_or_empty_chart_dir(target_charts, overwrite)
            publish_charts = Path(tempfile.mkdtemp(prefix=".so2-charts-", dir=target_charts.parent))
            try:
                for key, source_pdf in chart_pdfs.items():
                    shutil.copy2(source_pdf, publish_charts / f"{key}.pdf")
            except OSError as exc:
                shutil.rmtree(publish_charts, ignore_errors=True)
                _discard_publication_stages(publication_stages)
                raise InputError(f"Could not stage chart PDFs for publication: {exc}") from exc

        published: list[Path] = []
        backups: dict[Path, Path] = {}
        chart_dir_was_empty = target_charts is not None and target_charts.exists()

        def replace_file(source: Path, target: Path) -> None:
            """Publish one file while retaining its prior value for rollback."""
            if target.exists():
                descriptor, backup_name = tempfile.mkstemp(
                    prefix=".so2-report-backup-", dir=target.parent,
                )
                os.close(descriptor)
                backup = Path(backup_name)
                os.replace(target, backup)
                backups[target] = backup
            os.replace(source, target)
            published.append(target)

        def replace_directory(source: Path, target: Path) -> None:
            """Publish a staged directory while retaining its prior directory for rollback."""
            if target.exists():
                backup = Path(tempfile.mkdtemp(prefix=".so2-charts-backup-", dir=target.parent))
                backup.rmdir()
                os.replace(target, backup)
                backups[target] = backup
            os.replace(source, target)
            published.append(target)

        try:
            replace_file(staged_tex, target_tex)
            if target_pdf is not None and staged_pdf is not None:
                replace_file(staged_pdf, target_pdf)
            if target_charts is not None and publish_charts is not None:
                replace_directory(publish_charts, target_charts)
        except OSError as exc:
            rollback_failures: list[OSError] = []
            for target in reversed(published):
                try:
                    if target.is_dir():
                        shutil.rmtree(target)
                    else:
                        target.unlink(missing_ok=True)
                except OSError as rollback_exc:
                    rollback_failures.append(rollback_exc)
            for target, backup in backups.items():
                try:
                    os.replace(backup, target)
                except OSError as rollback_exc:
                    rollback_failures.append(rollback_exc)
            if chart_dir_was_empty and target_charts is not None:
                try:
                    if not target_charts.exists():
                        target_charts.mkdir()
                    elif not target_charts.is_dir() or any(target_charts.iterdir()):
                        raise OSError(f"chart directory has unexpected contents: {target_charts}")
                except OSError as rollback_exc:
                    rollback_failures.append(rollback_exc)
            if publish_charts is not None:
                try:
                    if publish_charts.exists():
                        shutil.rmtree(publish_charts)
                except OSError as rollback_exc:
                    rollback_failures.append(rollback_exc)
            _discard_publication_stages(publication_stages)
            if rollback_failures:
                details = "; ".join(str(rollback_exc) for rollback_exc in rollback_failures)
                raise InputError(
                    "Could not publish descriptive report outputs: "
                    f"{exc}; rollback could not establish a clean state: {details}"
                ) from exc
            raise InputError(f"Could not publish descriptive report outputs: {exc}") from exc
        _discard_publication_stages(publication_stages)
        for backup in backups.values():
            if backup.is_dir():
                shutil.rmtree(backup)
            else:
                backup.unlink(missing_ok=True)
