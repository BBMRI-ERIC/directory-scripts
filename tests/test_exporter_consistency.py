"""Test exporter consistency behavior."""

import copy
import io
import json
import runpy
import sys
import types
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path


REPO_ROOT = Path(__file__).resolve().parent.parent


class SharedDirectoryStub:
    """Directory stub shared across exporter consistency tests.
    """

    BASE_BIOBANKS = [
        {
            "id": "bbmri-eric:ID:CZ_BB1",
            "name": "Biobank 1",
            "country": "CZ",
            "contact": {"id": "ct1"},
            "network": [{"id": "net1"}],
            "longitude": "14.42076",
            "latitude": "50.08804",
            "withdrawn": False,
        },
        {
            "id": "bbmri-eric:ID:EXT_BB2",
            "name": "Biobank 2",
            "country": "DE",
            "contact": {"id": "ct2"},
            "network": [{"id": "net2"}],
            "longitude": "13.4050",
            "latitude": "52.5200",
            "withdrawn": False,
        },
        {
            "id": "bbmri-eric:ID:NL_BB3",
            "name": "Biobank 3",
            "country": "NL",
            "contact": {"id": "ct3"},
            "network": [{"id": "net3"}],
            "withdrawn": True,
        },
    ]
    BASE_COLLECTIONS = [
        {
            "id": "col1",
            "name": "Collection 1",
            "biobank": {"id": "bbmri-eric:ID:CZ_BB1"},
            "country": "CZ",
            "contact": {"id": "ct1"},
            "networks": [{"id": "net1"}],
            "type": ["DISEASE_SPECIFIC"],
            "materials": ["SERUM"],
            "order_of_magnitude": 2,
            "order_of_magnitude_donors": 1,
            "size": 100,
            "number_of_donors": 10,
            "withdrawn": False,
            "facts": [{"id": "f1"}],
            "studies": [{"id": "study1"}, {"id": "study2"}],
            "longitude": "14.5000",
            "latitude": "50.1000",
        },
        {
            "id": "col2",
            "name": "Collection 2",
            "biobank": {"id": "bbmri-eric:ID:CZ_BB1"},
            "country": "CZ",
            "contact": {"id": "ct1"},
            "networks": [{"id": "net1"}],
            "type": ["CASE_CONTROL"],
            "materials": ["DNA"],
            "parent_collection": {"id": "col1"},
            "order_of_magnitude": 1,
            "size": 5,
            "number_of_donors": 3,
            "withdrawn": False,
            "facts": [{"id": "f2"}],
        },
        {
            "id": "col3",
            "name": "Collection 3",
            "biobank": {"id": "bbmri-eric:ID:EXT_BB2"},
            "country": "DE",
            "contact": {"id": "ct2"},
            "networks": [{"id": "net2"}],
            "type": ["POPULATION"],
            "materials": ["PLASMA"],
            "order_of_magnitude": 3,
            "order_of_magnitude_donors": 2,
            "withdrawn": False,
            "facts": [{"id": "f3"}],
            "studies": [{"id": "study2"}],
        },
        {
            "id": "col4",
            "name": "Collection 4",
            "biobank": {"id": "bbmri-eric:ID:EXT_BB2"},
            "country": "DE",
            "contact": {"id": "ct2"},
            "networks": [{"id": "net2"}],
            "type": ["LONGITUDINAL"],
            "materials": ["SERUM"],
            "order_of_magnitude": 1,
            "withdrawn": True,
            "facts": [{"id": "f4"}],
        },
        {
            "id": "col5",
            "name": "Collection 5",
            "biobank": {"id": "bbmri-eric:ID:NL_BB3"},
            "country": "NL",
            "contact": {"id": "ct3"},
            "networks": [{"id": "net3"}],
            "type": ["CASE_CONTROL"],
            "materials": ["SERUM"],
            "order_of_magnitude": 2,
            "order_of_magnitude_donors": 2,
            "size": 12,
            "number_of_donors": 8,
            "withdrawn": False,
            "facts": [{"id": "f5"}],
            "studies": [{"id": "study3"}],
        },
    ]
    BASE_CONTACTS = [
        {"id": "ct1", "email": "ct1@example.org", "country": "CZ"},
        {"id": "ct2", "email": "ct2@example.org", "country": "DE"},
        {"id": "ct3", "email": "ct3@example.org", "country": "NL"},
    ]
    BASE_NETWORKS = [
        {"id": "net1", "name": "Network 1", "country": {"id": "CZ"}, "contact": {"id": "ct1"}},
        {"id": "net2", "name": "Network 2", "country": {"id": "DE"}, "contact": {"id": "ct2"}},
        {"id": "net3", "name": "Network 3", "country": {"id": "NL"}, "contact": {"id": "ct3"}},
    ]
    BASE_FACTS = {
        "col1": [
            {
                "id": "f1",
                "sex": "*",
                "age_range": "*",
                "sample_type": "*",
                "disease": "*",
                "number_of_samples": 100,
                "number_of_donors": 10,
            }
        ],
        "col2": [
            {
                "id": "f2",
                "sex": "*",
                "age_range": "*",
                "sample_type": "*",
                "disease": "*",
                "number_of_samples": 5,
                "number_of_donors": 3,
            }
        ],
        "col3": [
            {
                "id": "f3",
                "sex": "*",
                "age_range": "*",
                "sample_type": "*",
                "disease": "*",
                "number_of_samples": 1000,
                "number_of_donors": 100,
            }
        ],
        "col4": [
            {
                "id": "f4",
                "sex": "*",
                "age_range": "*",
                "sample_type": "*",
                "disease": "*",
                "number_of_samples": 10,
                "number_of_donors": 5,
            }
        ],
        "col5": [
            {
                "id": "f5",
                "sex": "*",
                "age_range": "*",
                "sample_type": "*",
                "disease": "*",
                "number_of_samples": 12,
                "number_of_donors": 8,
            }
        ],
    }
    BASE_SERVICES = [
        {"id": "svc1", "biobank": {"id": "bbmri-eric:ID:CZ_BB1"}, "serviceTypes": ["SEQUENCING"]},
        {"id": "svc2", "biobank": {"id": "bbmri-eric:ID:NL_BB3"}, "serviceTypes": ["BIOANALYTICAL_SERVICES"]},
    ]
    BASE_STUDIES = [
        {
            "id": "study1",
            "title": "Study 1",
            "longitude": "14.6000",
            "latitude": "50.2000",
        },
        {
            "id": "study2",
            "title": "Study 2",
        },
        {
            "id": "study3",
            "title": "Study 3",
        },
    ]

    def __init__(self, *args, **kwargs):
        """Deep-copy entity fixtures and configure schema, target, and withdrawal scope.

        Args:
            *args: Unused positional Directory constructor arguments, accepted for CLI compatibility.
            **kwargs: Directory options: schema, directory_url, and include/only_withdrawn_entities; remaining keys are ignored.
        """
        self.include_withdrawn_entities = kwargs.get("include_withdrawn_entities", False) or kwargs.get("only_withdrawn_entities", False)
        self.only_withdrawn_entities = kwargs.get("only_withdrawn_entities", False)
        self._schema = kwargs.get("schema", "ERIC")
        self._directory_url = kwargs.get("directory_url", "https://directory.example.test")
        self.biobanks = copy.deepcopy(self.BASE_BIOBANKS)
        self.collections = copy.deepcopy(self.BASE_COLLECTIONS)
        self.contacts = copy.deepcopy(self.BASE_CONTACTS)
        self.networks = copy.deepcopy(self.BASE_NETWORKS)
        self.services = copy.deepcopy(self.BASE_SERVICES)
        self.studies = copy.deepcopy(self.BASE_STUDIES)
        self.collectionFactMap = copy.deepcopy(self.BASE_FACTS)
        self.contactHashmap = {contact["id"]: contact for contact in self.contacts}

    def _matches_withdrawn_scope(self, is_withdrawn):
        """Return whether the configured withdrawal state is in scope.

        Args:
            is_withdrawn: Fixture withdrawal flag evaluated against the requested scope.

        Returns:
            True when the flag satisfies only/include-withdrawn settings; only-withdrawn takes precedence.
        """
        if self.only_withdrawn_entities:
            return is_withdrawn
        if self.include_withdrawn_entities:
            return True
        return not is_withdrawn

    def getSchema(self):
        """Return schema string selected for the fake Directory session.

        Returns:
            The schema string selected for the fake Directory session.
        """
        return self._schema

    def getDirectoryUrl(self):
        """Return endpoint string assigned to the fake Directory session.

        Returns:
            Fixed synthetic endpoint URL; no connection is opened.
        """
        return self._directory_url

    def isBiobankWithdrawn(self, biobank_id):
        """Report the fixture marks the requested biobank as withdrawn.

        Args:
            biobank_id: Biobank identifier whose fixture withdrawal status this stub reports.

        Returns:
            Whether the fixture marks the requested biobank as withdrawn.
        """
        biobank = next(
            biobank for biobank in self.biobanks if biobank["id"] == biobank_id
        )
        return bool(biobank.get("withdrawn"))

    def isCollectionWithdrawn(self, collection_id):
        """Report the fixture marks the requested collection as withdrawn.

        Args:
            collection_id: Collection identifier whose fixture withdrawal status this stub reports.

        Returns:
            Whether the fixture marks the requested collection as withdrawn.
        """
        collection = next(
            collection for collection in self.collections if collection["id"] == collection_id
        )
        if bool(collection.get("withdrawn")):
            return True
        if self.isBiobankWithdrawn(collection["biobank"]["id"]):
            return True
        parent = collection.get("parent_collection")
        if parent is not None:
            return self.isCollectionWithdrawn(parent["id"])
        return False

    def getBiobanks(self):
        """Return synthetic biobank records available to the code under test.

        Returns:
            The synthetic biobank records available to the code under test.
        """
        return [
            biobank
            for biobank in self.biobanks
            if self._matches_withdrawn_scope(self.isBiobankWithdrawn(biobank["id"]))
        ]

    def getBiobanksCount(self):
        """Return number of synthetic biobank records visible in the selected scope.

        Returns:
            Number of biobank records exposed by this stub.
        """
        return len(self.getBiobanks())

    def getBiobankById(self, biobank_id, raise_on_missing=False):
        """Return synthetic biobank record selected by the requested identifier, or `None` when absent.

        Args:
            biobank_id: Biobank identifier whose fixture record this stub returns or omits.
            raise_on_missing: Whether the fixture lookup should raise instead of returning a missing record.

        Returns:
            The synthetic biobank record selected by the requested identifier, or `None` when absent.
        """
        for biobank in self.biobanks:
            if biobank["id"] == biobank_id:
                if not self._matches_withdrawn_scope(self.isBiobankWithdrawn(biobank_id)):
                    break
                return biobank
        if raise_on_missing:
            raise KeyError(biobank_id)
        return None

    def getLoadedBiobankById(self, biobank_id, raise_on_missing=False):
        """Return synthetic loaded biobank record selected by the requested identifier, or `None` when absent.

        Args:
            biobank_id: Biobank identifier whose loaded fixture record this stub returns or omits.
            raise_on_missing: Whether the fixture lookup should raise instead of returning a missing record.

        Returns:
            The synthetic loaded biobank record selected by the requested identifier, or `None` when absent.
        """
        for biobank in self.biobanks:
            if biobank["id"] == biobank_id:
                return biobank
        if raise_on_missing:
            raise KeyError(biobank_id)
        return None

    def getBiobankCountry(self, biobank_id):
        """Return fixture country code for the requested biobank.

        Args:
            biobank_id: Biobank identifier whose fixture country code this stub returns.

        Returns:
            The fixture country code for the requested biobank.
        """
        return self.getBiobankById(biobank_id)["country"]

    def getCollections(self):
        """Return synthetic collection records available to the code under test.

        Returns:
            The synthetic collection records available to the code under test.
        """
        return [
            collection
            for collection in self.collections
            if self._matches_withdrawn_scope(self.isCollectionWithdrawn(collection["id"]))
        ]

    def getCollectionsCount(self):
        """Return number of synthetic collection records visible in the selected scope.

        Returns:
            Number of collection records exposed by this stub.
        """
        return len(self.getCollections())

    def getCollectionBiobankId(self, collection_id):
        """Return fixture parent-biobank identifier for the requested collection.

        Args:
            collection_id: Collection identifier whose fixture parent-biobank ID this stub returns.

        Returns:
            The fixture parent-biobank identifier for the requested collection.
        """
        return self.getCollectionById(collection_id)["biobank"]["id"]

    def getCollectionById(self, collection_id, raise_on_missing=False):
        """Return synthetic collection record selected by the requested identifier, or `None` when absent.

        Args:
            collection_id: Collection identifier whose fixture record this stub returns or omits.
            raise_on_missing: Whether the fixture lookup should raise instead of returning a missing record.

        Returns:
            The synthetic collection record selected by the requested identifier, or `None` when absent.
        """
        for collection in self.collections:
            if collection["id"] == collection_id:
                if not self._matches_withdrawn_scope(
                    self.isCollectionWithdrawn(collection_id)
                ):
                    break
                return collection
        if raise_on_missing:
            raise KeyError(collection_id)
        return None

    def getLoadedCollectionById(self, collection_id, raise_on_missing=False):
        """Return synthetic loaded collection record selected by the requested identifier, or `None` when absent.

        Args:
            collection_id: Collection identifier whose loaded fixture record this stub returns or omits.
            raise_on_missing: Whether the fixture lookup should raise instead of returning a missing record.

        Returns:
            The synthetic loaded collection record selected by the requested identifier, or `None` when absent.
        """
        for collection in self.collections:
            if collection["id"] == collection_id:
                return collection
        if raise_on_missing:
            raise KeyError(collection_id)
        return None

    def getLoadedCollections(self):
        """Return synthetic unfiltered collection records retained by the Directory fixture.

        Returns:
            The synthetic unfiltered collection records retained by the Directory fixture.
        """
        return list(self.collections)

    def getCollectionFacts(self, collection_id):
        """Return synthetic fact rows associated with the requested collection.

        Args:
            collection_id: Collection identifier whose fixture fact rows this stub returns.

        Returns:
            The synthetic fact rows associated with the requested collection.
        """
        return self.collectionFactMap.get(collection_id, [])

    def getCollectionCountry(self, collection_id):
        """Return fixture country code for the requested collection.

        Args:
            collection_id: Collection identifier whose fixture country code this stub returns.

        Returns:
            The fixture country code for the requested collection.
        """
        collection = self.getCollectionById(collection_id)
        return collection["country"]

    def getCollectionStudies(self, collection_id):
        """Return fixture studies linked to the requested collection.

        Args:
            collection_id: Collection identifier whose fixture linked studies this stub returns.

        Returns:
            The fixture studies linked to the requested collection.
        """
        collection = self.getCollectionById(collection_id)
        if collection is None:
            return []
        study_ids = {
            study_ref["id"]
            for study_ref in collection.get("studies", [])
            if isinstance(study_ref, dict) and study_ref.get("id")
        }
        return [
            study
            for study in self.getStudies()
            if study["id"] in study_ids
        ]

    def getCollectionStudyIds(self, collection_id):
        """Return fixture study identifiers linked to the requested collection.

        Args:
            collection_id: Collection identifier whose fixture linked study identifiers this stub returns.

        Returns:
            The fixture study identifiers linked to the requested collection.
        """
        return [study["id"] for study in self.getCollectionStudies(collection_id)]

    def getContact(self, contact_id):
        """Return synthetic contact record selected by the requested identifier.

        Args:
            contact_id: Contact identifier whose fixture contact record this stub returns.

        Returns:
            The synthetic contact record selected by the requested identifier.
        """
        return self.contactHashmap[contact_id]

    def getContacts(self):
        """Return synthetic contact records available to the code under test.

        Returns:
            The synthetic contact records available to the code under test.
        """
        return self.contacts

    def getNetworks(self):
        """Return synthetic network records available to the code under test.

        Returns:
            The synthetic network records available to the code under test.
        """
        return self.networks

    def getServices(self):
        """Return synthetic service records available to the code under test.

        Returns:
            The synthetic service records available to the code under test.
        """
        return [
            service
            for service in self.services
            if self._matches_withdrawn_scope(
                self.isBiobankWithdrawn(service["biobank"]["id"])
            )
        ]

    def getServiceBiobankId(self, service_id):
        """Return fixture parent-biobank identifier for the requested service.

        Args:
            service_id: Service identifier whose fixture parent-biobank ID this stub returns.

        Returns:
            The fixture parent-biobank identifier for the requested service.
        """
        return next(
            service["biobank"]["id"]
            for service in self.services
            if service["id"] == service_id
        )

    def getStudies(self):
        """Return synthetic study records available to the code under test.

        Returns:
            The synthetic study records available to the code under test.
        """
        visible_study_ids = []
        for collection in self.getCollections():
            for study_ref in collection.get("studies", []):
                if not isinstance(study_ref, dict):
                    continue
                study_id = study_ref.get("id")
                if study_id and study_id not in visible_study_ids:
                    visible_study_ids.append(study_id)
        return [
            study
            for study in self.studies
            if study["id"] in visible_study_ids
        ]

    def getStudyById(self, study_id, raise_on_missing=False):
        """Return fixture study record selected by the requested identifier, or `None` when absent.

        Args:
            study_id: Study identifier whose fixture record this stub returns or omits.
            raise_on_missing: Whether the fixture lookup should raise instead of returning a missing record.

        Returns:
            The fixture study record selected by the requested identifier, or `None` when absent.
        """
        for study in self.getStudies():
            if study["id"] == study_id:
                return study
        if raise_on_missing:
            raise KeyError(study_id)
        return None

    def getStudyCollectionIds(self, study_id):
        """Return fixture collection identifiers linked to the requested study.

        Args:
            study_id: Study identifier whose fixture collection identifiers this stub returns.

        Returns:
            The fixture collection identifiers linked to the requested study.
        """
        collection_ids = []
        for collection in self.getCollections():
            study_ids = {
                study_ref["id"]
                for study_ref in collection.get("studies", [])
                if isinstance(study_ref, dict) and study_ref.get("id")
            }
            if study_id in study_ids:
                collection_ids.append(collection["id"])
        return collection_ids

    def getStudyBiobankIds(self, study_id):
        """Return fixture biobank identifiers linked to the requested study.

        Args:
            study_id: Study identifier whose fixture biobank identifiers this stub returns.

        Returns:
            The fixture biobank identifiers linked to the requested study.
        """
        biobank_ids = []
        for collection_id in self.getStudyCollectionIds(study_id):
            biobank_id = self.getCollectionBiobankId(collection_id)
            if biobank_id not in biobank_ids:
                biobank_ids.append(biobank_id)
        return biobank_ids

    def getBiobankStudies(self, biobank_id):
        """Return fixture studies linked to the requested biobank.

        Args:
            biobank_id: Biobank identifier whose fixture linked studies this stub returns.

        Returns:
            The fixture studies linked to the requested biobank.
        """
        return [
            study
            for study in self.getStudies()
            if biobank_id in self.getStudyBiobankIds(study["id"])
        ]

    def getStudyCountries(self, study_id):
        """Return sorted fixture country codes represented by the requested study.

        Args:
            study_id: Study identifier whose fixture country codes this stub returns.

        Returns:
            The sorted fixture country codes represented by the requested study.
        """
        return sorted(
            {
                self.getCollectionCountry(collection_id)
                for collection_id in self.getStudyCollectionIds(study_id)
            }
        )

    def isTopLevelCollection(self, collection_id):
        """Report the fixture collection has no parent collection.

        Args:
            collection_id: Collection identifier whose fixture hierarchy position this stub evaluates.

        Returns:
            Whether the fixture collection has no parent collection.
        """
        return "parent_collection" not in self.getCollectionById(collection_id)

    def isCountableCollection(self, collection_id, metric):
        """Report the fixture collection is countable for the requested metric.

        Args:
            collection_id: Collection identifier whose fixture counting eligibility this stub evaluates.
            metric: Fact-sheet metric whose countability is evaluated by the fixture.

        Returns:
            Whether the fixture collection is countable for the requested metric.
        """
        collection = self.getCollectionById(collection_id)
        if metric not in collection or not isinstance(collection[metric], int):
            return False
        parent = collection.get("parent_collection")
        while parent is not None:
            parent_collection = self.getLoadedCollectionById(parent["id"])
            if metric in parent_collection and isinstance(parent_collection[metric], int):
                return False
            parent = parent_collection.get("parent_collection")
        return True


def _run_script(monkeypatch, script_name, argv, directory_class=SharedDirectoryStub):
    """Execute an exporter in-process with a fake Directory and captured output streams.

    Args:
        monkeypatch: Pytest monkeypatch fixture; temporary dependency and environment overrides are undone after the test.
        script_name: Repository-relative Python CLI filename run as __main__.
        argv: Command-line vector supplied to the isolated script invocation.
        directory_class: Directory double injected into the exporter entry point.

    Returns:
        Tuple of (script globals dictionary, captured stdout text, captured stderr text).
    """
    fake_directory_module = types.ModuleType("directory")
    fake_directory_module.Directory = directory_class
    monkeypatch.setitem(sys.modules, "directory", fake_directory_module)
    monkeypatch.setattr(sys, "argv", [script_name, *argv])

    stdout = io.StringIO()
    stderr = io.StringIO()
    with redirect_stdout(stdout), redirect_stderr(stderr):
        globals_dict = runpy.run_path(
            str(REPO_ROOT / script_name),
            run_name="__main__",
        )
    return globals_dict, stdout.getvalue(), stderr.getvalue()


class WithdrawnCohortUnderActiveBiobankDirectoryStub(SharedDirectoryStub):
    """Directory stub with a withdrawn cohort collection under an active biobank.
    """

    BASE_COLLECTIONS = copy.deepcopy(SharedDirectoryStub.BASE_COLLECTIONS)
    BASE_COLLECTIONS[3]["type"] = ["COHORT"]


class CohortTotalsDirectoryStub(SharedDirectoryStub):
    """Directory stub with cohort collections covering explicit and OoM totals.
    """

    BASE_COLLECTIONS = copy.deepcopy(SharedDirectoryStub.BASE_COLLECTIONS)
    BASE_COLLECTIONS[0]["type"] = ["COHORT"]
    BASE_COLLECTIONS[1]["type"] = ["COHORT"]
    BASE_COLLECTIONS[2]["type"] = ["POPULATION_BASED"]


class CohortNoStarFallbackDirectoryStub(CohortTotalsDirectoryStub):
    """Cohort fixture with one fully concrete fact row and no margins.
    """

    BASE_FACTS = copy.deepcopy(CohortTotalsDirectoryStub.BASE_FACTS)
    BASE_FACTS["col1"].append(
        {
            "id": "f1-concrete",
            "sex": "FEMALE",
            "age_range": "Adult",
            "sample_type": "SERUM",
            "disease": "C50",
            "number_of_samples": 20,
            "number_of_donors": 8,
        }
    )


class FactSheetEmulationDirectoryStub(SharedDirectoryStub):
    """Directory fixture with one material-partitioned sibling family.
    """

    PARENT_ID = "bbmri-eric:ID:CZ_BB1:collection:legacy"
    BASE_COLLECTIONS = [
        {
            "id": PARENT_ID,
            "name": "Legacy collection",
            "biobank": {"id": "bbmri-eric:ID:CZ_BB1"},
            "country": "CZ",
            "contact": {"id": "ct1"},
            "storage_temperatures": ["temperatureRoom"],
            "license": "https://example.test/license",
            "type": ["SAMPLE"],
            "materials": ["SERUM", "DNA"],
            "size": 100,
            "withdrawn": False,
        },
        {
            "id": f"{PARENT_ID}:serum",
            "name": "Legacy collection - Serum",
            "biobank": {"id": "bbmri-eric:ID:CZ_BB1"},
            "parent_collection": {"id": PARENT_ID},
            "country": "CZ",
            "contact": {"id": "ct1"},
            "storage_temperatures": ["temperatureRoom"],
            "license": "https://example.test/license",
            "type": ["SAMPLE"],
            "materials": ["SERUM"],
            "size": 60,
            "number_of_donors": 40,
            "withdrawn": False,
        },
        {
            "id": f"{PARENT_ID}:dna",
            "name": "Legacy collection - DNA",
            "biobank": {"id": "bbmri-eric:ID:CZ_BB1"},
            "parent_collection": {"id": PARENT_ID},
            "country": "CZ",
            "contact": {"id": "ct1"},
            "storage_temperatures": ["temperatureRoom"],
            "license": "https://example.test/license",
            "type": ["SAMPLE"],
            "materials": ["DNA"],
            "size": 40,
            "number_of_donors": 30,
            "withdrawn": False,
        },
    ]
    BASE_FACTS = {
        PARENT_ID: [],
        f"{PARENT_ID}:serum": [],
        f"{PARENT_ID}:dna": [],
    }


class FactSheetEmulationReviewDirectoryStub(FactSheetEmulationDirectoryStub):
    """Emulation fixture with anatomy evidence needing external review.
    """

    BASE_COLLECTIONS = copy.deepcopy(FactSheetEmulationDirectoryStub.BASE_COLLECTIONS)
    BASE_COLLECTIONS[1]["body_part_examined"] = ["T-28000"]
    BASE_COLLECTIONS[1]["type"].append("IMAGE")
    BASE_COLLECTIONS[2]["body_part_examined"] = ["T-04000"]
    BASE_COLLECTIONS[2]["type"].append("IMAGE")


def test_directory_stats_matches_exporter_all_active_totals(monkeypatch):
    """Verify directory stats matches exporter all active totals.

    Args:
        monkeypatch: Pytest monkeypatch fixture; temporary dependency and environment overrides are undone after the test.

    Returns:
        None. Verifies directory stats matches exporter all active totals.
    """
    stats_globals, _, _ = _run_script(
        monkeypatch,
        "directory-stats.py",
        ["-N"],
    )
    exporter_globals, _, _ = _run_script(
        monkeypatch,
        "exporter-all.py",
        ["-N"],
    )

    summary = stats_globals["summary"]
    assert summary["withdrawn_biobanks"] == 0
    assert summary["biobanks_total"] == len(exporter_globals["allBiobanks"])
    assert summary["collection_records_total"] == len(exporter_globals["allCollections"])
    assert summary["samples_explicit"] == exporter_globals["allCollectionSamplesExplicit"]
    assert summary["donors_explicit"] == exporter_globals["allCollectionDonorsExplicit"]
    assert summary["samples_total"] == exporter_globals["allCollectionSamplesIncOoM"]
    assert summary["donors_total"] == exporter_globals["allCollectionDonorsIncOoM"]


def test_exporter_all_collects_services_and_studies_in_active_scope(monkeypatch):
    """Verify exporter all collects services and studies in active scope.

    Args:
        monkeypatch: Pytest monkeypatch fixture; temporary dependency and environment overrides are undone after the test.

    Returns:
        None. Verifies exporter all collects services and studies in active scope.
    """
    exporter_globals, _, _ = _run_script(
        monkeypatch,
        "exporter-all.py",
        ["-N"],
    )

    assert [service["id"] for service in exporter_globals["allServices"]] == ["svc1"]
    assert [study["id"] for study in exporter_globals["allStudies"]] == ["study1", "study2"]
    assert [contact["id"] for contact in exporter_globals["allContacts"]] == ["ct1", "ct2"]
    assert [network["id"] for network in exporter_globals["allNetworks"]] == ["net1", "net2"]


def test_exporter_all_can_append_withdrawn_sheets_to_main_workbook(monkeypatch, tmp_path):
    """Verify exporter all can append withdrawn sheets to main workbook.

    Args:
        monkeypatch: Pytest monkeypatch fixture; temporary dependency and environment overrides are undone after the test.
        tmp_path: Pytest-managed temporary filesystem directory for this test case.

    Returns:
        None. Verifies exporter all can append withdrawn sheets to main workbook.
    """
    workbook = tmp_path / "all.xlsx"
    _run_script(
        monkeypatch,
        "exporter-all.py",
        [
            "-N",
            "-w",
            "-X",
            str(workbook),
            "--include-withdrawn-sheets-in-output",
        ],
    )

    import pandas as pd

    sheet_names = pd.ExcelFile(workbook).sheet_names
    assert sheet_names == [
        "Biobanks",
        "Collections",
        "Services",
        "Studies",
        "Contacts",
        "Networks",
        "Fact sheet summary",
        "Fact sheet all-star rows",
        "Fact sheet distributions",
        "Fact sheet margin rows",
        "Withdrawn biobanks",
        "Withdrawn collections",
        "Withdrawn services",
        "Withdrawn studies",
        "Withdrawn contacts",
        "Withdrawn networks",
    ]


def test_exporter_all_writes_clickable_id_hyperlinks(monkeypatch, tmp_path):
    """Verify exporter all writes clickable id hyperlinks.

    Args:
        monkeypatch: Pytest monkeypatch fixture; temporary dependency and environment overrides are undone after the test.
        tmp_path: Pytest-managed temporary filesystem directory for this test case.

    Returns:
        None. Verifies exporter all writes clickable id hyperlinks.
    """
    workbook = tmp_path / "links.xlsx"
    _run_script(
        monkeypatch,
        "exporter-all.py",
        [
            "-N",
            "-X",
            str(workbook),
        ],
    )

    from openpyxl import load_workbook

    wb = load_workbook(workbook)
    biobank_formula = wb["Biobanks"]["B2"].value
    collection_formula = wb["Collections"]["B2"].value
    service_formula = wb["Services"]["B2"].value
    study_formula = wb["Studies"]["B2"].value
    contact_formula = wb["Contacts"]["B2"].value
    network_formula = wb["Networks"]["B2"].value

    assert biobank_formula == '=HYPERLINK("https://directory.example.test/ERIC/directory/#/biobank/bbmri-eric:ID:CZ_BB1","bbmri-eric:ID:CZ_BB1")'
    assert collection_formula == '=HYPERLINK("https://directory.example.test/ERIC/directory/#/collection/col1","col1")'
    assert service_formula == '=HYPERLINK("https://directory.example.test/ERIC/directory/#/service/svc1","svc1")'
    assert study_formula == '=HYPERLINK("https://directory.example.test/ERIC/directory/#/study/study1","study1")'
    assert contact_formula == '=HYPERLINK("https://directory.example.test/ERIC/directory/#/person/ct1","ct1")'
    assert network_formula == '=HYPERLINK("https://directory.example.test/ERIC/directory/#/network/net1","net1")'


def test_exporter_cohorts_only_withdrawn_handles_active_parent_biobank(
    monkeypatch,
    tmp_path,
):
    """Verify exporter cohorts only withdrawn handles active parent biobank.

    Args:
        monkeypatch: Pytest monkeypatch fixture; temporary dependency and environment overrides are undone after the test.
        tmp_path: Pytest-managed temporary filesystem directory for this test case.

    Returns:
        None. Verifies exporter cohorts only withdrawn handles active parent biobank.
    """
    workbook = tmp_path / "cohorts.xlsx"

    exporter_globals, _, _ = _run_script(
        monkeypatch,
        "exporter-cohorts.py",
        ["-N", "--only-withdrawn", "-X", str(workbook)],
        directory_class=WithdrawnCohortUnderActiveBiobankDirectoryStub,
    )

    assert [collection["id"] for collection in exporter_globals["cohortCollections"]] == [
        "col4"
    ]
    assert exporter_globals["cohortBiobankIds"] == {"bbmri-eric:ID:EXT_BB2"}
    assert exporter_globals["cohortCountries"] == {"DE"}


def test_exporter_cohorts_reports_explicit_and_oom_totals(monkeypatch):
    """Verify exporter cohorts reports explicit and oom totals.

    Args:
        monkeypatch: Pytest monkeypatch fixture; temporary dependency and environment overrides are undone after the test.

    Returns:
        None. Verifies exporter cohorts reports explicit and oom totals.
    """
    exporter_globals, stdout, _ = _run_script(
        monkeypatch,
        "exporter-cohorts.py",
        [],
        directory_class=CohortTotalsDirectoryStub,
    )

    assert [collection["id"] for collection in exporter_globals["cohortCollections"]] == [
        "col1",
        "col2",
        "col3",
    ]
    assert exporter_globals["cohortCollectionSamplesExplicit"] == 100
    assert exporter_globals["cohortCollectionDonorsExplicit"] == 10
    assert exporter_globals["cohortCollectionSamplesIncOoM"] == 1100
    assert exporter_globals["cohortCollectionDonorsIncOoM"] == 110
    assert "Fact-sheet summary:" in stdout
    assert "- collections with fact sheets: 3 / 3" in stdout
    assert "sample/donor values are fact-row observations" not in stdout
    assert (
        "- all-star totals for collections with populated all-but-one-star rows: "
        "0 samples / 0 donors (from 0 collections with one populated all-star row)"
    ) in stdout
    assert (
        "- total of samples/donors advertised explicitly in cohort collections: "
        "100 / 10"
    ) in stdout
    assert (
        "- total of samples/donors advertised in cohort collections including "
        "OoM estimates: 1100 / 110"
    ) in stdout


def test_exporter_cohorts_writes_fact_sheet_summary_sheets(monkeypatch, tmp_path):
    """Verify exporter cohorts writes fact sheet summary sheets.

    Args:
        monkeypatch: Pytest monkeypatch fixture; temporary dependency and environment overrides are undone after the test.
        tmp_path: Pytest-managed temporary filesystem directory for this test case.

    Returns:
        None. Verifies exporter cohorts writes fact sheet summary sheets.
    """
    workbook = tmp_path / "cohorts.xlsx"

    _run_script(
        monkeypatch,
        "exporter-cohorts.py",
        ["-N", "-X", str(workbook)],
        directory_class=CohortTotalsDirectoryStub,
    )

    import pandas as pd

    sheet_names = pd.ExcelFile(workbook).sheet_names
    assert "Fact sheet summary" in sheet_names
    assert "Fact sheet all-star rows" in sheet_names
    assert "Fact sheet distributions" in sheet_names
    summary = pd.read_excel(workbook, sheet_name="Fact sheet summary").iloc[0]
    assert summary["collections"] == 3
    assert summary["collections_with_fact_sheets"] == 3
    assert summary["populated_all_star_rows"] == 3


def test_exporter_cohorts_no_star_fallback_is_warned_and_separate(
    monkeypatch,
    tmp_path,
):
    """Verify exporter cohorts no star fallback is warned and separate.

    Args:
        monkeypatch: Pytest monkeypatch fixture; temporary dependency and environment overrides are undone after the test.
        tmp_path: Pytest-managed temporary filesystem directory for this test case.

    Returns:
        None. Verifies exporter cohorts no star fallback is warned and separate.
    """
    workbook = tmp_path / "cohorts-fallback.xlsx"

    _, stdout, _ = _run_script(
        monkeypatch,
        "exporter-cohorts.py",
        ["--allow-no-star-fact-sums", "-X", str(workbook)],
        directory_class=CohortNoStarFallbackDirectoryStub,
    )

    assert "WARNING: No-star fact-sheet fallback is enabled" in stdout
    assert "UNSAFE no-star fallback distributions" in stdout

    import pandas as pd

    sheet_names = pd.ExcelFile(workbook).sheet_names
    assert "Fact sheet no-star fallback" in sheet_names
    fallback = pd.read_excel(workbook, sheet_name="Fact sheet no-star fallback")
    female = fallback[(fallback["dimension"] == "sex") & (fallback["value_id"] == "FEMALE")].iloc[0]
    assert female["number_of_samples"] == 20
    assert female["number_of_donors"] == 8
    assert female["no_star_fallback_collections"] == 1
    assert bool(female["assumption_violating"]) is True

    authoritative = pd.read_excel(workbook, sheet_name="Fact sheet distributions")
    assert authoritative.empty


def test_exporter_cmdr_lists_biobanks_collections_and_studies(monkeypatch):
    """Verify exporter cmdr lists biobanks collections and studies.

    Args:
        monkeypatch: Pytest monkeypatch fixture; temporary dependency and environment overrides are undone after the test.

    Returns:
        None. Verifies exporter cmdr lists biobanks collections and studies.
    """
    exporter_globals, stdout, _ = _run_script(
        monkeypatch,
        "exporter-cMDR.py",
        [],
    )

    assert [biobank["id"] for biobank in exporter_globals["cmdrBiobanks"]] == [
        "bbmri-eric:ID:CZ_BB1",
        "bbmri-eric:ID:EXT_BB2",
    ]
    assert [collection["id"] for collection in exporter_globals["cmdrCollections"]] == [
        "col1",
        "col3",
    ]
    assert [study["id"] for study in exporter_globals["cmdrStudies"]] == [
        "study1",
        "study2",
    ]
    assert "CZ" in stdout
    assert "bbmri-eric:ID:CZ_BB1 - Biobank 1 [studies: study1,study2]" in stdout
    assert "col1 - Collection 1 [studies: study1,study2]" in stdout
    assert "DE" in stdout
    assert "bbmri-eric:ID:EXT_BB2 - Biobank 2 [studies: study2]" in stdout
    assert "study2 - Study 2 [collections: col1,col3]" in stdout
    assert "Per-country summary:" in stdout
    assert "- CZ: biobanks linked to studies = 1, collections linked to studies = 1, studies linked to collections = 2" in stdout
    assert "- DE: biobanks linked to studies = 1, collections linked to studies = 1, studies linked to collections = 1" in stdout
    assert "- CZ,DE:" not in stdout


def test_exporter_cmdr_writes_sorted_hyperlinked_workbook(monkeypatch, tmp_path):
    """Verify exporter cmdr writes sorted hyperlinked workbook.

    Args:
        monkeypatch: Pytest monkeypatch fixture; temporary dependency and environment overrides are undone after the test.
        tmp_path: Pytest-managed temporary filesystem directory for this test case.

    Returns:
        None. Verifies exporter cmdr writes sorted hyperlinked workbook.
    """
    workbook = tmp_path / "cmdr.xlsx"
    _run_script(
        monkeypatch,
        "exporter-cMDR.py",
        [
            "-N",
            "-X",
            str(workbook),
        ],
    )

    from openpyxl import load_workbook

    wb = load_workbook(workbook)
    assert wb.sheetnames == [
        "Biobanks",
        "Collections",
        "Studies",
        "Fact sheet summary",
        "Fact sheet all-star rows",
        "Fact sheet distributions",
        "Fact sheet margin rows",
    ]
    assert wb["Biobanks"]["A2"].value == "CZ"
    assert wb["Biobanks"]["B2"].value == '=HYPERLINK("https://directory.example.test/ERIC/directory/#/biobank/bbmri-eric:ID:CZ_BB1","bbmri-eric:ID:CZ_BB1")'
    assert wb["Biobanks"]["B3"].value == '=HYPERLINK("https://directory.example.test/ERIC/directory/#/biobank/bbmri-eric:ID:EXT_BB2","bbmri-eric:ID:EXT_BB2")'
    assert wb["Collections"]["A2"].value == "CZ"
    assert wb["Collections"]["B2"].value == '=HYPERLINK("https://directory.example.test/ERIC/directory/#/collection/col1","col1")'
    assert wb["Collections"]["B3"].value == '=HYPERLINK("https://directory.example.test/ERIC/directory/#/collection/col3","col3")'
    assert wb["Studies"]["A2"].value == "CZ"
    assert wb["Studies"]["B2"].value == '=HYPERLINK("https://directory.example.test/ERIC/directory/#/study/study1","study1")'
    assert wb["Studies"]["I2"].value == '=HYPERLINK("https://directory.example.test/ERIC/directory/#/biobank/bbmri-eric:ID:CZ_BB1","bbmri-eric:ID:CZ_BB1")'
    assert wb["Studies"]["M2"].value == '=HYPERLINK("https://directory.example.test/ERIC/directory/#/collection/col1","col1")'
    assert wb["Studies"]["A3"].value == "CZ,DE"
    assert wb["Studies"]["I3"].value == '=HYPERLINK("https://directory.example.test/ERIC/directory/#/biobank/bbmri-eric:ID:CZ_BB1","bbmri-eric:ID:CZ_BB1")'
    assert wb["Studies"]["K3"].value == '=HYPERLINK("https://directory.example.test/ERIC/directory/#/biobank/bbmri-eric:ID:EXT_BB2","bbmri-eric:ID:EXT_BB2")'
    assert wb["Studies"]["M3"].value == '=HYPERLINK("https://directory.example.test/ERIC/directory/#/collection/col1","col1")'


def test_exporter_cmdr_writes_geojson_with_entity_and_biobank_fallback(monkeypatch, tmp_path):
    """Verify exporter cmdr writes geojson with entity and biobank fallback.

    Args:
        monkeypatch: Pytest monkeypatch fixture; temporary dependency and environment overrides are undone after the test.
        tmp_path: Pytest-managed temporary filesystem directory for this test case.

    Returns:
        None. Verifies exporter cmdr writes geojson with entity and biobank fallback.
    """
    output_file = tmp_path / "cmdr.geojson"
    _run_script(
        monkeypatch,
        "exporter-cMDR.py",
        [
            "-N",
            "-G",
            str(output_file),
        ],
    )

    payload = json.loads(output_file.read_text(encoding="utf-8"))
    assert payload["type"] == "FeatureCollection"
    by_key = {
        (feature["properties"]["entity_type"], feature["properties"]["id"]): feature
        for feature in payload["features"]
    }

    collection_feature = by_key[("collection", "col1")]
    assert collection_feature["geometry"]["coordinates"] == [14.5, 50.1]
    assert collection_feature["properties"]["coordinate_source"] == "collection"

    collection_feature_fallback = by_key[("collection", "col3")]
    assert collection_feature_fallback["geometry"]["coordinates"] == [13.405, 52.52]
    assert collection_feature_fallback["properties"]["coordinate_source"] == "biobank"

    study_feature = by_key[("study", "study1")]
    assert study_feature["geometry"]["coordinates"] == [14.6, 50.2]
    assert study_feature["properties"]["coordinate_source"] == "study"

    study_feature_fallback = by_key[("study", "study2")]
    assert study_feature_fallback["geometry"]["coordinates"] == [14.5, 50.1]
    assert study_feature_fallback["properties"]["coordinate_source"] == "collection"
    assert study_feature_fallback["properties"]["coordinate_source_id"] == "col1"


def test_directory_stats_can_include_withdrawn_biobanks(monkeypatch):
    """Verify directory stats can include withdrawn biobanks.

    Args:
        monkeypatch: Pytest monkeypatch fixture; temporary dependency and environment overrides are undone after the test.

    Returns:
        None. Verifies directory stats can include withdrawn biobanks.
    """
    default_globals, _, _ = _run_script(
        monkeypatch,
        "directory-stats.py",
        ["-N"],
    )
    include_globals, _, _ = _run_script(
        monkeypatch,
        "directory-stats.py",
        ["-N", "-w"],
    )

    default_summary = default_globals["summary"]
    include_summary = include_globals["summary"]

    assert default_summary["biobanks_total"] == 2
    assert include_summary["biobanks_total"] == 3
    assert default_summary["withdrawn_biobanks"] == 0
    assert include_summary["withdrawn_biobanks"] == 1
    assert include_summary["collection_records_total"] == default_summary["collection_records_total"] + 2
    assert include_summary["samples_explicit"] == default_summary["samples_explicit"] + 12
    assert include_summary["donors_explicit"] == default_summary["donors_explicit"] + 8
    assert include_summary["services_total"] == default_summary["services_total"] + 1


def test_directory_stats_can_select_only_withdrawn_biobanks(monkeypatch):
    """Verify directory stats can select only withdrawn biobanks.

    Args:
        monkeypatch: Pytest monkeypatch fixture; temporary dependency and environment overrides are undone after the test.

    Returns:
        None. Verifies directory stats can select only withdrawn biobanks.
    """
    only_globals, _, _ = _run_script(
        monkeypatch,
        "directory-stats.py",
        ["-N", "--only-withdrawn"],
    )

    only_summary = only_globals["summary"]

    assert only_summary["biobanks_total"] == 1
    assert only_summary["withdrawn_biobanks"] == 1
    assert only_summary["collection_records_total"] == 1


def test_directory_stats_matches_exporter_all_when_oom_policy_changes(
    monkeypatch,
):
    """Verify directory stats matches exporter all when oom policy changes.

    Args:
        monkeypatch: Pytest monkeypatch fixture; temporary dependency and environment overrides are undone after the test.

    Returns:
        None. Verifies directory stats matches exporter all when oom policy changes.
    """
    monkeypatch.setenv("DIRECTORY_OOM_UPPER_BOUND_COEFFICIENT", "0.3")

    stats_globals, _, _ = _run_script(
        monkeypatch,
        "directory-stats.py",
        ["-N"],
    )
    exporter_globals, _, _ = _run_script(
        monkeypatch,
        "exporter-all.py",
        ["-N"],
    )

    summary = stats_globals["summary"]
    assert summary["oom_upper_bound_coefficient"] == 0.3
    assert summary["samples_total"] == exporter_globals["allCollectionSamplesIncOoM"]
    assert summary["donors_total"] == exporter_globals["allCollectionDonorsIncOoM"]


def test_directory_stats_script_applies_country_and_staging_area_filters(monkeypatch):
    """Verify directory stats script applies country and staging area filters.

    Args:
        monkeypatch: Pytest monkeypatch fixture; temporary dependency and environment overrides are undone after the test.

    Returns:
        None. Verifies directory stats script applies country and staging area filters.
    """
    globals_dict, _, _ = _run_script(
        monkeypatch,
        "directory-stats.py",
        ["-N", "-c", "DE", "-A", "EXT"],
    )

    stats_df = globals_dict["stats_df"]
    summary = globals_dict["summary"]

    assert stats_df["id"].tolist() == ["bbmri-eric:ID:EXT_BB2"]
    assert summary["biobanks_total"] == 1
    assert summary["country_filter"] == "DE"
    assert summary["staging_area_filter"] == "EXT"


def test_directory_stats_script_supports_comma_delimited_filters_and_collection_types(
    monkeypatch,
):
    """Verify directory stats script supports comma delimited filters and collection types.

    Args:
        monkeypatch: Pytest monkeypatch fixture; temporary dependency and environment overrides are undone after the test.

    Returns:
        None. Verifies directory stats script supports comma delimited filters and collection types.
    """
    globals_dict, _, _ = _run_script(
        monkeypatch,
        "directory-stats.py",
        ["-N", "-c", "CZ,DE", "-T", "CASE_CONTROL,POPULATION"],
    )

    stats_df = globals_dict["stats_df"]
    summary = globals_dict["summary"]

    assert stats_df["id"].tolist() == [
        "bbmri-eric:ID:CZ_BB1",
        "bbmri-eric:ID:EXT_BB2",
    ]
    assert summary["country_filter"] == "CZ,DE"
    assert summary["collection_type_filter"] == "CASE_CONTROL,POPULATION"
    assert summary["collection_records_total"] == 2


def test_fact_sheet_emulation_exporter_writes_complete_hyperlinked_workbook(
    monkeypatch,
    tmp_path,
):
    """Verify fact sheet emulation exporter writes complete hyperlinked workbook.

    Args:
        monkeypatch: Pytest monkeypatch fixture; temporary dependency and environment overrides are undone after the test.
        tmp_path: Pytest-managed temporary filesystem directory for this test case.

    Returns:
        None. Verifies fact sheet emulation exporter writes complete hyperlinked workbook.
    """
    workbook = tmp_path / "emulation.xlsx"

    _, stdout, _ = _run_script(
        monkeypatch,
        "exporter-fact-sheet-emulation.py",
        ["-X", str(workbook)],
        directory_class=FactSheetEmulationDirectoryStub,
    )

    from openpyxl import load_workbook

    wb = load_workbook(workbook, data_only=False)
    assert wb.sheetnames == [
        "Candidate families",
        "Source collections",
        "Proposed facts",
        "Unrepresentable data",
        "Migration mapping",
        "Dimension candidates",
        "Dimension values",
    ]
    assert "Field comparison" not in wb.sheetnames
    assert "Boundary evidence" not in wb.sheetnames

    source_sheet = wb["Source collections"]
    headers = [cell.value for cell in source_sheet[1]]
    collection_column = headers.index("collection_id") + 1
    assert source_sheet.cell(2, collection_column).value.startswith('=HYPERLINK("')

    proposed_sheet = wb["Proposed facts"]
    proposed_headers = [cell.value for cell in proposed_sheet[1]]
    proposed_rows = [
        dict(zip(proposed_headers, row))
        for row in proposed_sheet.iter_rows(min_row=2, values_only=True)
    ]
    assert len(proposed_rows) == 2
    assert {row["row_kind"] for row in proposed_rows} == {"all_but_one_star"}
    assert {row["sample_type"] for row in proposed_rows} == {"SERUM", "DNA"}
    assert {row["sex"] for row in proposed_rows} == {"*"}
    assert {row["age_range"] for row in proposed_rows} == {"*"}
    assert {row["disease"] for row in proposed_rows} == {"*"}
    assert {row["number_of_samples"] for row in proposed_rows} == {60, 40}
    assert {row["number_of_donors"] for row in proposed_rows} == {40, 30}

    candidate_headers = [cell.value for cell in wb["Candidate families"][1]]
    assert {
        "discovery_rule",
        "description_classification",
        "operational_boundary_categories",
        "abstention_reason",
    }.issubset(candidate_headers)
    assert "source_collection_ids" not in [
        cell.value for cell in wb["Candidate families"][1]
    ]
    assert "Per-country summary:" in stdout
    assert "- CZ: families = 1, source collections = 2" in stdout
    assert "Per-biobank summary:" in stdout
    assert "- CZ / bbmri-eric:ID:CZ_BB1: families = 1, source collections = 2" in stdout
    assert "ready_current_fact_schema = 1" in stdout


def test_fact_sheet_emulation_exporter_adds_diagnostics_only_to_advanced_workbook(
    monkeypatch,
    tmp_path,
):
    """Verify fact sheet emulation exporter adds diagnostics only to advanced workbook.

    Args:
        monkeypatch: Pytest monkeypatch fixture; temporary dependency and environment overrides are undone after the test.
        tmp_path: Pytest-managed temporary filesystem directory for this test case.

    Returns:
        None. Verifies fact sheet emulation exporter adds diagnostics only to advanced workbook.
    """
    workbook = tmp_path / "emulation-advanced.xlsx"

    _run_script(
        monkeypatch,
        "exporter-fact-sheet-emulation.py",
        ["-N", "--advanced-reporting", "-X", str(workbook)],
        directory_class=FactSheetEmulationDirectoryStub,
    )

    from openpyxl import load_workbook

    wb = load_workbook(workbook, data_only=False)
    assert "Field comparison" in wb.sheetnames
    assert "Boundary evidence" in wb.sheetnames

    comparison_headers = [cell.value for cell in wb["Field comparison"][1]]
    assert "distinct_value_count" in comparison_headers
    assert "distinct_values" not in comparison_headers
    assert "missing_collection_count" in comparison_headers
    assert "missing_collection_ids" not in comparison_headers
    comparison_rows = [
        dict(zip(comparison_headers, row))
        for row in wb["Field comparison"].iter_rows(min_row=2, values_only=True)
    ]
    material_rows = [
        row
        for row in comparison_rows
        if row["field"] == "materials" and row["role"] == "characterization"
    ]
    assert len(material_rows) == 2
    assert {row["distinct_value_count"] for row in material_rows} == {2}
    assert {row["missing_collection_count"] for row in material_rows} == {0}
    purpose_rows = [
        row
        for row in comparison_rows
        if row["field"] == "purpose" and row["role"] == "operational"
    ]
    assert len(purpose_rows) == 2
    assert {row["distinct_value_count"] for row in purpose_rows} == {0}
    assert {row["missing_collection_count"] for row in purpose_rows} == {2}


def test_fact_sheet_emulation_exporter_writes_matching_external_review_packets(
    monkeypatch,
    tmp_path,
):
    """Verify fact sheet emulation exporter writes matching external review packets.

    Args:
        monkeypatch: Pytest monkeypatch fixture; temporary dependency and environment overrides are undone after the test.
        tmp_path: Pytest-managed temporary filesystem directory for this test case.

    Returns:
        None. Verifies fact sheet emulation exporter writes matching external review packets.
    """
    prefix = tmp_path / "emulation"

    _run_script(
        monkeypatch,
        "exporter-fact-sheet-emulation.py",
        ["-N", "--ai-review-prefix", str(prefix)],
        directory_class=FactSheetEmulationReviewDirectoryStub,
    )

    json_path = tmp_path / "emulation-ai-review.json"
    markdown_path = tmp_path / "emulation-ai-review.md"
    packet = json.loads(json_path.read_text(encoding="utf-8"))
    markdown = markdown_path.read_text(encoding="utf-8")

    assert packet["schema_version"] == "1.2"
    assert len(packet["cases"]) == 1
    assert packet["cases"][0]["family"]["country"] == "CZ"
    assert "boundary_evidence" in packet["cases"][0]
    assert "field_summary" in packet["cases"][0]
    assert "anatomical_site" in markdown
    assert "Do not sum" in markdown
