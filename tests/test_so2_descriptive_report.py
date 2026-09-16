"""Tests for the self-contained SO2 descriptive-workbook reader."""

import importlib
from pathlib import Path

import openpyxl
import pytest


module = importlib.import_module("so2_descriptive_report")
WORKTREE = Path(__file__).resolve().parents[1]
PRODUCTION_SCHEMA = WORKTREE / "survey-mappings" / "so2_2025_descriptive_report.json"
PRODUCTION_WORKBOOK = Path("/storage/emulated/0/BBMRI-ERIC/directory-scripts/Content_Export_SO2_2025_20260313.xlsx")
LEGAL_GDPR = "What types of legal barriers have you faced?: Data protection regulations (e.g., GDPR)"
LEGAL_LICENSING = "What types of legal barriers have you faced?: Licensing restrictions"
LEGAL_OWNERSHIP = "What types of legal barriers have you faced?: Lack of clarity on data ownership"
LEGAL_CROSS_BORDER = "What types of legal barriers have you faced?: Cross-border data sharing restrictions"
OTHER_LEGAL = "If there are other legal barriers, please specify:"
OTHER_ORGANISATIONAL = "If there are other organisational barriers, please specify:"
RETURNED_DATA_STORAGE = "How much disk space do you currently require storing returned data overall (for all returned data altogether)?"
ORDINAL_BARRIER_CATEGORIES = ["never", "rarely", "occasionally", "frequently"]
TRACEABILITY_READINESS = (
    "How would you rate your repository’s technical readiness for integrating full digital "
    "traceability? Please choose what corresponds the most to your current situation. By “digital "
    "traceability,” we mean the structured, automated, and system-supported capture and linkage of "
    "associated data (e.g. preanalytical, clinical, research-derived), enabling traceability of a "
    "sample across its full lifecycle, from collection through processing, storage, use, and data "
    "generation. This includes the use of dedicated software (e.g. BIMS/LIMS), digital metadata "
    "standards, audit trails, and secure linkage to clinical data or research data."
)


HEADERS = [
    "Name of Institution",
    "Country",
    "Creation date",
    "Last update",
    "Digital maturity",
]


def minimal_schema():
    """Return a complete schema for the controlled workbook fixture."""
    return {
        "schema_version": "1",
        "input": {
            "worksheet": "Survey responses",
            "alias": "SO2_2025",
            "header_row": 4,
        },
        "columns": {
            "respondent_context": ["Name of Institution", "Country"],
            "administrative_exclusions": ["Creation date", "Last update"],
        },
        "questions": [
            {
                "question_id": "digital_maturity",
                "column": "Digital maturity",
                "question_type": "ordinal",
                "label": "Digital maturity",
                "categories": ["Low", "High"],
            }
        ],
        "output": {"source_row_column": "source_row"},
    }


def write_descriptive_workbook(tmp_path, rows=(), mutator=None):
    """Write a workbook matching the survey service envelope."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Survey responses"
    sheet.cell(1, 1, "SO2_2025")
    sheet.cell(2, 1, "2026-03-13T07:22:45")
    for column_number, header in enumerate(HEADERS, start=1):
        sheet.cell(4, column_number, header)
    for row_number, row in enumerate(rows, start=5):
        for column_number, header in enumerate(HEADERS, start=1):
            sheet.cell(row_number, column_number, row.get(header))
    if mutator is not None:
        mutator(sheet)
    path = tmp_path / "descriptive.xlsx"
    workbook.save(path)
    return path


def wrong_alias_row(sheet):
    """Break the row-1 alias envelope field."""
    sheet.cell(1, 1, "WRONG")


def missing_export_date(sheet):
    """Break the row-2 export-date envelope field."""
    sheet.cell(2, 1).value = None


def invalid_export_date(sheet):
    """Break the ISO-like row-2 export-date envelope field."""
    sheet.cell(2, 1, "not-a-date")


def nonblank_separator(sheet):
    """Break the required blank separator row."""
    sheet.cell(3, 1, "not blank")


def wrong_header(sheet):
    """Break the first cell of the row-4 header row."""
    sheet.cell(4, 1, "Institution")


def test_reader_preserves_service_metadata_and_only_response_rows(tmp_path):
    """Reader preserves service provenance and excludes blank data records."""
    workbook = write_descriptive_workbook(
        tmp_path,
        rows=[
            {"Name of Institution": "Alpha", "Country": "Austria"},
            {"Creation date": "2026-01-01", "Last update": "2026-01-02"},
            {},
        ],
    )

    result = module.read_descriptive_workbook(workbook, minimal_schema())

    assert result.alias == "SO2_2025"
    assert result.export_date == "2026-03-13T07:22:45"
    assert result.header_row == 4
    assert result.worksheet == "Survey responses"
    assert result.responses["Name of Institution"].tolist() == ["Alpha"]
    assert result.responses["source_row"].tolist() == [5]
    assert result.total_data_rows == 3
    assert result.excluded_blank_rows == 2
    assert len(result.source_sha256) == 64


@pytest.mark.parametrize(
    ("mutator", "message"),
    [
        (wrong_alias_row, "row 1"),
        (missing_export_date, "row 2"),
        (invalid_export_date, "row 2"),
        (nonblank_separator, "row 3"),
        (wrong_header, "header row 4"),
    ],
)
def test_reader_rejects_invalid_service_envelope(tmp_path, mutator, message):
    """Reader rejects each service-envelope row when it is malformed."""
    with pytest.raises(module.InputError, match=message):
        module.read_descriptive_workbook(
            write_descriptive_workbook(tmp_path, mutator=mutator), minimal_schema()
        )


def test_reader_converts_corrupt_xlsx_to_input_error(tmp_path):
    """Reader reports a non-ZIP XLSX file as a user-facing input error."""
    workbook = tmp_path / "corrupt.xlsx"
    workbook.write_bytes(b"not a ZIP workbook")

    with pytest.raises(module.InputError, match="Could not open descriptive workbook"):
        module.read_descriptive_workbook(workbook, minimal_schema())


def test_reader_rejects_missing_schema_fields(tmp_path):
    """Reader reports incomplete schemas as actionable input errors."""
    schema = minimal_schema()
    del schema["input"]

    with pytest.raises(module.InputError, match="input"):
        module.read_descriptive_workbook(write_descriptive_workbook(tmp_path), schema)


def test_schema_rejects_missing_or_unclassified_columns():
    """Schema cannot omit a source column or declare one that is absent."""
    with pytest.raises(module.InputError, match="unclassified source column"):
        module.validate_descriptive_schema(
            minimal_schema(), HEADERS + ["Extra"]
        )
    with pytest.raises(module.InputError, match="declared column is absent"):
        module.validate_descriptive_schema(
            minimal_schema(), ["Name of Institution", "Country"]
        )


def test_schema_rejects_duplicate_classification_and_invalid_question_values():
    """Schema requires one classification per column and valid question fields."""
    schema = minimal_schema()
    schema["columns"]["respondent_context"].append("Digital maturity")
    with pytest.raises(module.InputError, match="classified more than once"):
        module.validate_descriptive_schema(schema, HEADERS)

    schema = minimal_schema()
    schema["questions"][0]["question_type"] = "numeric"
    with pytest.raises(module.InputError, match="question_type"):
        module.validate_descriptive_schema(schema, HEADERS)


def test_schema_rejects_absent_parent_column():
    """Question parent references must name an existing source column."""
    schema = minimal_schema()
    schema["questions"][0]["parent_columns"] = ["Missing parent"]

    with pytest.raises(module.InputError, match="declared column is absent"):
        module.validate_descriptive_schema(schema, HEADERS)


def test_load_schema_rejects_duplicate_json_keys_and_invalid_output(tmp_path):
    """Schema loading rejects ambiguous JSON and unusable output provenance."""
    duplicate_keys = tmp_path / "duplicate-keys.json"
    duplicate_keys.write_text('{"schema_version": "1", "schema_version": "1"}')
    with pytest.raises(module.InputError, match="duplicate JSON key"):
        module.load_descriptive_schema(duplicate_keys)

    schema = minimal_schema()
    schema["output"] = {"source_row_column": "Name of Institution"}
    with pytest.raises(module.InputError, match="output.source_row_column"):
        module.validate_descriptive_schema(schema, HEADERS)


def test_schema_returns_frozen_question_definitions():
    """Validated questions expose the stable immutable public data class."""
    questions = module.validate_descriptive_schema(minimal_schema(), HEADERS)

    assert questions[0] == module.QuestionDefinition(
        question_id="digital_maturity",
        column="Digital maturity",
        question_type="ordinal",
        label="Digital maturity",
        categories=("Low", "High"),
        delimiter=None,
        parent_columns=(),
        applicability=None,
    )
    with pytest.raises(AttributeError):
        questions[0].label = "Changed"


def test_nonblank_response_row_ignores_administrative_values():
    """Creation and update timestamps alone never create a response record."""
    pandas = pytest.importorskip("pandas")
    row = pandas.Series(
        {
            "Name of Institution": " ",
            "Country": None,
            "Creation date": "2026-01-01",
            "Last update": "2026-01-02",
            "Digital maturity": "\t",
        }
    )

    assert not module.is_nonblank_response_row(
        row, ["Name of Institution", "Country", "Digital maturity"]
    )


def read_row_four_headers(workbook_path):
    """Return the exact production workbook row-4 headers."""
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        return [cell.value for cell in workbook.active[4]]
    finally:
        workbook.close()


def questions_by_column(schema):
    """Index raw schema questions by their explicit source header."""
    return {question["column"]: question for question in schema["questions"]}


def test_production_schema_accounts_for_every_header():
    """The production registry classifies each current source header exactly once."""
    schema = module.load_descriptive_schema(PRODUCTION_SCHEMA)
    headers = read_row_four_headers(PRODUCTION_WORKBOOK)
    module.validate_descriptive_schema(schema, headers)
    classified = set(schema["columns"]["respondent_context"])
    classified |= set(schema["columns"]["administrative_exclusions"])
    classified |= {item["column"] for item in schema["questions"]}

    assert classified == set(headers)
    assert len(classified) == len(headers)


def test_production_schema_overlapping_barrier_columns_are_ordinal_with_parents():
    """Explicit ordinal precedence retains matrix context for both other-barrier columns."""
    questions = questions_by_column(module.load_descriptive_schema(PRODUCTION_SCHEMA))

    assert questions[LEGAL_GDPR]["question_type"] == "ordinal"
    assert questions[OTHER_LEGAL]["question_type"] == "ordinal"
    assert questions[OTHER_LEGAL]["categories"] == ORDINAL_BARRIER_CATEGORIES
    assert questions[OTHER_LEGAL]["parent_columns"] == [
        LEGAL_GDPR,
        LEGAL_LICENSING,
        LEGAL_OWNERSHIP,
        LEGAL_CROSS_BORDER,
    ]
    assert questions[OTHER_ORGANISATIONAL]["question_type"] == "ordinal"
    assert questions[OTHER_ORGANISATIONAL]["categories"] == ORDINAL_BARRIER_CATEGORIES
    assert questions[OTHER_ORGANISATIONAL]["parent_columns"] == [
        "What organizational barriers have you encountered? (Select all that apply): "
        "Lack of data sharing agreements",
        "What organizational barriers have you encountered? (Select all that apply): "
        "Bureaucratic delays",
        "What organizational barriers have you encountered? (Select all that apply): "
        "Lack of contact points or unclear responsibilities",
        "What organizational barriers have you encountered? (Select all that apply): "
        "Institutional reluctance to share data",
    ]


def test_production_schema_preserves_declared_zero_count_ordinal_category_order():
    """Declared ordinal storage bands remain ordered even when no response selects some bands."""
    question = questions_by_column(module.load_descriptive_schema(PRODUCTION_SCHEMA))[RETURNED_DATA_STORAGE]

    assert question["question_type"] == "ordinal"
    assert question["categories"] == [
        "I don't know",
        "<1TB",
        "1TB-10TB",
        "10TB-100TB",
        "100TB-1PB",
        ">1PB",
    ]


def test_production_schema_semicolon_inside_declared_ordinal_is_not_split():
    """An ordinal label containing semicolons is not treated as a multi-choice answer."""
    question = questions_by_column(module.load_descriptive_schema(PRODUCTION_SCHEMA))[TRACEABILITY_READINESS]

    assert question["question_type"] == "ordinal"
    assert question.get("delimiter") is None


def test_production_schema_declares_every_observed_non_other_structured_value():
    """Ordinary structured responses have an explicit declared category."""
    schema = module.load_descriptive_schema(PRODUCTION_SCHEMA)
    questions = questions_by_column(schema)
    workbook = openpyxl.load_workbook(PRODUCTION_WORKBOOK, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(min_row=4, values_only=True)
        headers = list(next(rows))
        observed = [set() for _ in headers]
        for row in rows:
            for index, value in enumerate(row):
                if value not in (None, ""):
                    observed[index].add(value)
    finally:
        workbook.close()
    for index, header in enumerate(headers):
        question = questions.get(header)
        if question is None or question["question_type"] == "free_text" or question.get("parent_columns"):
            continue
        values = observed[index]
        if question["question_type"] == "multi_choice":
            values = {choice for value in values for choice in value.split(question["delimiter"])}
        assert values <= set(question["categories"]), header
