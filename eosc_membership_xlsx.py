"""Read EOSC-A membership workbooks and write values-only match workbooks.

The helpers in this module deliberately do not perform matching. They keep
the source worksheet's values and structure available to the matching and CLI
layers while making the generated workbook safe to publish as a new file.
"""

from __future__ import annotations

import ctypes
import errno
import hashlib
import logging
import os
import tempfile
import unicodedata
from pathlib import Path
from typing import Any, Iterable, Mapping

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


log = logging.getLogger(__name__)

SUMMARY_SHEET_NAME = "Matched institutions"
CONTRIBUTOR_TAG = "EOSC Node BBMRI-ERIC"
SUMMARY_HEADERS = (
    "Organisation ID",
    "EOSC-A Name",
    "BBMRI-ERIC Name",
    "List of biobankIDs",
)
EXCEL_MAX_CELL_CHARS = 32767
_EXCEL_ERROR_VALUES = frozenset(
    {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#GETTING_DATA"}
)


def _normalise_heading(value: Any) -> str:
    """Return a case-insensitive heading key without spacing or punctuation."""
    if not isinstance(value, str):
        return ""
    value = unicodedata.normalize("NFKC", value).casefold()
    return "".join(character for character in value if character.isalnum())


def _is_blank(value: Any) -> bool:
    """Return whether a source value is empty or whitespace-only."""
    return value is None or (isinstance(value, str) and not value.strip())


def _is_excel_error(value: Any) -> bool:
    """Return whether a value is an Excel error literal, not identity data."""
    return isinstance(value, str) and value.strip().upper() in _EXCEL_ERROR_VALUES


def _is_external(value: Any) -> bool:
    """Return whether a source identifier marks an external row."""
    return isinstance(value, str) and value.strip().casefold() == "external"


def _identifier_text(value: Any, *, row_number: int) -> str:
    """Convert a non-empty spreadsheet identifier to its stable text form."""
    if isinstance(value, bool):
        raise ValueError(f"Invalid Organisation ID at row {row_number}: boolean value")
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not value.is_integer():
            raise ValueError(
                f"Invalid Organisation ID at row {row_number}: non-integral number"
            )
        return str(int(value))
    if isinstance(value, str):
        identifier = value.strip()
        if identifier:
            return identifier
    raise ValueError(
        f"Invalid Organisation ID at row {row_number}: expected text or number"
    )


def _text_value(value: Any) -> str:
    """Return a required organisation field as text without normalising it."""
    if isinstance(value, str):
        return value
    return str(value)


def _heading_role(value: Any) -> str | None:
    """Classify a supported source heading, if it is recognisable."""
    key = _normalise_heading(value)
    if key in {"organisationid", "organizationid"}:
        return "organisation_id"
    if (
        key in {"name", "organisationname", "organizationname"}
        or key.endswith("organisationname")
        or key.endswith("organizationname")
    ):
        return "name"
    if key in {
        "acronym",
        "orgacronym",
        "organisationacronym",
        "organizationacronym",
    } or key.endswith("acronym"):
        return "acronym"
    if key == "country" or key.endswith("country"):
        return "country"
    if (
        key in {"type", "membershiptype", "membershipstatustype"}
        or key.endswith("membershiptype")
        or key.endswith("membershipstatustype")
    ):
        return "membership_type"
    if (
        key in {"status", "membershipstatus"}
        or key.endswith("membershipstatus")
    ):
        return "membership_status"
    if key in {"nodecontributor", "contributor"} or key.endswith("nodecontributor"):
        return "contributor"
    return None


def _select_sheet(workbook, *, sheet: str | None, sheet_index: int | None):
    """Select one worksheet using the public one-based selection contract."""
    if sheet is not None and sheet_index is not None:
        raise ValueError("sheet and sheet_index are mutually exclusive")
    if not workbook.worksheets:
        raise ValueError("Workbook contains no worksheets")
    if sheet is not None:
        try:
            return workbook[sheet]
        except KeyError as error:
            raise ValueError(f"Worksheet not found: {sheet!r}") from error
    if sheet_index is None:
        return workbook.worksheets[0]
    if not isinstance(sheet_index, int) or isinstance(sheet_index, bool) or sheet_index < 1:
        raise ValueError("sheet_index must be a positive one-based integer")
    if sheet_index > len(workbook.worksheets):
        raise ValueError(
            f"sheet_index {sheet_index} is outside the workbook's 1..{len(workbook.worksheets)} range"
        )
    return workbook.worksheets[sheet_index - 1]


def _workbook_sha256(path: Path) -> str:
    """Return the SHA-256 digest of a workbook file."""
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _formula_coordinates(worksheet) -> set[tuple[int, int]]:
    """Return one-based coordinates of occupied formula cells."""
    return {
        (cell.row, cell.column)
        for row in worksheet.iter_rows()
        for cell in row
        if cell.data_type == "f"
    }


def read_membership(
    path: str | os.PathLike[str], *, sheet: str | None = None, sheet_index: int | None = None
) -> dict[str, Any]:
    """Read one EOSC-A membership worksheet into a validated literal snapshot.

    Args:
        path: XLSX input path.
        sheet: Exact worksheet name to select.
        sheet_index: One-based worksheet index to select.

    Returns:
        A dictionary containing validated organisation records, all worksheet
        rows as literal or cached values, formula coordinates, and source
        provenance.

    Raises:
        ValueError: If worksheet selection, headings, identity values, or IDs
            are missing or ambiguous.
    """
    source_path = Path(path).expanduser().resolve()
    if not source_path.is_file():
        raise ValueError(f"Input workbook does not exist or is not a file: {source_path}")

    formulas_workbook = load_workbook(source_path, data_only=False)
    cached_workbook = load_workbook(source_path, data_only=True)
    try:
        formula_sheet = _select_sheet(
            formulas_workbook, sheet=sheet, sheet_index=sheet_index
        )
        cached_sheet = cached_workbook[formula_sheet.title]
        formula_cells = _formula_coordinates(formula_sheet)
        max_row = formula_sheet.max_row
        max_column = formula_sheet.max_column
        rows = [
            [cached_sheet.cell(row, column).value for column in range(1, max_column + 1)]
            for row in range(1, max_row + 1)
        ]

        id_header_rows = [
            row
            for row in range(1, max_row + 1)
            if _heading_role(formula_sheet.cell(row, 1).value) == "organisation_id"
        ]
        if len(id_header_rows) != 1:
            raise ValueError(
                "Expected exactly one Organisation/Organization ID header in column A; "
                f"found {len(id_header_rows)}"
            )
        header_row = id_header_rows[0]
        role_columns: dict[str, list[int]] = {}
        for column in range(1, max_column + 1):
            role = _heading_role(formula_sheet.cell(header_row, column).value)
            if role is not None:
                role_columns.setdefault(role, []).append(column)

        required_roles = (
            "name",
            "country",
            "membership_type",
            "membership_status",
        )
        missing_roles = [role for role in required_roles if role not in role_columns]
        if missing_roles:
            raise ValueError(
                "Missing required membership heading(s): "
                + ", ".join(missing_roles)
            )
        duplicate_roles = [
            role
            for role in ("name", "country", "membership_type", "membership_status", "acronym")
            if len(role_columns.get(role, [])) > 1
        ]
        if duplicate_roles:
            raise ValueError(
                "Ambiguous duplicated membership heading(s): "
                + ", ".join(duplicate_roles)
            )

        name_column = role_columns["name"][0]
        country_column = role_columns["country"][0]
        type_column = role_columns["membership_type"][0]
        status_column = role_columns["membership_status"][0]
        acronym_column = role_columns.get("acronym", [None])[0]
        identity_columns = {
            1,
            name_column,
            country_column,
            type_column,
            status_column,
        }
        if acronym_column is not None:
            identity_columns.add(acronym_column)

        organisations: list[dict[str, str | int]] = []
        seen_ids: dict[str, int] = {}
        nonfatal_formula_cells: set[tuple[int, int]] = set()
        missing_cache_cells: set[tuple[int, int]] = set()

        for row_number in range(header_row + 1, max_row + 1):
            raw_id = formula_sheet.cell(row_number, 1).value
            cached_id = rows[row_number - 1][0]
            if (row_number, 1) in formula_cells:
                if _is_blank(cached_id):
                    raise ValueError(
                        f"Missing cached required identity value at "
                        f"{formula_sheet.title}!A{row_number}"
                    )
                id_value = cached_id
            else:
                id_value = raw_id
            if _is_excel_error(id_value):
                raise ValueError(
                    f"Invalid Excel error identity value at "
                    f"{formula_sheet.title}!A{row_number}"
                )
            if _is_blank(id_value):
                # Blank IDs are non-membership/source rows, even when they
                # contain notes or other worksheet data.
                continue
            if _is_external(id_value):
                nonfatal_formula_cells.update(
                    coordinate
                    for coordinate in formula_cells
                    if coordinate[0] == row_number
                )
                continue
            organisation_id = _identifier_text(id_value, row_number=row_number)
            if organisation_id in seen_ids:
                raise ValueError(
                    f"Ambiguous duplicate Organisation ID {organisation_id!r} at rows "
                    f"{seen_ids[organisation_id]} and {row_number}"
                )
            seen_ids[organisation_id] = row_number

            field_values = {
                "name": rows[row_number - 1][name_column - 1],
                "country": rows[row_number - 1][country_column - 1],
                "membership_type": rows[row_number - 1][type_column - 1],
                "membership_status": rows[row_number - 1][status_column - 1],
            }
            for field_name, column in (
                ("name", name_column),
                ("country", country_column),
                ("membership_type", type_column),
                ("membership_status", status_column),
            ):
                field_value = field_values[field_name]
                if _is_excel_error(field_value):
                    raise ValueError(
                        f"Invalid Excel error identity value at "
                        f"{formula_sheet.title}!{get_column_letter(column)}{row_number}"
                    )
                if (
                    field_name != "country"
                    and (row_number, column) in formula_cells
                    and _is_blank(field_value)
                ):
                    raise ValueError(
                        f"Missing cached required identity value at "
                        f"{formula_sheet.title}!{get_column_letter(column)}{row_number}"
                    )
            if acronym_column is not None:
                acronym_value = rows[row_number - 1][acronym_column - 1]
                if _is_excel_error(acronym_value):
                    raise ValueError(
                        f"Invalid Excel error identity value at "
                        f"{formula_sheet.title}!{get_column_letter(acronym_column)}{row_number}"
                    )
            missing_fields = [
                name
                for name, value in field_values.items()
                if name != "country" and _is_blank(value)
            ]
            if missing_fields:
                raise ValueError(
                    f"Missing required identity value(s) at {formula_sheet.title}!row {row_number}: "
                    + ", ".join(missing_fields)
                )
            if _is_blank(field_values["country"]):
                log.warning(
                    "Missing country value at %s!%s%d; treating it as unknown country.",
                    formula_sheet.title,
                    get_column_letter(country_column),
                    row_number,
                )
            country_value = (
                "" if _is_blank(field_values["country"])
                else _text_value(field_values["country"])
            )

            for coordinate in formula_cells:
                if coordinate[0] == row_number and coordinate[1] not in identity_columns:
                    nonfatal_formula_cells.add(coordinate)
            for coordinate in formula_cells:
                if coordinate[0] == row_number and rows[coordinate[0] - 1][coordinate[1] - 1] is None:
                    missing_cache_cells.add(coordinate)

            organisations.append(
                {
                    "organisation_id": organisation_id,
                    "name": _text_value(field_values["name"]),
                    "acronym": (
                        ""
                        if acronym_column is None
                        or _is_blank(rows[row_number - 1][acronym_column - 1])
                        else _text_value(rows[row_number - 1][acronym_column - 1])
                    ),
                    "country": country_value,
                    "membership_type": _text_value(field_values["membership_type"]),
                    "membership_status": _text_value(field_values["membership_status"]),
                    "source_row": row_number,
                }
            )

        nonfatal_formula_cells.update(formula_cells)
        for coordinate in formula_cells:
            if rows[coordinate[0] - 1][coordinate[1] - 1] is None:
                missing_cache_cells.add(coordinate)
        if nonfatal_formula_cells:
            log.warning(
                "Worksheet %s contains %d non-fatal formula cell(s); cached values "
                "are returned as literals and formula cells remain occupied for export.",
                formula_sheet.title,
                len(nonfatal_formula_cells),
            )
        if missing_cache_cells:
            log.warning(
                "Worksheet %s contains %d formula cell(s) without cached values; "
                "they are returned as blank literals and cannot support identity.",
                formula_sheet.title,
                len(missing_cache_cells),
            )

        return {
            "organisations": organisations,
            "rows": rows,
            "formula_cells": [list(coordinate) for coordinate in sorted(formula_cells)],
            "sheet_name": formula_sheet.title,
            "header_row": header_row,
            "workbook_sha256": _workbook_sha256(source_path),
            "source_path": str(source_path),
            "contributor_columns": role_columns.get("contributor", []),
        }
    finally:
        formulas_workbook.close()
        cached_workbook.close()


def _write_literal(cell, value: Any) -> None:
    """Write a value without allowing strings beginning with ``=`` to become formulas."""
    cell.value = value
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"


def _validate_cell_lengths(rows: Iterable[Iterable[Any]], sheet_name: str) -> None:
    """Reject values Excel cannot represent instead of silently truncating them."""
    for row_number, row in enumerate(rows, start=1):
        for column_number, value in enumerate(row, start=1):
            if isinstance(value, str) and len(value) > EXCEL_MAX_CELL_CHARS:
                raise ValueError(
                    f"Value at {sheet_name}!{get_column_letter(column_number)}{row_number} "
                    f"exceeds Excel's {EXCEL_MAX_CELL_CHARS}-character cell limit"
                )


def _format_header(worksheet, row_number: int, column_count: int) -> None:
    """Apply restrained header formatting and a filter to a worksheet."""
    fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    for column in range(1, column_count + 1):
        cell = worksheet.cell(row_number, column)
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    worksheet.freeze_panes = f"A{row_number + 1}"
    worksheet.auto_filter.ref = (
        f"A{row_number}:{get_column_letter(column_count)}{worksheet.max_row}"
    )


def _unique_source_sheet_name(source_name: str) -> str:
    """Choose a valid source-sheet name that does not conflict with the summary."""
    if source_name.casefold() != SUMMARY_SHEET_NAME.casefold():
        return source_name
    return "Source"


def _biobank_ids_text(value: Any) -> str:
    """Render a match's biobank ID collection as one comma-separated cell."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    try:
        return ",".join(str(item) for item in value)
    except TypeError:
        return str(value)


def _validate_matches(matches: list[Mapping[str, Any]]) -> None:
    """Validate the small match interface before creating any output."""
    required = {
        "organisation_id",
        "eosc_name",
        "directory_juridical_person",
        "biobank_ids",
    }
    for index, match in enumerate(matches):
        if not isinstance(match, Mapping):
            raise ValueError(f"Match {index} is not a mapping")
        missing = required - set(match)
        if missing:
            raise ValueError(
                f"Match {index} is missing required key(s): {', '.join(sorted(missing))}"
            )


def _validate_written_workbook(
    path: Path, source_sheet_name: str, source_rows: list[list[Any]]
) -> None:
    """Reopen a staged workbook and verify its essential values-only contract."""
    workbook = load_workbook(path, data_only=False)
    try:
        if workbook.sheetnames != [SUMMARY_SHEET_NAME, source_sheet_name]:
            raise ValueError(
                "Generated workbook has unexpected sheet names: "
                + repr(workbook.sheetnames)
            )
        summary = workbook[SUMMARY_SHEET_NAME]
        if [summary.cell(1, column).value for column in range(1, 5)] != list(
            SUMMARY_HEADERS
        ):
            raise ValueError("Generated workbook has an invalid summary header")
        source = workbook[source_sheet_name]
        if source.max_row != len(source_rows):
            raise ValueError("Generated workbook changed the source row count")
        for row_number, expected_row in enumerate(source_rows, start=1):
            for column_number, expected_value in enumerate(expected_row, start=1):
                actual_value = source.cell(row_number, column_number).value
                semantically_blank = {None, ""}
                if expected_value in semantically_blank and actual_value in semantically_blank:
                    continue
                if actual_value != expected_value:
                    raise ValueError(
                        "Generated workbook changed a source literal at "
                        f"{source_sheet_name}!{get_column_letter(column_number)}{row_number}"
                    )
        for worksheet in workbook.worksheets:
            if any(cell.data_type == "f" for row in worksheet.iter_rows() for cell in row):
                raise ValueError("Generated workbook unexpectedly contains a formula")
    finally:
        workbook.close()



def _publish_without_overwrite(temporary_path: Path, destination: Path) -> None:
    """Publish a validated temporary file atomically without replacing a file."""
    unsupported_link_errors = {
        error_number
        for error_number in (errno.EOPNOTSUPP, errno.ENOSYS, errno.EPERM)
        if error_number is not None
    }
    if hasattr(os, "link"):
        try:
            os.link(temporary_path, destination)
            return
        except FileExistsError as error:
            raise FileExistsError(
                f"Output workbook appeared during publication: {destination}"
            ) from error
        except OSError as error:
            if error.errno not in unsupported_link_errors:
                raise

    try:
        libc = ctypes.CDLL(None, use_errno=True)
        renameat2 = libc.renameat2
    except AttributeError:
        renameat2 = None
    if renameat2 is not None:
        renameat2.restype = ctypes.c_int
        renameat2.argtypes = [
            ctypes.c_int,
            ctypes.c_char_p,
            ctypes.c_int,
            ctypes.c_char_p,
            ctypes.c_uint,
        ]
        result = renameat2(
            -100,
            os.fsencode(temporary_path),
            -100,
            os.fsencode(destination),
            1,
        )
        if result == 0:
            return
        error_number = ctypes.get_errno()
        if error_number == errno.EEXIST:
            raise FileExistsError(
                f"Output workbook appeared during publication: {destination}"
            )
        if error_number not in unsupported_link_errors:
            raise OSError(
                error_number,
                f"Could not atomically publish output workbook: {os.strerror(error_number)}",
            )

    raise OSError(
        "Cannot atomically publish output workbook without a supported "
        "no-overwrite primitive (hard link or renameat2)"
    )

def write_matches_xlsx(
    workbook: Mapping[str, Any],
    matches: list[Mapping[str, Any]],
    path: str | os.PathLike[str],
) -> None:
    """Write a fresh, values-only summary and source rendition workbook.

    Args:
        workbook: The dictionary returned by :func:`read_membership`.
        matches: Match mappings requiring ``organisation_id``, ``eosc_name``,
            ``directory_juridical_person``, and ``biobank_ids``.
        path: New XLSX destination. Existing files and source aliases are rejected.

    Raises:
        ValueError: If the destination is unsafe, the match interface is
            invalid, or source values exceed Excel's cell limit.
        FileExistsError: If a destination appears during atomic publication.
    """
    if not isinstance(workbook, Mapping):
        raise ValueError("workbook must be the mapping returned by read_membership")
    match_list = list(matches)
    _validate_matches(match_list)

    source_path = Path(str(workbook["source_path"])).expanduser().resolve()
    destination_input = Path(path).expanduser()
    if destination_input.is_symlink():
        raise FileExistsError(
            f"Output path is an existing symlink and is rejected: {destination_input}"
        )
    destination = destination_input.resolve()
    if destination == source_path:
        raise ValueError("Output workbook must not alias the source workbook")
    if destination.exists():
        raise FileExistsError(f"Output workbook already exists: {destination}")
    if not destination.parent.is_dir():
        raise ValueError(f"Output directory does not exist: {destination.parent}")

    try:
        source_rows = [list(row) for row in workbook["rows"]]
        header_row = int(workbook["header_row"])
        source_name = str(workbook["sheet_name"])
        contributor_columns = [int(column) for column in workbook["contributor_columns"]]
        formula_cells = {
            (int(coordinate[0]), int(coordinate[1]))
            for coordinate in workbook["formula_cells"]
        }
    except (KeyError, TypeError, ValueError, IndexError) as error:
        raise ValueError("workbook is missing a valid read_membership field") from error
    if not source_rows or not 1 <= header_row <= len(source_rows):
        raise ValueError("workbook has no valid source rows/header_row")
    if any(column < 1 for column in contributor_columns):
        raise ValueError("workbook contains an invalid contributor column")

    source_column_count = max(
        max((len(row) for row in source_rows), default=0),
        max(contributor_columns, default=0),
    )
    for row in source_rows:
        row.extend([None] * (source_column_count - len(row)))

    _validate_cell_lengths(source_rows, source_name)
    summary_rows: list[list[Any]] = []
    for match in match_list:
        summary_rows.append(
            [
                str(match["organisation_id"]),
                match["eosc_name"],
                match["directory_juridical_person"],
                _biobank_ids_text(match["biobank_ids"]),
            ]
        )
    _validate_cell_lengths([SUMMARY_HEADERS, *summary_rows], SUMMARY_SHEET_NAME)

    target_ids = {str(match["organisation_id"]).strip() for match in match_list}
    source_id_rows: dict[str, list[int]] = {}
    for row_number in range(header_row + 1, len(source_rows) + 1):
        raw_id = source_rows[row_number - 1][0]
        if _is_blank(raw_id) or _is_external(raw_id):
            continue
        try:
            organisation_id = _identifier_text(raw_id, row_number=row_number)
        except ValueError:
            continue
        source_id_rows.setdefault(organisation_id, []).append(row_number)
    target_rows = [
        row_number
        for organisation_id in sorted(target_ids)
        for row_number in source_id_rows.get(organisation_id, [])
    ]

    if formula_cells:
        log.warning(
            "Formula cells from source sheet will be written as cached literal values; "
            "occupied formula slots are not treated as empty contributor slots."
        )

    pending_new_column_rows: list[int] = []
    for row_number in target_rows:
        row = source_rows[row_number - 1]
        if any(
            _normalise_heading(row[column - 1]) == _normalise_heading(CONTRIBUTOR_TAG)
            for column in contributor_columns
            if row[column - 1] is not None
        ):
            continue
        empty_column = next(
            (
                column
                for column in contributor_columns
                if (row_number, column) not in formula_cells
                and _is_blank(row[column - 1])
            ),
            None,
        )
        if empty_column is not None:
            row[empty_column - 1] = CONTRIBUTOR_TAG
        else:
            pending_new_column_rows.append(row_number)
    if pending_new_column_rows:
        new_column = source_column_count + 1
        for row in source_rows:
            row.append(None)
        source_rows[header_row - 1][new_column - 1] = "Node Contributor"
        for row_number in pending_new_column_rows:
            source_rows[row_number - 1][new_column - 1] = CONTRIBUTOR_TAG

    _validate_cell_lengths(source_rows, source_name)
    source_sheet_name = _unique_source_sheet_name(source_name)
    output_workbook = Workbook()
    summary_sheet = output_workbook.active
    summary_sheet.title = SUMMARY_SHEET_NAME
    source_sheet = output_workbook.create_sheet(source_sheet_name)
    for row_number, row in enumerate([list(SUMMARY_HEADERS), *summary_rows], start=1):
        for column_number, value in enumerate(row, start=1):
            _write_literal(summary_sheet.cell(row_number, column_number), value)
    _format_header(summary_sheet, 1, len(SUMMARY_HEADERS))
    for column in range(1, 5):
        summary_sheet.column_dimensions[get_column_letter(column)].width = (
            18 if column == 1 else 42 if column in (2, 3) else 60
        )
    for row_number in range(2, summary_sheet.max_row + 1):
        for column in range(1, 5):
            summary_sheet.cell(row_number, column).alignment = Alignment(
                vertical="top", wrap_text=True
            )

    for row_number, row in enumerate(source_rows, start=1):
        for column_number, value in enumerate(row, start=1):
            _write_literal(source_sheet.cell(row_number, column_number), value)
        if all(_is_blank(value) for value in row):
            # Openpyxl omits entirely empty trailing rows unless a row
            # dimension keeps their position in the saved worksheet.
            source_sheet.row_dimensions[row_number].height = 15
    _format_header(source_sheet, header_row, len(source_rows[header_row - 1]))
    for column in range(1, source_sheet.max_column + 1):
        source_sheet.column_dimensions[get_column_letter(column)].width = 24
    for row_number in range(1, source_sheet.max_row + 1):
        for column in range(1, source_sheet.max_column + 1):
            source_sheet.cell(row_number, column).alignment = Alignment(
                vertical="top", wrap_text=True
            )

    temporary_path: Path | None = None
    try:
        descriptor, temporary_name = tempfile.mkstemp(
            prefix=f".{destination.stem}.", suffix=".tmp.xlsx", dir=destination.parent
        )
        os.close(descriptor)
        temporary_path = Path(temporary_name)
        output_workbook.save(temporary_path)
        _validate_written_workbook(temporary_path, source_sheet_name, source_rows)
        _publish_without_overwrite(temporary_path, destination)
    finally:
        output_workbook.close()
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()
